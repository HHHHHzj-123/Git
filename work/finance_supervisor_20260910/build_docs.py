from pathlib import Path
import json
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent; O=R/'deliverables'; I=R/'images'; I.mkdir(exist_ok=True)
D=json.loads((R/'content.json').read_text(encoding='utf-8')); NAVY='17365D'; PALE='F3F6F9'; AMBER='FFF2CC'
def fnt(s):
    p=Path(r'C:\Windows\Fonts\msyh.ttc'); return ImageFont.truetype(str(p),s,index=0) if p.exists() else ImageFont.load_default()
def flow(name,title,levels):
    w=1500;m=60;bh=95;gap=42;h=145+len(levels)*(bh+gap);im=Image.new('RGB',(w,h),'white');d=ImageDraw.Draw(im);d.text((m,28),title,font=fnt(38),fill='black');y=105
    for ri,row in enumerate(levels):
        bw=(w-2*m-(len(row)-1)*25)/len(row)
        for j,s in enumerate(row):
            x=m+j*(bw+25);d.rounded_rectangle((x,y,x+bw,y+bh),12,fill=(232,242,250) if ri%2==0 else (245,247,249),outline=(47,117,181),width=3)
            lines=[];cur=''
            for ch in s:
                if d.textlength(cur+ch,font=fnt(21))>bw-20:lines.append(cur);cur=ch
                else:cur+=ch
            if cur:lines.append(cur)
            for k,t in enumerate(lines[:3]):tw=d.textlength(t,font=fnt(21));d.text((x+(bw-tw)/2,y+22+k*27),t,font=fnt(21),fill='black')
            if j<len(row)-1:d.line((x+bw+4,y+bh/2,x+bw+18,y+bh/2),fill=(90,90,90),width=3)
        if ri<len(levels)-1:d.line((w/2,y+bh+4,w/2,y+bh+28),fill=(90,90,90),width=3)
        y+=bh+gap
    im.save(I/name,dpi=(180,180))
flow('upgrade.png','总账会计到财务主管',[['自己做','自己核对','自己结账'],['分配任务','管理进度','风险复核'],['协调业务','升级问题','验证关闭'],['建立制度','优化流程','培养团队'],['对财务结果负责']])
flow('loop.png','财务主管管理闭环',[['确定目标'],['拆分任务','指定R/A'],['截止时间','完成证据'],['过程监控','风险复核'],['例外与升级'],['Closure验证'],['根因与流程优化']])
flow('close.png','月结Close Management',[['T-5截止通知'],['T-3业务缺口'],['T-2关键路径'],['T-1最终输入'],['T业务截止'],['T+1子模块'],['T+2成本税务GL'],['T+3报表锁账']])
flow('issue.png','Issue到Closure',[['Open 确认问题'],['In Progress 执行'],['Pending 等待外部条件'],['Resolved 动作完成'],['验证账务/系统/资料/报表税务'],['Closed 正式关闭'],['RCA 防止重复']])
flow('control.png','企业内控落地',[['风险'],['控制目标'],['控制活动'],['责任人'],['系统与权限'],['执行证据'],['主管复核'],['例外与改进']])
flow('payment.png','100万元付款控制链',[['业务申请','预算/合同'],['采购/法务审核'],['收货或验收'],['AP三单匹配'],['分级审批'],['资金制单'],['授权人银行复核'],['回单/核销/主管抽查']])
flow('emergency.png','紧急事项处置',[['第一小时：止损、保全证据、限制权限'],['第一天：量化影响、内部升级、外部联络'],['后续：账务税务处理、追偿、根因整改'],['主管验证Closure']])

def setup(title,subtitle):
    doc=Document();sec=doc.sections[0];sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.05);sec.right_margin=Cm(2.0)
    for st,sz in [('Normal',10),('Title',25),('Heading 1',18),('Heading 2',14),('Heading 3',11.2)]:
        s=doc.styles[st];s.font.name='Microsoft YaHei';s._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑');s.font.size=Pt(sz)
        if st!='Normal':s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.keep_with_next=True
    doc.styles['Normal'].paragraph_format.line_spacing=1.22;doc.styles['Normal'].paragraph_format.space_after=Pt(4)
    h=sec.header.paragraphs[0];h.text=title;h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in h.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(100,100,100)
    ft=sec.footer.paragraphs[0];ft.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');ft._p.append(fld)
    doc.add_paragraph().paragraph_format.space_after=Pt(75);p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(title)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(subtitle);r.bold=True;r.font.size=Pt(17)
    doc.add_paragraph().paragraph_format.space_after=Pt(70)
    for s in ['案例企业：启辰智能硬件科技有限公司','版本日期：2026年9月10日','适用岗位：总账会计、财务主管、财务经理']:
        p=doc.add_paragraph(s);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break();return doc
def shade(c,color):
    pr=c._tc.get_or_add_tcPr();e=pr.find(qn('w:shd'))
    if e is None:e=OxmlElement('w:shd');pr.append(e)
    e.set(qn('w:fill'),color)
def tab(doc,heads,rows,fs=7.4):
    t=doc.add_table(rows=1,cols=len(heads));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(heads):
        c=t.rows[0].cells[i];c.text=str(h);shade(c,NAVY);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(fs)
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text='' if v is None else str(v);cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2:shade(cs[i],PALE)
            for p in cs[i].paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.03
                for r in p.runs:r.font.size=Pt(fs)
    pr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');rep.set(qn('w:val'),'true');pr.append(rep);doc.add_paragraph().paragraph_format.space_after=Pt(0);return t
def p(doc,text,lead=None):
    x=doc.add_paragraph();
    if lead:x.add_run(lead).bold=True
    x.add_run(text);return x
def bullets(doc,items):
    for x in items:doc.add_paragraph(x,style='List Bullet')
def pic(doc,name,cap):
    x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.add_run().add_picture(str(I/name),width=Cm(16.3));x=doc.add_paragraph(cap);x.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in x.runs:r.font.size=Pt(8);r.font.italic=True;r.font.color.rgb=RGBColor(90,90,90)
def toc(doc):
    x=doc.add_paragraph();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u');x._p.append(fld)
def page(doc):doc.add_page_break()

# Main handbook
doc=setup('财务主管管理、内控与跨部门协同实操手册','让一个财务团队稳定地把账做对')
doc.add_heading('使用说明',0);p(doc,'本手册以责任人、截止时间、复核证据、例外、升级、Issue和Closure为主线。文中“建议”属于企业管理或内控设计，不等同于法律强制要求；具体权限金额应按企业规模、治理结构和风险承受能力制定。')
doc.add_heading('目录',0);toc(doc);page(doc)
doc.add_heading('第一篇 从总账会计升级为财务主管',0)
doc.add_heading('第1章 能力升级图',1);pic(doc,'upgrade.png','图1 总账会计到财务主管的能力升级')
tab(doc,['维度','优秀总账会计','财务主管','主管的完成证据'],[['工作对象','自己的凭证、对账和报表','整个团队和跨部门依赖','工作分配表、RACI'],['时间','自己按期完成','管理关键路径和阻断事项','Close Calendar'],['质量','发现和纠正错误','通过流程、权限和复核降低错误','Review Checklist'],['问题','解决具体差异','管理Issue、升级和Closure','Issue Log'],['改进','提高个人效率','消除重复问题和无效控制','RCA、Action Plan']],7.7)
doc.add_heading('第2章 财务主管管理闭环',1);pic(doc,'loop.png','图2 主管工作的最小管理闭环')
doc.add_heading('第3章 授权与亲自介入',1)
tab(doc,['事项特征','建议方式','主管动作'],[['重复、标准、低风险','授权执行','定义输入、输出和抽查规则'],['金额大但系统稳定','下属执行+抽查','查看控制总数和异常'],['判断复杂','主管参与','核实Fact、准则和备选方案'],['法律税务或资金高风险','主管重点复核','必要时找法务/税务顾问'],['重大异常或管理层指令','主管负责','书面留痕、升级、拒绝违法执行']],7.6)
doc.add_heading('第4章 30组错误做法与成熟做法',1);tab(doc,['编号','错误做法','成熟做法'],[[x['id'],x['wrong'],x['mature']] for x in D['comparisons']],7.2);page(doc)

doc.add_heading('第二篇 团队、边界和RACI',0)
doc.add_heading('第5章 案例公司和组织架构',1);tab(doc,['项目','设定'],[[k,v] for k,v in D['company'].items()],8)
doc.add_heading('第6章 财务岗位职责与边界',1);tab(doc,['岗位','主要职责','主管复核重点','不可模糊边界'],D['roles'],7.2)
doc.add_heading('第7章 财务流程RACI',1);p(doc,'R是执行，A是最终对结果负责，C在决策前被咨询，I在结果形成后获知。同一任务可以有多个R，但原则上只有一个清晰的A。')
tab(doc,['流程','R执行','A最终负责','C协商','I知会'],D['raci'],7.1)
doc.add_heading('第8章 Definition of Done',1)
tab(doc,['岗位','“做完”不能只意味着','主管应取得的证据'],[['AP','发票已录入','AP/GL核对、暂估账龄、重复付款检查、异常清单'],['AR','客户余额已核对','AR/GL核对、未认领回款、账龄、争议和期后回款'],['成本','成本程序运行成功','工单状态、负库存、成本差异、毛利桥、存货/GL核对'],['税务','申报按钮已提交','账票税桥、申报回执、缴款凭证、差异解释'],['总账','报表能打开','模块对账、调整清单、报表勾稽、未决事项批准']],7.4)
doc.add_heading('第9章 工作日历与精力分配',1)
tab(doc,['周期','高频关注','主管输出','精力特点'],[['每日','资金、重大付款、关键Issue、异常凭证','决策、升级或授权','时间短但风险高'],['每周','应收、挂账、付款计划、Action','周会行动清单','主要用于防止月末堆积'],['月初','报表、税务、复盘','最终报表和改进任务','最集中'],['月中','清理、流程、培训','中期预警','适合处理根因'],['月末','截止、库存、成本、工资','Close状态和例外批准','最耗主管精力'],['季度/年度','盘点、所得税、审计、预算','专项项目计划','低频但复杂']],7.3);page(doc)

doc.add_heading('第三篇 月结Close Management',0)
doc.add_heading('第10章 月结全景',1);pic(doc,'close.png','图3 T-5至T+3月结关键路径')
doc.add_heading('第11章 Close Calendar详表',1);tab(doc,['时间','团队任务','R','证据','异常标准','升级'],D['close'],6.8)
doc.add_heading('第12章 任务依赖和阻断判断',1)
tab(doc,['前置任务','后续任务','可否暂缓','主管判断'],[['工资锁版','薪酬接口/计提','通常不可','影响工资、个税和费用归属'],['库存锁定/工单关闭','成本运行','不可','数据仍变动时成本结果不稳定'],['AP/AR子模块完成','GL对账','不可','总账余额仍可能变化'],['税务测算','最终税费凭证','视申报日和重要性','需可靠估计并记录差异'],['非重大资料补充','报表','可经批准暂留','必须量化影响并设Owner和期限']],7.5)
doc.add_heading('第13章 贯穿式完整月份',1)
month=[
['T-5','发布日历；识别本月大额销售、预付款和奖金','主管核对Owner、请假覆盖、系统窗口','无Owner任务立即补齐','Calendar和重大事项清单'],['T-3','暂估缺3项、验收缺失、负库存','采购/AP、销售/AR、仓库/IT分三条Issue','预计T前不能完成则升部门负责人','缺口清单和第一次催收'],['T-1','工资表修改、工单未关','HR提交差异版；生产按工单逐项承诺','阻断成本或发薪则升财务经理/厂长','锁版记录和工单状态'],['T','库存截止和例外处理','只批准有原因、期限和补救的例外','无批准跨期操作不得执行','锁库日志和例外清单'],['T+1','AP完成；120万元回款未认领；BOM异常','AR和销售认领；工程确认BOM；成本评估重跑','重大收入/成本差异立即升级','子模块对账和异常证据'],['T+2','销项差异；应付借方','税务做账票税桥；AP查余额实质','影响申报/报表真实性不得带过','税务桥和余额处理结论'],['T+3','出报表','复核重大调整、毛利、现金流、Issue','无法量化的重大事项阻断出表','最终报表、复核签字、锁账']]
tab(doc,['时点','发生事项','主管先做什么','何时升级','当天Closure证据'],month,6.9)
doc.add_heading('第14章 Close Daily与月结复盘',1);p(doc,'Close Daily只讨论红黄事项、决策和下一动作，建议10至15分钟。正常完成项不逐条汇报。月结复盘只选择重复发生、影响关键路径或产生重大风险的问题进入流程优化。')
doc.add_heading('【面试可能怎么问】',2);bullets(doc,['你怎么保证月结及时？','成本延误会如何管理？','哪些事项可以作为未决事项带入下月？','如何判断是否延后出表？']);page(doc)

doc.add_heading('第四篇 风险导向复核',0)
doc.add_heading('第15章 五维复核模型',1);p(doc,'金额、风险、判断、异常变化和历史问题共同决定复核深度。金额小但涉及关联方、老板私人支出、虚假发票或越权付款，仍应作为高风险事项。')
tab(doc,['因素','高风险信号','复核方法'],[['金额重大','超过内部重要性或单笔异常','100%检查事实和批准'],['风险重大','资金、税务、法律、舞弊','提高层级并检查原始证据'],['判断复杂','收入、减值、资本化','Accounting Memo'],['异常变化','环比/同比或业务规律异常','趋势分析后追明细'],['历史问题','反复出现或曾被审计调整','验证整改控制']],7.6)
doc.add_heading('第16章 月度复核清单',1);tab(doc,['项目','为什么复核','看什么','异常标准','处理','Closure证据'],D['reviews'],6.4)
doc.add_heading('第17章 主管抽样原则',1);bullets(doc,['全检：重大、异常、高判断、关联方、管理层指令和历史问题。','抽查：高金额但标准化、系统自动且控制稳定的事项。','分析复核：数量大、规律稳定的薪酬、折旧、标准费用。','轮换复核：低风险流程按月轮换，确保全年覆盖。'])
doc.add_heading('【面试可能怎么问】',2);bullets(doc,['主管是否需要逐张复核凭证？','如何复核成本会计说“成本已跑完”？','应付借方余额应如何处理？']);page(doc)

doc.add_heading('第五篇 跨部门协同和催收',0)
doc.add_heading('第18章 各部门需要提供什么',1)
tab(doc,['部门','财务需要','时间','不提供的影响','提前预防'],[['销售','合同、发货、签收验收、折扣、退货','T-3前','收入截止、应收和税务错误','合同ID和验收Owner'],['采购','PO、收货、暂估、预计到票、预付进度','T-3前','漏债、成本、资金风险','系统自动未收票清单'],['仓库','出入库、盘点、负库存和截止确认','T/T+1','库存和成本无法稳定','锁库和例外流程'],['生产','工单、报工、BOM和在制','T-1','成本无法运行','日常预警而非月底催'],['HR','人员、工资、奖金、离职和成本中心','T-1锁版','薪酬、个税和费用错误','变更版本和差异表'],['研发','项目、领料、工时、状态和预算','T-2','研发费用和资本化失真','项目字段必填'],['法务','合同变更、诉讼和法律意见','事件触发/月度','预计负债、付款和披露遗漏','财务会签及月度清单'],['IT','接口、权限、变更和运行日志','每日/关键期','漏传、重复、月结中断','财务冻结窗口']],6.9)
doc.add_heading('第19章 资料催收与升级SOP',1)
tab(doc,['阶段','表达重点','留痕','升级'],[['第一次提醒','资料、责任人、准确截止时间','工作群或系统任务','通常不升级'],['第二次提醒','剩余缺口及影响','邮件/OA','必要时抄送直接负责人'],['临近截止','明确会阻断什么','Issue Log','部门负责人'],['超过截止','登记逾期、下一动作和新期限','正式邮件+Issue','财务主管/部门负责人'],['影响月结申报','量化金额和后果','Close状态表','财务经理/CFO'],['涉嫌违法或重大资金','停止执行并保全证据','书面记录','法务、CFO、治理层']],7.2)
doc.add_heading('第20章 沟通原则',1);p(doc,'有效催收不是重复说“尽快”，而是写清楚事项、金额或范围、准确时间、业务影响、需要对方完成的动作和无法完成时的反馈要求。')
doc.add_heading('【面试可能怎么问】',2);bullets(doc,['业务部门不配合怎么办？','如何避免财务被认为“卡流程”？','资料晚一天是否都要升级？']);page(doc)

doc.add_heading('第六篇 Issue、升级与Closure',0)
doc.add_heading('第21章 Issue生命周期',1);pic(doc,'issue.png','图4 问题从Open到Closed')
doc.add_heading('第22章 Issue Log字段',1)
tab(doc,['字段组','核心字段','主管用途'],[['识别','编号、发现日、问题、流程、科目、金额','确定问题边界'],['风险','等级、月结/税务/报表/现金影响','排序和升级'],['责任','责任部门、Owner、财务Owner','避免大家以为别人会做'],['措施','临时措施、永久措施、下一步','止损与根因分开'],['时间','截止日、预计/实际完成日','跟踪逾期'],['关闭','证据、复核人、关闭状态和日期','验证Closure']],7.4)
doc.add_heading('第23章 问题升级矩阵',1)
tab(doc,['触发因素','一般','重要','重大/紧急'],[['金额','低于部门阈值','超过主管阈值','超过CFO阈值或资金安全'],['时间','不影响关键路径','可能影响月结','影响申报、出表或付款'],['法律税务','一般资料瑕疵','可能补税/处罚','涉嫌违法或故意虚假'],['跨部门','首次延迟','连续两次或拒绝','部门负责人仍不处理'],['报表','可量化不重大','重大但可修复','无法可靠量化或重大错报'],['升级对象','业务Owner','部门负责人/财务经理','CFO/总经理/法务/外部专家']],7.2)
doc.add_heading('第24章 根因分析',1);p(doc,'“采购没给暂估”只是现象。继续追问：为什么没给→没有Owner→PO无预计到票日→ERP无预警→财务月底才催。永久措施应落在PO字段、预警、责任人和提前催收，而不是再发一封更严厉的邮件。')
doc.add_heading('【面试可能怎么问】',2);bullets(doc,['什么叫问题真正关闭？','长期挂账如何推进？','如何处理持续Pending的问题？']);page(doc)

doc.add_heading('第七篇 内控、审批和权限',0)
doc.add_heading('第25章 内控落地模型',1);pic(doc,'control.png','图5 从风险到复核证据')
p(doc,'《企业内部控制基本规范》及配套指引为纳入实施范围的企业提供规范框架；非上市中型企业可结合规模和风险参照设计。手册中的具体审批层级和金额是管理建议，不是法定统一标准。')
doc.add_heading('第26章 采购付款内控',1);pic(doc,'payment.png','图6 100万元付款审批与银行执行链')
tab(doc,['风险','关键控制','可按金额简化','不可轻易简化'],[['虚假供应商','准入、税号和账户验证','低额标准供应商可批量复核','银行账户变更独立核验'],['重复付款','发票唯一性、三单匹配、付款前扫描','低额报销可系统去重','同一申请制单和最终支付'],['无合同/验收','合同、PO、验收和例外审批','低额零星采购可简化PO','大额预付和陌生供应商'],['紧急付款','原因、临时批准、补件期限、事后复盘','审批路径可加速','账户核验和支付授权']],7.1)
doc.add_heading('第27章 销售收款内控',1);p(doc,'客户主数据、信用额度、销售价格、发货、收入判断、退款和核销需要合理分工。销售可以提供事实和商业理由，但财务应独立判断收入确认。')
doc.add_heading('第28章 资金、费用、存货和固定资产内控',1)
tab(doc,['模块','高风险权限/事项','主管复核'],[['资金','制单、审核、付款、U盾、印鉴、账户','大额/异常付款、未达和权限冲突'],['费用','老板支出、招待礼品、咨询服务','真实性、受益人、税务和审批'],['存货生产','收货、领料、BOM、工单、盘点、报废','负库存、异常调整、差异和废料'],['固定资产','验收、卡片、调拨、报废、出售','账实、转固、闲置和个人占用']],7.3)
doc.add_heading('第29章 税务内控',1);p(doc,'准备、复核、提交、缴税和归档应有明确责任。主管不必重新填报每个单元格，但要取得账票税桥、异常差异、申报回执和缴款凭证。')
doc.add_heading('第30章 ERP权限矩阵',1)
tab(doc,['权限组合','风险','处理'],[['供应商账户维护+付款执行','可将款项导向异常账户','职责分离；小团队设置独立回拨和事后日志复核'],['凭证录入+审核','可自行完成未经复核的调整','分离或限定低风险范围'],['反结账+报表发布','可改变已发布结果','反结账需专项批准和变更日志'],['系统管理员+业务审批','可修改权限并执行交易','管理员不承担业务审批'],['税务申报+税款最终支付','错误或舞弊缺少独立发现','主管复核申报，授权人支付']],7.2)
doc.add_heading('第31章 内控与效率平衡',1);p(doc,'控制强度取决于金额、风险、交易类型、系统能力和人员成本。低额、标准、重复交易可用系统规则、抽查和额度简化；账户变更、关联方、重大预付、退款、资产处置和异常管理层指令即使金额不大也需要强化。')
doc.add_heading('第32章 例外事项管理',1);p(doc,'例外不是无控制。每项例外必须有真实原因、授权人、适用期限、补救动作、失效日期和关闭证据。长期重复出现的“例外”说明标准流程需要调整或业务在规避控制。');page(doc)

doc.add_heading('第八篇 团队、判断、审计和应急',0)
doc.add_heading('第33章 团队管理和交叉备岗',1);tab(doc,['管理动作','实际做法','证据'],[['工作分配','按周期、风险和容量分配','任务表'],['优先级','资金/税务/报表/关键路径优先','周会Action'],['请假覆盖','每个关键任务有主备岗','备岗矩阵和实操记录'],['新人培养','业务→系统→科目→独立操作→月结','30天计划'],['错误管理','区分能力、粗心、流程、系统、资料和故意','RCA与改进计划']],7.5)
doc.add_heading('第34章 新员工30天计划',1);tab(doc,['阶段','总账/AP/AR/成本共同重点','通过标准'],[['第1周','业务流程、单据链、系统和权限','能从一笔交易追到GL'],['第2周','标准操作和台账','在监督下完成日常任务'],['第3周','对账和异常处理','解释差异并保留证据'],['第4周','参与月结','完成指定模块并通过主管复核']],7.5)
doc.add_heading('第35章 重大会计判断和Accounting Memo',1);p(doc,'流程为Fact→适用准则→会计问题→备选方案→财务和税务影响→建议→审批→实际入账→后续复核。主管不知道答案时，应根据问题性质找业务、IT、税务顾问、审计师或律师，不应装作全部掌握。')
doc.add_heading('第36章 异常管理层指令',1);p(doc,'先确认事实和指令，再解释会计、税务、法律和公司治理后果，提出合法替代方案，书面留痕并按层级升级。明显违法事项应拒绝执行并保护会计资料。')
doc.add_heading('第37章 年度审计管理',1);p(doc,'由财务主管统一接收PBC、指定资料Owner、控制版本、集中回复审计Issue。审计师可与资料负责人沟通事实，但新增资料要求和审计调整应回到统一清单，避免无限分散向团队要资料。')
doc.add_heading('第38章 紧急事项',1);pic(doc,'emergency.png','图7 财务紧急事项处理节奏')
tab(doc,['事项','第一小时','第一天','后续'],[['被骗付款','停止支付、联系银行、保全通信和权限','升级CFO/法务并报案评估','追偿、账务、控制整改'],['账户冻结','确认范围和原因、保护其他账户','资金预测和法律税务联络','解除、披露及替代支付'],['税务逾期','确认税种金额、停止继续错误','联系税务机关并准备更正','缴税滞纳、RCA'],['ERP宕机','冻结线下变更、启用应急记录','恢复和控制总数校验','补录、对账和灾备改进'],['重大盘亏/舞弊','保护现场和系统日志','独立调查并限制权限','损失、追责、披露和整改']],6.8);page(doc)

doc.add_heading('第九篇 流程优化和主管能力',0)
doc.add_heading('第39章 20个流程优化案例',1);tab(doc,['编号','主题','原流程/痛点','新流程','减少工作','增加控制'],[[x['id'],x['topic'],x['old'],x['new'],x['reduced'],x['control']] for x in D['optimizations']],6.5)
doc.add_heading('第40章 三种规模企业的主管差异',1);tab(doc,['企业','主管现实','重点','不应照搬'],[['50人小企业','兼总账、税务和部分资金','关键资金、申报、月结、基础分工','大型集团复杂会议和审批'],['300-500人制造企业','管理AP/AR/成本/总账等专业岗位','Close、复核、协同、Issue和流程','自己完成全部明细工作'],['大型集团','共享中心、BP、集团报告和中心职能分层','SLA、政策、合并、权限和数据治理','用个人催收替代系统和治理']],7.2)
doc.add_heading('第41章 财务主管20项核心能力',1);bullets(doc,['任务拆分和RACI','Close Calendar','关键路径管理','Definition of Done','风险导向复核','重大异常识别','Issue状态管理','Closure验证','问题升级','跨部门催收','业务影响沟通','内控效率平衡','例外管理','审批授权','ERP权限','Accounting Memo','根因分析','团队培养','交叉备岗','现实约束下的风险排序'])
doc.add_heading('第42章 90天上岗行动计划',1);tab(doc,['阶段','目标','关键输出'],[['1-30天','理解业务、系统、团队和历史问题','组织图、RACI、月结地图、Top 10风险'],['31-60天','接管月结和复核','Close Calendar、Review Checklist、Issue Log'],['61-90天','稳定团队并改进重复问题','备岗计划、流程优化、管理看板']],7.6)
doc.add_heading('附录A 官方依据与边界',0);tab(doc,['文件','机关/文号','本手册用途','官方链接'],D['sources'],6.6)
doc.add_heading('附录B 核心面试问题',0);bullets(doc,['你如何保证月结及时？','业务部门不配合怎么办？','你如何复核下属工作？','发现长期挂账如何推进？','如何平衡财务控制和业务效率？','什么情况下必须升级CFO？','问题已经Resolved，为什么还不能Closed？','主管应该亲自做什么、授权什么？'])
main=O/'财务主管管理、内控与跨部门协同实操手册.docx';doc.save(main)

# Case & communication library
c=setup('财务跨部门沟通与问题处理案例库','50个跨部门案例 + 50个沟通模板')
c.add_heading('使用说明',0);p(c,'案例按照问题、风险、业务原因、处理、沟通、提醒、升级和关闭编写。沟通模板应替换为具体事项、金额、单据号、准确时间和责任人，不能机械复制。')
c.add_heading('目录',0);toc(c);page(c)
c.add_heading('第一篇 跨部门问题处理案例',0)
groups=[]
for x in D['cases']:
    if x['category'] not in groups:groups.append(x['category'])
for g in groups:
    c.add_heading(g,1)
    for x in [z for z in D['cases'] if z['category']==g]:
        c.add_heading(f"{x['id']} {x['scenario']}",2)
        tab(c,['环节','处理内容'],[['财务风险',x['risk']],['业务为什么这样做',x['business_reason']],['主管怎么处理',x['action']],['怎么说',x['words']],['提醒时间',x['timing']],['升级条件',x['escalation']],['关闭标准',x['closure']]],7.5)
    c.add_heading('【面试可能怎么问】',2);p(c,f'如果{g}持续不配合或反复发生同类问题，你会如何分配责任、升级并验证关闭？')
page(c)
c.add_heading('第二篇 50个沟通模板',0)
groups=[]
for x in D['templates']:
    if x['category'] not in groups:groups.append(x['category'])
for g in groups:
    c.add_heading(g,1)
    tab(c,['编号','事项','第一次提醒','临近截止','逾期/升级','留痕'],[[x['id'],x['topic'],x['first'],x['deadline'],x['overdue'],x['record']] for x in D['templates'] if x['category']==g],6.7)
c.add_heading('第三篇 沟通使用规则',0)
bullets(c,['写具体单据、项目和金额，不发泛泛的“请尽快”。','写明确日期和时间，并说明无法按时完成时需要反馈什么。','说明业务影响，不把所有要求包装成“财务规定”。','首次普通延迟可柔性沟通；影响月结、申报、资金、报表或涉嫌违规时必须书面留痕。','升级针对问题和风险，不针对个人情绪。','收到“已经处理”后仍应检查凭证、系统、资料和报表税务影响。'])
casefile=O/'财务跨部门沟通与问题处理案例库.docx';c.save(casefile)
print(main);print(casefile)
