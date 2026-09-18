from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE=Path(__file__).resolve().parent
OUT=BASE/'deliverables'
D=json.loads((BASE/'content.json').read_text(encoding='utf-8'))
NAV='17365D'; LIGHT='D9EAF7'; PALE='F4F7FA'; LINE='D9D9D9'; RED='C00000'; GREEN='548235'

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def borders(table):
    tblPr=table._tbl.tblPr; tb=tblPr.find(qn('w:tblBorders'))
    if tb is None: tb=OxmlElement('w:tblBorders'); tblPr.append(tb)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{edge}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),LINE); tb.append(el)

def set_repeat(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def table(doc, headers, rows, widths=None, font=8.2):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); shade(c,NAV); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(font)
    set_repeat(t.rows[0])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text='' if v is None else str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],PALE)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.05
                for r in p.runs: r.font.size=Pt(font); r.font.name='Microsoft YaHei'
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    borders(t); doc.add_paragraph().paragraph_format.space_after=Pt(2); return t

def heading(doc,text,level=1):
    p=doc.add_heading(text,level=level); p.paragraph_format.keep_with_next=True; return p

def para(doc,text='',bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.25
    if bold_lead:
        r=p.add_run(bold_lead); r.bold=True
    p.add_run(text)
    return p

def bullets(doc, items):
    for x in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(x)

def flow(doc, steps):
    rows=[]
    for i,s in enumerate(steps):
        rows.append([str(i+1),s,'↓' if i<len(steps)-1 else '完成'])
    table(doc,['序号','流程节点','流转'],rows,[1.2,12.5,2.0],8.5)

def add_toc(doc):
    p=doc.add_paragraph(); run=p.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u'); run._r.addnext(fld)

def setup():
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(2.1); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2.2); sec.right_margin=Cm(2.0)
    styles=doc.styles
    styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size=Pt(10)
    for n,size in [('Title',26),('Subtitle',12),('Heading 1',18),('Heading 2',14),('Heading 3',11.5)]:
        s=styles[n]; s.font.name='Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
        if n.startswith('Heading'): s.font.bold=True; s.paragraph_format.space_before=Pt(10); s.paragraph_format.space_after=Pt(5); s.paragraph_format.keep_with_next=True
    return doc

def project_detail(doc,p):
    heading(doc,f"{p['id']} {p['name']}",2)
    para(doc,f"该项目属于{p['process']}流程，目标是把人工整理、匹配或检查转化为可重复运行的规则，同时保留异常判断和审批责任。")
    table(doc,['项目要素','具体设计'],[
        ['业务背景',f"华辰智能当前依赖Excel人工完成{p['name']}，结果受个人经验、字段变化和截止时间影响。"],
        ['当前人工操作','导出数据→复制粘贴→清洗字段→逐项查找→标记异常→人工复核→保存底稿。'],
        ['每月工作量',f"模拟估计原型上线后每月可节约约{p['saved']}小时；正式立项前必须用真实笔数重新测量。"],
        ['输入数据',f"ERP或业务模块明细、主数据、期间和规则参数。核心规则：{p['rules']}。"],
        ['输出结果','处理结果、未匹配或异常清单、原因代码、人工处理状态、运行日志和复核记录。'],
        ['推荐工具',p['tool']],
        ['选择原因',f"按规则、数据量和维护能力选择Level {p['level']}方案；不为了展示AI增加不确定性。"],
        ['异常情况','空值、重复键、日期异常、金额尾差、字段改名、主数据缺失、期间跨越和接口中断。'],
        ['人工复核点','高金额、低置信度、主体不一致、业务实质不明、影响记账或报表的项目。'],
        ['数据安全','只读取完成任务所需字段；原始数据只读；敏感字段脱敏；运行和复核分权。'],
        ['预计开发难度',f"Level {p['level']}，财务人员可先做最小可用原型。"],
        ['预计开发时间',f"原型、测试和修改合计约{p['hours']}小时。"],
        ['ROI',f"月节约{p['saved']}小时为案例假设。以实际人工成本、软件费、维护费和错误减少收益计算回收期。"],
        ['上线前测试','正常、重复、空值、极端金额、日期边界、错误格式、尾差和系统字段变化。'],
        ['上线后维护','每月抽样复核；字段或规则变更必须登记版本；至少保留一个可执行的手工Fallback。'],
    ],[3.4,12.3],8.6)
    flow(doc,['保留原始输入','检查字段和数据质量','执行确定性规则','输出通过项和异常项','人工复核异常','授权人员批准后执行','归档输入、规则、输出和证据'])

def build_doc():
    doc=setup()
    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('AI 财务自动化实战手册')
    p=doc.add_paragraph(style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('从 Excel 到 Power Query Python RPA 与 AI Agent')
    doc.add_paragraph(); p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('适用岗位：总账会计 财务主管 财务经理').bold=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('版本日期：2026年9月11日')
    doc.add_paragraph(); para(doc,'本手册面向有审计基础、企业财务实操和自动化开发经验相对不足的学习者。它从业务、数据、规则、控制和投入产出出发，帮助读者判断一项工作是否值得自动化，以及应该选择哪一种工具。')
    para(doc,'核心原则：先稳定流程和数据，再自动化；能用确定性规则解决时优先确定性规则；自动化负责准备、匹配、计算和提示，财务人员负责判断、复核、批准和承担责任。',bold_lead='核心结论：')
    doc.add_page_break(); heading(doc,'目录',1); add_toc(doc); doc.add_page_break()

    heading(doc,'第一章 企业财务自动化全景',1)
    heading(doc,'1.1 贯穿案例公司',2)
    table(doc,['项目','设定'],[[k,D['company'][k]] for k in ['name','profile','systems','volumes']],[3,12.7],9)
    heading(doc,'1.2 从业务到自动化结果',2)
    flow(doc,['业务系统及外部数据','数据获取和格式标准化','确定性规则处理','必要时AI辅助理解','形成对账表和异常清单','财务人工复核','授权人员批准','ERP处理和资料归档'])
    heading(doc,'1.3 自动化机会全景',2)
    table(doc,['业务领域','典型任务','重复','规则','判断','价值','工具','人工保留'],D['panorama'],[2.0,3.2,1.0,1.0,1.0,1.0,2.6,4.0],7.5)
    para(doc,'系统计算结果为零差异，只能说明预设字段和规则没有发现差异，不能证明业务真实性、会计判断和截止性一定正确。',bold_lead='实务提醒：')

    heading(doc,'第二章 自动化可行性和ROI',1)
    heading(doc,'2.1 立项前四道门槛',2)
    table(doc,['门槛','必须回答的问题','不满足时怎么做'],[
        ['流程','现状流程、Owner和截止时间是否明确','先梳理流程和RACI'],['数据','数据是否合法、稳定、可追溯','先治理字段、主数据和归档'],['权限','是否突破付款、申报或ERP审批','重新划定自动化边界'],['责任','谁制表、谁复核、谁批准','明确岗位和异常升级路径']], [2.1,7.2,6.4],8.8)
    heading(doc,'2.2 100分评分模型',2)
    table(doc,['维度','指标','权重','评分方法'],[
        ['收益价值','月度人工耗时',15,'月度工时越高越优先'],['收益价值','发生频率',8,'日频和周频高于月频'],['收益价值','错误风险降低',12,'能减少漏项、重复和金额错误'],['收益价值','复用范围',8,'可覆盖多个主体和期间'],['收益价值','月结关键程度',7,'关键路径任务优先'],['实施可行性','规则清晰度',12,'能够写成明确条件和阈值'],['实施可行性','数据结构化',10,'表格和接口数据优于图片自由文本'],['实施可行性','来源稳定性',8,'字段和导出格式稳定'],['实施可行性','异常比例',5,'异常越少越适合自动处理'],['实施可行性','工具可用性',5,'可以合法导出或连接'],['可维护性','维护难度',5,'财务人员能够修改'],['可维护性','用户接受',3,'输入人员愿意按标准维护'],['可维护性','回退能力',2,'失败时可以手工完成']], [2.6,5.2,1.5,6.4],8.5)
    para(doc,'项目优先得分 = 收益价值 + 实施可行性 + 可维护性 - 风险扣分。70分以上优先，55至69分条件成熟后试点，40至54分先优化数据或流程，40分以下暂不自动化。')
    heading(doc,'2.3 ROI和总体拥有成本',2)
    bullets(doc,['月度人工成本 = 月发生次数 × 单次耗时 ÷ 60 × 综合小时人工成本。','首年TCO = 开发工时成本 + 软件及API + 测试 + 培训 + 首年维护。','首年ROI =（首年收益 - 首年TCO）÷ 首年TCO。','回收期超过12个月时，通常先寻找Excel或Power Query等更轻的方案。'])
    table(doc,['案例','月度节约','开发投入','建议'],[['每月只省20分钟','0.33小时','三天Python开发','通常不立项'],['月结关键路径每月省8小时','8小时','20小时Power Query','适合试点'],['低频但可避免重大重复付款','时间收益较低','规则检查8小时','结合风险降低收益判断']],[4,3,4,4.7],9)

    heading(doc,'第三章 工具选择',1)
    table(doc,['工具','适用问题','典型场景','不适合'],D['tools'],[2.2,5,4.3,4.3],8.5)
    heading(doc,'3.1 决策顺序',2); flow(doc,['检查ERP原生功能','判断Excel是否足够','重复清洗使用Power Query','复杂或大数据使用Python','无API且界面固定时考虑RPA','非结构化理解时使用AI','流程权限成熟后再设计Agent'])
    heading(doc,'3.2 为什么不能所有工作都用AI',2)
    para(doc,'金额匹配、余额计算、字段转换和报表勾稽需要确定、稳定、可复算的规则。大语言模型的概率输出会增加不必要的不确定性。AI更适合合同、邮件、说明等非结构化内容的初步理解，并且必须保留来源和人工复核。')

    heading(doc,'第四章 五个优先练习项目',1)
    for p in D['projects'][:1]+[D['projects'][4],D['projects'][7],D['projects'][8],D['projects'][15]]: project_detail(doc,p)
    heading(doc,'4.6 银行对账匹配层级',2)
    table(doc,['层级','规则','是否自动通过'],[['一级','金额、日期、账户和唯一键完全一致','可建议自动通过，仍需抽样'],['二级','金额一致、日期相差1至3天、户名或摘要相似','人工确认'],['三级','一对多或多对一组合金额一致','人工确认'],['四级','手续费、利息、未知收款、尾差','人工判断及可能补账']],[2,8,5.7],9)
    heading(doc,'4.7 暂估匹配的会计边界',2)
    para(doc,'工具只能识别暂估、PO、入库和发票之间的关系。是否冲销、如何处理价格差异、是否属于跨期以及进项税何时确认，仍应由财务根据真实业务、发票和企业政策判断。')
    heading(doc,'4.8 异常余额不能机械重分类',2)
    para(doc,'例如应收账款贷方可能是预收、退款未付、核销错误、串户或重复收款。规则库只能提示“方向异常”。总账必须追到客户、合同、发货、开票和收款，再决定调账、重分类或保持原列报。')

    heading(doc,'第五章 其他十五个可落地项目',1)
    priority_ids={'AUT-01','AUT-05','AUT-08','AUT-09','AUT-16'}
    for p in D['projects']:
        if p['id'] not in priority_ids: project_detail(doc,p)

    heading(doc,'第六章 自动化项目生命周期与UAT',1)
    flow(doc,['发现痛点','记录现状与人工耗时','完成评分和ROI','定义输入输出和规则','准备并治理数据','开发最小可用原型','开发测试','业务UAT','人工并行验证','批准上线','持续监控','版本维护'])
    heading(doc,'6.1 测试设计',2)
    table(doc,['测试类型','案例','预期'],[['正常','唯一匹配、标准日期和正确金额','按规则通过'],['边界','截止日、阈值等于边界、零和负数','按书面规则稳定处理'],['异常','重复、主数据缺失、方向错误','进入异常清单'],['脏数据','空值、文本金额、错误日期、列名变化','停止或明确报错'],['极端','超大金额、大量数据、组合关系','性能可接受且不静默漏项'],['回归','规则修改后重跑历史测试集','旧功能没有被破坏']],[2.8,6,6.9],9)
    heading(doc,'6.2 结果可追溯',2)
    bullets(doc,['保存原始输入的文件名、期间和哈希或版本号。','保存规则编号、阈值、启用日期、修改人和审批人。','区分自动通过、自动建议、人工判断和最终批准。','任何结果能够追到输入记录及所用规则。','规则变更后重新执行回归测试并保留证据。'])
    heading(doc,'6.3 自动化失败的Fallback',2)
    flow(doc,['识别失败范围','停止使用不完整输出','恢复原始数据和手工模板','优先完成月结关键任务','记录影响和临时措施','修复并重新测试','人工并行核对','批准恢复'])

    heading(doc,'第七章 数据安全和权限',1)
    table(doc,['数据类型','等级','示例','AI原则','脱敏','权限','复核'],D['security'],[2.2,1,3.2,3.2,3.2,3.2,3.2],7.8)
    heading(doc,'7.1 禁止交给AI或普通自动化的内容',2)
    bullets(doc,['网银密码、U盾凭据和支付动态口令。','税务登录凭证、系统管理员密码和个人私钥。','未经授权的完整身份证、银行卡和工资信息。','未经企业批准上传的商业秘密、合同和未公开报表。'])
    heading(doc,'7.2 付款和申报红线',2)
    para(doc,'AI和RPA可以检查资料、准备建议、识别异常和生成待办，但不能默认自主批准或执行重大付款，也不能绕过纳税申报责任人的复核与提交。系统必须保留授权、审批、复核、职责分离和日志。')

    heading(doc,'第八章 Excel Power Query Power BI Python与RPA学习路径',1)
    table(doc,['阶段','学习内容','财务练习','完成标准'],[
        ['Excel','SUMIFS、XLOOKUP、COUNTIFS、动态数组、透视表、验证和条件格式','报表勾稽、预提、异常规则','能解释公式并处理空值重复'],['Power Query','文件夹合并、类型、清洗、合并、追加、分组和刷新','银行、暂估、往来、模块对账','换月文件后可稳定刷新'],['Power BI','模型、关系、维度、指标、DAX和刷新','应收、库存、月结和经营指标','指标能追到源数据'],['Python','pandas、Excel读写、匹配、汇总、异常检测和批处理','五个优先项目增强','能运行、读日志和修改参数'],['RPA','录制、元素、等待、异常和凭证管理','无API系统下载及归档','界面变化时安全停止'],['AI Agent','工具调用、状态、权限、日志和停止条件','受控Close Agent','只编排成熟工具并等待审批']],[2.2,5.0,5.2,4.4],8.5)
    heading(doc,'8.1 三个月路线',2)
    table(doc,['周','重点','成果'],[[1,'自动化测量、ROI和Excel规则','报表勾稽原型'],[2,'Power Query导入与清洗','统一字段模板'],[3,'银行自动对账','对账工具和未匹配清单'],[4,'长期挂账扫描','账龄和责任清单'],[5,'Python与pandas','能读写和合并Excel'],[6,'暂估自动匹配','到票和差异结果'],[7,'TB异常规则库','异常余额清单'],[8,'UAT和并行验证','测试证据与版本记录'],[9,'Power BI模型','月结异常Dashboard'],[10,'API和RPA边界','接口决策记录'],[11,'AI辅助','合同摘要和异常说明'],[12,'受控Close Agent概念验证','读取任务、运行检查、等待批准']],[1,6.8,8],8.7)

    heading(doc,'第九章 Finance Close Agent 放在最后',1)
    heading(doc,'9.1 Agent可以做什么',2)
    bullets(doc,['读取Close Tracker和任务依赖。','调用已经验证的对账、异常扫描和报表校验工具。','整理未完成事项和Issue。','生成提醒草稿和初步异常报告。','等待责任人反馈并更新状态。'])
    heading(doc,'9.2 Agent不能自行做什么',2)
    bullets(doc,['调整重大会计分录。','改变会计政策。','确认复杂收入和重大估计。','提交税务申报。','批准或执行重大付款。','关闭未经责任人验证的重大异常。'])
    flow(doc,['流程和主数据先稳定','确定性检查工具完成UAT','明确系统权限和日志','Agent只调度已批准工具','低风险事项生成建议','高风险事项停止并请求人工','授权人确认后继续'])

    heading(doc,'第十章 主管如何推进自动化',1)
    table(doc,['阶段','主管动作','交付'],[['发现','让执行人记录笔数、耗时、错误和截止影响','现状流程与基线'],['评估','组织评分、ROI和风险门槛','立项或不立项结论'],['设计','确定字段、规则、例外、Owner和审批','业务需求和控制设计'],['开发','限制范围，要求可读和可维护','最小可用原型'],['UAT','业务人员提供正常、边界和异常案例','测试证据'],['上线','人工并行、复核、授权和回退','上线批准'],['运行','监控异常率、失败率和实际节约时间','月度运行报告'],['维护','管理规则和字段版本','变更记录和回归测试']],[2.2,8.3,5.3],9)
    para(doc,'自动化项目的成功标准不是“代码能运行”，而是数据来源稳定、规则可解释、结果可追溯、异常有人处理、失败可以回退，并且实际节约时间足以覆盖持续维护成本。',bold_lead='财务主管检查点：')

    heading(doc,'附录A 20个项目机会清单',1)
    table(doc,['编号','项目','流程','工具','等级','投入小时','月省小时','核心规则'],[[p['id'],p['name'],p['process'],p['tool'],p['level'],p['hours'],p['saved'],p['rules']] for p in D['projects']],[1.3,2.8,2.2,3.2,1,1.4,1.4,4.2],7.3)
    heading(doc,'附录B 项目包使用顺序',1)
    bullets(doc,['先阅读每个项目README和字段说明。','打开练习工作簿，理解输入、规则、输出和UAT。','用Power Query方案完成可刷新版本。','运行Python增强脚本并与预期结果比较。','故意修改日期、金额、重复键和空值，观察异常。','把自己的测试结果记录进UAT工作簿。'])

    out=OUT/'AI＋财务自动化实战手册.docx'; OUT.mkdir(parents=True,exist_ok=True); doc.save(out); return out

def topic(title, children):
    return {'id':title,'title':title,'children':{'attached':[topic(x,[]) if isinstance(x,str) else topic(x[0],x[1]) for x in children]}} if children else {'id':title,'title':title}

def build_xmind():
    main=['自动化评估','Excel','Power Query','Power BI','Python','RPA','AI','Agent','银行','应收','应付','发票','税务','成本','存货','月结','报表','分析','资料管理','数据安全']
    sheets=[]
    root=topic('企业财务自动化全景',[(x,[]) for x in main]); sheets.append({'id':'s1','class':'sheet','title':'全景图','rootTopic':root})
    project_nodes=[]
    for level in range(1,6): project_nodes.append((f'Level {level}',[(p['name'],[]) for p in D['projects'] if p['level']==level]))
    sheets.append({'id':'s2','class':'sheet','title':'20个项目','rootTopic':topic('财务自动化项目地图',project_nodes)})
    route=[('第1个月',[('Excel规则',[]),('Power Query',[]),('银行对账',[]),('长期挂账',[])]),('第2个月',[('Python基础',[]),('暂估匹配',[]),('TB异常',[]),('UAT',[])]),('第3个月',[('Power BI',[]),('API与RPA',[]),('AI辅助',[]),('受控Agent',[])])]
    sheets.append({'id':'s3','class':'sheet','title':'学习路线','rootTopic':topic('3个月学习路线',route)})
    path=OUT/'企业财务自动化全景.xmind'
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        z.writestr('content.json',json.dumps(sheets,ensure_ascii=False))
        z.writestr('metadata.json',json.dumps({'creator':{'name':'Codex'},'activeSheetId':'s1'},ensure_ascii=False))
        z.writestr('manifest.json',json.dumps({'file-entries':{'content.json':{},'metadata.json':{}}}))
    return path

if __name__=='__main__':
    print(build_doc()); print(build_xmind())

