from pathlib import Path
from zipfile import ZipFile
import json, re
import pandas as pd
from openpyxl import load_workbook
from PIL import Image, ImageOps, ImageDraw

BASE=Path(__file__).resolve().parent; OUT=BASE/'deliverables'; QA=BASE/'qa'; CONTACT=QA/'contacts'; CONTACT.mkdir(parents=True,exist_ok=True)

def contact(files,out,cols=4,thumb=(340,440)):
    files=list(files); rows=(len(files)+cols-1)//cols
    canvas=Image.new('RGB',(cols*thumb[0],rows*thumb[1]),'white'); draw=ImageDraw.Draw(canvas)
    for i,p in enumerate(files):
        im=Image.open(p).convert('RGB'); im.thumbnail((thumb[0]-12,thumb[1]-28))
        x=(i%cols)*thumb[0]+6; y=(i//cols)*thumb[1]+20
        canvas.paste(im,(x+(thumb[0]-12-im.width)//2,y)); draw.text((x,3+(i//cols)*thumb[1]),p.stem,fill='black')
    canvas.save(out)

def validate_xlsx():
    reports=[]
    for p in OUT.rglob('*.xlsx'):
        wb=load_workbook(p,data_only=False,read_only=False)
        formulas=[]; errors=[]
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value,str) and c.value.startswith('='): formulas.append(f'{ws.title}!{c.coordinate}')
                    if isinstance(c.value,str) and re.search(r'#REF!|#DIV/0!|#VALUE!|#NAME\?|#N/A|#NUM!|#NULL!|#SPILL!|#CALC!',c.value): errors.append(f'{ws.title}!{c.coordinate}:{c.value}')
        reports.append({'file':str(p.relative_to(OUT)),'sheets':len(wb.sheetnames),'formulas':len(formulas),'errors':errors})
    return reports

def validate_scripts():
    pack=OUT/'财务自动化实战项目包'
    checks={}
    # Business result checks, intentionally do not require exact formatting.
    x=pd.read_excel(pack/'01-银行自动对账'/'运行结果.xlsx',sheet_name='匹配结果')
    checks['bank_rows']=len(x); checks['bank_exact']=int((x.match_type=='精确匹配').sum()); checks['bank_crossday']=int((x.match_type=='跨日候选').sum())
    x=pd.read_excel(pack/'02-暂估自动匹配'/'运行结果.xlsx'); checks['accrual_results']=x.result.value_counts().to_dict()
    x=pd.read_excel(pack/'03-长期挂账扫描'/'运行结果.xlsx'); checks['aging_high']=int((x.priority=='高').sum()); checks['aging_direction']=int(x.direction_flag.sum())
    x=pd.read_excel(pack/'04-全科目异常检查'/'运行结果.xlsx'); checks['tb_flags']=len(x); checks['tb_rules']=sorted(x.rule_id.unique().tolist())
    x=pd.read_excel(pack/'05-财务报表自动勾稽'/'运行结果.xlsx'); checks['fs_failed']=int((x.status=='未通过').sum()); checks['fs_passed']=int((x.status=='通过').sum())
    assert checks['bank_rows']==7 and checks['bank_exact']==2 and checks['bank_crossday']==1
    assert checks['accrual_results'].get('可冲销')==1 and checks['accrual_results'].get('差异待查')==1
    assert checks['aging_high']>=3 and checks['aging_direction']==1
    assert checks['tb_flags']==5
    assert checks['fs_failed']==3 and checks['fs_passed']==1
    return checks

def main():
    reports=validate_xlsx(); checks=validate_scripts()
    with ZipFile(OUT/'AI＋财务自动化实战手册.docx') as z: doc_bad=z.testzip(); doc_entries=len(z.namelist())
    with ZipFile(OUT/'企业财务自动化全景.xmind') as z: xm_bad=z.testzip(); xm_sheets=len(json.loads(z.read('content.json')))
    pages=sorted((QA/'docx_render').glob('page-*.png'),key=lambda p:int(re.search(r'(\d+)',p.stem).group()))
    for i in range(0,len(pages),12): contact(pages[i:i+12],CONTACT/f'doc_pages_{i+1:02d}_{min(i+12,len(pages)):02d}.png',cols=4)
    imgs=sorted([p for p in QA.glob('*.png') if p.is_file()])
    for i in range(0,len(imgs),12): contact(imgs[i:i+12],CONTACT/f'xlsx_sheets_{i+1:02d}_{min(i+12,len(imgs)):02d}.png',cols=3,thumb=(430,360))
    report={'xlsx':reports,'script_checks':checks,'docx':{'zip_bad':doc_bad,'entries':doc_entries,'pages':len(pages)},'xmind':{'zip_bad':xm_bad,'sheets':xm_sheets}}
    (QA/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2,default=str))

if __name__=='__main__': main()
