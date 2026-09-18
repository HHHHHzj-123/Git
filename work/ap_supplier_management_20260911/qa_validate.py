import json,zipfile
from pathlib import Path
from openpyxl import load_workbook
root=Path(r'C:\Users\HZJ\Desktop\Git\work\ap_supplier_management_20260911')
d=json.loads((root/'content.json').read_text(encoding='utf-8'))
p=root/'deliverables'/'应付账款及供应商管理工具包.xlsx'
# independent expected values from source transaction data
bal=sum(x[10] for x in d['ap'])
over=sum(x[10] for x in d['ap'] if x[10]>0 and x[4]<'2026-08-31')
future=sum(x[10] for x in d['ap'] if x[10]>0 and '2026-09-01'<=x[4]<='2026-09-30')
grni=sum(x[10] for x in d['ap'] if '暂估' in x[11])
dispute=sum(x[10] for x in d['ap'] if x[12]=='是')
negative=-sum(x[10] for x in d['ap'] if x[10]<0)
expected=[bal,over,over/bal,future,grni,dispute,negative,1,1,300000,2]
wb=load_workbook(p,data_only=True,read_only=True)
ws=wb['00-管理看板'];actual=[ws[f'B{i}'].value for i in range(6,17)]
for i,(a,e) in enumerate(zip(actual,expected),6):
    assert a is not None,(i,a,e)
    assert abs(float(a)-float(e))<0.01,(i,a,e)
wbf=load_workbook(p,data_only=False,read_only=True)
formula_count=0;errs=[]
for ws2 in wbf.worksheets:
    for row in ws2.iter_rows():
        for c in row:
            v=c.value
            if isinstance(v,str) and v.startswith('='): formula_count+=1
            if isinstance(v,str) and v in {'#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A'}:errs.append((ws2.title,c.coordinate,v))
assert not errs,errs
for q in [p,root/'deliverables'/'应付账款及供应商管理实战手册.docx',root/'deliverables'/'应付账款及供应商管理全景.xmind']:
    assert zipfile.is_zipfile(q),q
    with zipfile.ZipFile(q) as z: assert not z.testzip(),q
print({'dashboard':actual,'expected':expected,'sheets':len(wb.sheetnames),'formulas':formula_count,'formula_errors':errs})
