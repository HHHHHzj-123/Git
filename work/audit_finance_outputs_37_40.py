import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\HZJ\Desktop\Git\work")
EXPECTED = {
    37: ("72小时启动审计的绝地反击", "7498393544000867594", [17, 86]),
    38: ("服装制造财务接管的正确顺序", "7530205076594461986", [13, 67]),
    39: ("从3000万到2亿元（上）", "7538263260135427382", [14, 53]),
    40: ("从3000万到2亿元（下）", "7538368096617237817", [17, 81]),
}


def walk_topics(node):
    yield node
    children = node.get("children", {}).get("attached", [])
    for child in children:
        yield from walk_topics(child)


for episode, (title, video_id, expected_counts) in EXPECTED.items():
    docx_path = ROOT / f"齐昊老师_学员实战案例_第{episode}集_文字思维导图.docx"
    with zipfile.ZipFile(docx_path) as archive:
        assert archive.testzip() is None
        assert "word/document.xml" in archive.namelist()
        xml = archive.read("word/document.xml").decode("utf-8")
        assert "TODO" not in xml and "PLACEHOLDER" not in xml

    doc = Document(docx_path)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert title in all_text and video_id in all_text
    if episode == 39:
        assert "第40集" in all_text and "上篇" in all_text
    if episode == 40:
        assert "承接上集" in all_text and "两集共同的方法论" in all_text
    headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    lists = [
        p
        for p in doc.paragraphs
        if p._p.pPr is not None and p._p.pPr.numPr is not None
    ]
    assert headings and lists
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:tblW")).get(qn("w:w")) == "9360"
        assert tbl_pr.find(qn("w:tblInd")).get(qn("w:w")) == "120"
        grid = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
        assert sum(grid) == 9360
        for row in table.rows:
            assert [int(cell._tc.tcPr.tcW.get(qn("w:w"))) for cell in row.cells] == grid

    xmind_path = ROOT / f"齐昊老师_学员实战案例_第{episode}集_思维导图.xmind"
    with zipfile.ZipFile(xmind_path) as archive:
        assert archive.testzip() is None
        assert {"content.json", "manifest.json"}.issubset(archive.namelist())
        content = json.loads(archive.read("content.json"))
        assert len(content) == 2
        assert [sheet["title"] for sheet in content] == ["01 总览图", "02 详细知识树"]
        topic_sets = [list(walk_topics(sheet["rootTopic"])) for sheet in content]
        counts = [len(items) for items in topic_sets]
        assert counts == expected_counts
        ids = [topic["id"] for items in topic_sets for topic in items]
        assert len(ids) == len(set(ids))
        assert all(title in sheet["rootTopic"]["title"] for sheet in content)

    print(
        f"EP{episode}\theadings={len(headings)}\tlists={len(lists)}\t"
        f"tables={len(doc.tables)}\ttopics={counts}\tOK"
    )
