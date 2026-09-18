from pathlib import Path
import json, importlib.util
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent; OUT=ROOT/'deliverables'; OUT.mkdir(exist_ok=True)
DATA=json.loads((ROOT/'cases.json').read_text(encoding='utf-8'))
DOCX=OUT/'企业发票、报销与异常支出实战手册.docx'
base=ROOT.parent/'invoice_tax_manual_20260909'/'build_tax_manual.py'
spec=importlib.util.spec_from_file_location('base',base);m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
heading=m.heading;para=m.para;table=m.table;marker=m.marker;add_toc=m.add_toc;apply_styles=m.apply_styles;set_run=m.set_run
BLUE=m.BLUE;RED=m.RED;GREEN=m.GREEN;AMBER=m.AMBER

def link(p,text,url):
    rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid);r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    c=OxmlElement('w:color');c.set(qn('w:val'),BLUE);u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.extend([c,u]);r.append(rp)
    t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def laws(codes):
    a=[]
    for code in codes.replace('；','、').split('、'):
        s=DATA['sources'].get(code.strip())
        if s: a.append(f"{code} {s[0]}｜{s[1]}｜{s[2]}")
    return '\n'.join(a)

def footer(sec):
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;set_run(p.add_run('企业发票、报销与异常支出实战手册  '),8,False,'777777')
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)

def add_case(d,c):
    p=heading(d,f"{c['id']}　{c['title']}",2);p.paragraph_format.page_break_before=True
    color=RED if c['risk'].startswith('红') else AMBER if c['risk'].startswith('橙') else GREEN if c['risk'].startswith('蓝') else 'B8860B'
    marker(d,'风险等级',c['risk'],color)
    rows=[('1. 业务场景',c['scene']),('2. 核心问题',c['core']),('3. 是否违规与结论',c['status']+'。'+c['conclusion']),
      ('4. 正确操作方式',c['operation']),('5. 会计处理与分录',c['entries']),('6. 增值税处理',c['vat']),('7. 企业所得税处理',c['cit']),
      ('8. 个人所得税处理',c['iit']),('9. 发票要求',c['invoice']),('10. 支持性资料',c['docs']),('11. 法律法规依据',laws(c['laws'])),
      ('12. 怎么沟通',c['comm']),('13. 拒不整改/继续执行的后果',c['consequences']),('14. 可能责任主体与财务留痕',c['responsible'])]
    table(d,['判断项目','实务处理'],rows,[4.1,12.9],7.6)

def build():
    d=Document();apply_styles(d);sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.55);sec.bottom_margin=Cm(1.45);sec.left_margin=Cm(1.6);sec.right_margin=Cm(1.6);footer(sec)
    p=d.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(118);p.add_run('企业发票、报销与异常支出\n实战手册')
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;set_run(p.add_run('141个高频案例｜总账会计与财务主管工作版'),13,False,'444444')
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(128);set_run(p.add_run('V1.0｜政策核验日：2026年9月10日'),9.5,False,'666666')
    d.add_page_break();heading(d,'使用说明',1)
    para(d,'本手册用于中国大陆企业财务实务学习与内部处理参考。结论按会计、增值税、企业所得税、个人所得税、发票和公司治理分别判断；地区执行、事实证据和后续政策变化可能改变个案结论。涉及主动更正、重大税款、人员舞弊或刑事风险时，应由企业税务负责人、法务及外部专业人员结合完整事实处理。')
    marker(d,'核心结论','公司支付不等于公司费用；发票真实不等于交易真实；会计可入账、增值税可抵扣、企业所得税可扣除是三个不同问题。',BLUE)
    heading(d,'目录',1);add_toc(d)
    heading(d,'第一篇　异常业务判断总则',1)
    heading(d,'1.1 十一步判断链',2)
    table(d,['步骤','财务要问','形成的工作底稿'],[
      ['1 还原业务','发生了什么，为什么发生','业务事实说明'],['2 查受益人','最终谁使用、谁获益','受益人清单'],['3 穿透主体','合同、履约、发票、付款、收款是否一致','主体关系图'],['4 会计确认','资产、费用、薪酬、分配还是往来','会计判断与分录'],['5 增值税','应税、扣税凭证、用途限制、视同交易','进销项处理'],['6 企业所得税','真实、相关、合理、凭证、限额和跨期','扣除/调增台账'],['7 个人所得税','利益是否归属个人及所得性质','扣缴判断'],['8 发票','票面、开票方、内容、金额、唯一性','接受/退票/红冲'],['9 证据','资料能否形成闭环','证据包'],['10 权限','谁申请、审批、制单、复核','审批链'],['11 后果','若拒不整改，风险升级到哪一级','整改或升级记录']], [2.2,7,7.8],8)
    heading(d,'1.2 七档合规结论与四级风险',2)
    table(d,['类型','含义','典型动作'],[['①正常','真实经营且资料完整','按制度入账、申报和归档'],['②有条件可做','需合同、审批、清单、分拆或专项判断','补齐资料后处理'],['③会计可反映、税务不可扣','真实支出但不满足抵扣/扣除规则','如实入账并税务调整'],['④需纳税调整','限额、跨期或税会口径不同','建立税会差异台账'],['⑤高风险','主体、货款、时间或业务证据异常','暂停付款/抵扣，升级复核'],['⑥明显不应处理','私人消费或公司不应承担','拒绝费用化、追回或挂个人往来'],['⑦涉嫌违法','虚构、买票、虚抵、毁损资料等','停止执行、保全记录、法务税务介入']], [3,6,8],8)
    table(d,['颜色','含义','动作'],[['蓝色','正常','常规审批'],['黄色','注意规范','补资料、分拆、留痕'],['橙色','高风险','暂停关键动作、主管复核'],['红色','不应执行/疑似违法','拒绝虚假处理并按权限升级']], [3,6,8],8)
    heading(d,'1.3 拒不整改的后果阶梯',2)
    table(d,['层级','可能后果','常见责任主体'],[
      ['流程与劳动管理','退回、拒付、追回、取消授权；员工故意骗报可依合法公示制度处分，严重违纪或造成重大损害时依法处理劳动关系','申请人、审批人'],
      ['会计纠正','冲销、重分类、更正报表；参与伪造者可能被罚款并在法定期间不得从事会计工作','单位负责人、会计经办/复核'],
      ['税款后果','进项扣减、纳税调增、更正申报、补税；逾期税款按日万分之五加收滞纳金','纳税人、扣缴义务人'],
      ['税务行政处罚','逃税可处少缴税款50%至5倍罚款；发票违法按情节没收违法所得并处罚款','单位、直接负责主管人员、直接责任人员'],
      ['公司治理与民事责任','返还占用资金、赔偿公司损失、关联交易责任','股东、实际控制人、董监高'],
      ['刑事责任','只有达到犯罪构成、数额/数量标准并结合主观故意和危害结果时，才可能判处罚金、拘役或有期徒刑；问题发票不自动等于坐牢','单位及直接负责主管人员、其他直接责任人员']], [3.1,8.6,5.3],8)
    marker(d,'判断提醒','罚款、补税、职业禁业和刑事责任不能互相替代。主动纠正、补税和如实配合可能影响处理结果，但并非所有案件都当然免责。',AMBER)
    heading(d,'1.4 责任主体怎么区分',2)
    table(d,['角色','应承担的控制责任','风险动作'],[['老板/单位负责人','保证会计资料真实完整，不得授意造假','强令虚列、批准私人消费、毁损资料'],['员工/申请人','真实说明业务并提交本人取得的原始资料','替票、重复报销、伪造行程'],['业务审批人','验证必要性、预算、受益人和履约','明知虚假仍批准'],['财务经办','审核凭证、按实入账、提出异常','明知无业务仍制单抵扣'],['财务主管/复核','建立制度、复核重大异常、升级红线','压下异常、指使修改事实'],['开票/收款方','如实开票并按真实交易收款','无业务开票、资金回流']], [3.6,6.7,6.7],8)
    heading(d,'1.5 面对不合理指令的SOP',2)
    para(d,'明确真实业务 → 判断受益人和主体 → 分别形成会计与税务结论 → 引用有效法规 → 提供合法替代方案 → 在OA/邮件记录风险和选择 → 按授权矩阵升级 → 明显违法时拒绝执行 → 保留原始记录，不制作或补造假证据。')
    heading(d,'1.6 20条可直接使用的沟通表达',2)
    scripts=[
      '这笔款公司可以先支付，但真实受益人是个人，账务上需要挂个人往来并约定归还，不能直接计入公司费用。','发票验真只能证明票据存在，尚不能证明公司实际取得了货物或服务；请补合同、验收和付款资料。','会计上应如实反映这笔支出，企业所得税是否扣除需要另做判断，我会同步登记纳税调整。','这项用途依法不得抵扣进项税，发票可以入账，但税额需要计入成本费用。','当前合同、履约、发票和付款主体不一致，请先说明法律关系并补代付/委托资料。','这张票已在系统中出现过，需先确认是否重复报销；核实前暂不付款。','周末或异地发票不是自动违规，但需要行程、对象和业务目的形成闭环。','私人部分请从本次申请中剔除，公司只承担能证明与公务相关的部分。','如果公司最终承担这项个人利益，需要按工资薪金或股息红利评估个税，不能只改费用科目。','供应商不能及时开票不等于业务不能入账；我们先按真实义务确认，同时跟踪税前扣除凭证。','这项支出达到资产确认条件，应登记资产并按受益期折旧/摊销，不能为了当期利润任意费用化。','请让开票方按真实品名和金额红冲重开，财务不能在发票上手工修改。','这笔招待费可以入账，但所得税扣除受双限额约束，需有对象、目的和审批资料。','员工福利需要名单、制度和签收，同时判断福利费限额、进项抵扣限制与个税。','母公司可以代付款，但费用、进项和资产应归实际购买并使用的主体，双方同步挂往来。','目前资料不足以支持抵扣和付款，我先把事项登记为异常，不会把事实不清的业务硬做进账。','如果继续按虚假业务处理，风险会从内部退单升级为补税、滞纳金、罚款，达到法定条件还可能涉及刑事责任。','我不能制作虚假合同、验收或会议记录；可以协助按真实业务寻找合法的账务和税务处理。','请把指令和业务事实写入OA，由有权负责人复核；财务会同时保留专业判断和替代方案。','该事项已经触及红线，我将停止制单、付款或抵扣，并按公司授权矩阵提交财务负责人和法务。']
    for i,s in enumerate(scripts,1): para(d,f'{i}. {s}')
    heading(d,'第二篇至第十一篇　141个高频案例',1)
    grouped={c['category']:[] for c in DATA['cases']}
    for c in DATA['cases']: grouped.setdefault(c['category'],[]).append(c)
    for idx,(prefix,cat) in enumerate(DATA['categories'].items(),2):
        name,coverage,_count=cat; heading(d,f'第{idx}篇　{name}',1)
        para(d,coverage)
        for c in grouped.get(name,[]): add_case(d,c)
    heading(d,'第十二篇　法规与官方资料索引',1)
    for code,s in DATA['sources'].items():
        p=d.add_paragraph();set_run(p.add_run(f"{code}　{s[0]}｜{s[1]}｜{s[2]}　"),8.4);link(p,'官方链接',s[3])
    heading(d,'附录　使用与维护',1)
    para(d,'建议在实际工作中使用Excel按关键词、风险等级、税种和案例编号检索，再回到本手册阅读完整判断。每季度复核官方法规状态；发生政策变化、地方执行差异或典型案例时，在版本记录中追加，不直接覆盖历史判断。')
    d.save(DOCX);print(DOCX)

if __name__=='__main__': build()
