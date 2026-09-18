from pathlib import Path
from docx import Document
from zipfile import ZipFile
import json,re
from openpyxl import load_workbook
root=Path(__file__).parent; out=root/'deliverables'
doc=Document(out/'企业发票、报销与异常支出实战手册.docx')
heads=[p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
caseheads=[x for x in heads if re.match(r'^(BOSS|ODD|EXP|WEL|MKT|NOP|INV|PAY|ASSET|RED)-\d{3}',x)]
text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
assert len(caseheads)==141,(len(caseheads),caseheads[-3:])
for term in ['拒不整改/继续执行的后果','可能责任主体与财务留痕','按日万分之五','50%至5倍','并非出现一张问题发票就自动入罪','20条可直接使用的沟通表达']: assert term in text,term
with ZipFile(out/'企业异常费用与发票问题.xmind') as z:
    content=json.loads(z.read('content.json')); assert len(content)==3
    raw=json.dumps(content,ensure_ascii=False); assert raw.count('风险：')==141; assert '刑事责任' in raw
wb=load_workbook(out/'异常发票与费用处理速查表.xlsx',read_only=False,data_only=False)
assert len(wb.sheetnames)==13
ws=wb['01-141案例速查']; assert ws.max_row==146 and ws.max_column==17,(ws.max_row,ws.max_column)
headers=[ws.cell(5,c).value for c in range(1,18)]
for h in ['拒不整改后果','可能责任主体','财务升级/留痕']: assert h in headers
allvals='\n'.join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None)
assert allvals.count('RED-0')>=10 and '可能被追究刑事责任' in allvals
print(f'docx_cases={len(caseheads)} tables={len(doc.tables)} headings={len(heads)}')
print(f'xmind_sheets={len(content)} xmind_cases={raw.count("风险：")}')
print(f'xlsx_sheets={len(wb.sheetnames)} main_rows={ws.max_row-5} main_cols={ws.max_column}')
