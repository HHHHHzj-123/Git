from pathlib import Path
import json, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/"deliverables"
IMG=ROOT/"images"
IMG.mkdir(exist_ok=True)
data=json.loads((ROOT/"content.json").read_text(encoding="utf-8"))

NAVY="#17365D"; BLUE="#2F75B5"; LIGHT="#D9EAF7"; PALE="#F3F6F9"; LINE="#D9E1E8"

def font(size,bold=False):
    paths=[Path(r"C:\Windows\Fonts\msyh.ttc"),Path(r"C:\Windows\Fonts\simhei.ttf")]
    for p in paths:
        if p.exists(): return ImageFont.truetype(str(p),size,index=0)
    return ImageFont.load_default()

def draw_flow(filename,title,rows):
    w=1500; margin=70; boxh=100; gap=42
    h=150+len(rows)*(boxh+gap)+40
    im=Image.new("RGB",(w,h),"white"); d=ImageDraw.Draw(im)
    d.text((margin,35),title,font=font(38),fill=(0,0,0))
    cols=max(len(r) for r in rows)
    usable=w-2*margin; bw=(usable-(cols-1)*35)/cols
    y=115
    for ri,row in enumerate(rows):
        offset=(cols-len(row))*(bw+35)/2
        for ci,label in enumerate(row):
            x=margin+offset+ci*(bw+35)
            fill=(230,241,250) if ri%2==0 else (243,246,249)
            d.rounded_rectangle((x,y,x+bw,y+boxh),radius=14,fill=fill,outline=(47,117,181),width=3)
            lines=[]; cur=""
            for ch in label:
                if d.textlength(cur+ch,font=font(23))>bw-24:
                    lines.append(cur); cur=ch
                else: cur+=ch
            if cur: lines.append(cur)
            total=len(lines)*30
            for li,t in enumerate(lines[:3]):
                tw=d.textlength(t,font=font(23)); d.text((x+(bw-tw)/2,y+(boxh-total)/2+li*30),t,font=font(23),fill=(0,0,0))
            if ci<len(row)-1:
                ax=x+bw+5; ay=y+boxh/2
                d.line((ax,ay,ax+22,ay),fill=(90,100,110),width=4); d.polygon([(ax+22,ay),(ax+12,ay-8),(ax+12,ay+8)],fill=(90,100,110))
        if ri<len(rows)-1:
            cx=w/2; sy=y+boxh+4
            d.line((cx,sy,cx,sy+25),fill=(90,100,110),width=4); d.polygon([(cx,sy+25),(cx-8,sy+15),(cx+8,sy+15)],fill=(90,100,110))
        y+=boxh+gap
    im.save(IMG/filename,dpi=(180,180))

draw_flow("01_landscape.png","企业业务系统与财务系统全景图",[
 ["销售 CRM","采购 SRM","仓储 WMS","生产 MES","费用 OA","人力 HR","资金 银企","发票税务"],
 ["销售","采购","库存","生产","成本","固定资产"],
 ["AR 应收","AP 应付","存货核算","产品成本","FA 资产","薪酬","资金","税务数据"],
 ["GL 总账"],["模块对账与月结"],["财务报表","税务申报","经营分析"]])
draw_flow("02_posting.png","业务单据进入总账的底层逻辑",[["业务单据"],["业务事件"],["审核与系统过账"],["会计凭证"],["财务子模块"],["GL总账"],["报表与分析"]])
draw_flow("03_p2p.png","采购到付款单据链",[["采购申请","采购订单"],["收货质检","采购入库"],["暂估 GR IR"],["供应商发票","三单匹配"],["应付未清项"],["付款申请","银企支付"],["银行回单","应付核销"]])
draw_flow("04_o2c.png","销售到收款单据链",[["客户准入","销售订单"],["交货拣货","销售出库"],["签收验收","收入判断"],["开票","应收未清项"],["银行收款","收款认领"],["应收核销"]])
draw_flow("05_close.png","ERP月结依赖关系",[["业务截止","接口完整性"],["AP与AR"],["库存与生产"],["存货核算与产品成本"],["固定资产","薪酬","资金","税务"],["子模块与GL对账"],["GL调整与损益结转"],["报表校验与期间关闭"]])

doc=Document()
sec=doc.sections[0]
sec.top_margin=Cm(2.3); sec.bottom_margin=Cm(2.1); sec.left_margin=Cm(2.3); sec.right_margin=Cm(2.1)

styles=doc.styles
styles["Normal"].font.name="Microsoft YaHei"; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑"); styles["Normal"].font.size=Pt(10.5)
styles["Normal"].paragraph_format.line_spacing=1.35; styles["Normal"].paragraph_format.space_after=Pt(5)
for name,size in [("Title",26),("Heading 1",18),("Heading 2",14),("Heading 3",11.5)]:
    s=styles[name]; s.font.name="Microsoft YaHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑"); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0); s.font.bold=True
styles["Heading 1"].paragraph_format.space_before=Pt(18); styles["Heading 1"].paragraph_format.space_after=Pt(10)
styles["Heading 2"].paragraph_format.space_before=Pt(13); styles["Heading 2"].paragraph_format.space_after=Pt(7)
styles["Heading 3"].paragraph_format.space_before=Pt(8); styles["Heading 3"].paragraph_format.space_after=Pt(4)

def set_cell_shading(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"),color.replace("#",""))
def set_cell_margins(cell,top=90,start=90,bottom=90,end=90):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=tcMar.find(qn(f"w:{m}"))
        if node is None: node=OxmlElement(f"w:{m}"); tcMar.append(node)
        node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa")
def add_table(headers,rows,widths=None,font_size=8.2):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style="Table Grid"
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); set_cell_shading(c,NAVY); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.font.bold=True; r.font.size=Pt(font_size); r.font.name="Microsoft YaHei"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑")
        c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; set_cell_margins(c)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text="" if v is None else str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cells[i])
            if ri%2: set_cell_shading(cells[i],PALE)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.1
                for r in p.runs: r.font.size=Pt(font_size); r.font.name="Microsoft YaHei"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑")
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    # repeat header
    trPr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement("w:tblHeader"); rep.set(qn("w:val"),"true"); trPr.append(rep)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def add_para(text,boldlead=None):
    p=doc.add_paragraph()
    if boldlead:
        r=p.add_run(boldlead); r.bold=True
    p.add_run(text)
    return p
def bullet(text):
    p=doc.add_paragraph(style="List Bullet"); p.add_run(text); return p
def numbered(text):
    p=doc.add_paragraph(style="List Number"); p.add_run(text); return p
def add_pic(name,caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(IMG/name),width=Cm(16.2))
    c=doc.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in c.runs: r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(89,89,89)
def pagebreak(): doc.add_page_break()
def add_toc():
    p=doc.add_paragraph(); run=p.add_run(); fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),'TOC \\o "1-3" \\h \\z \\u'); run._r.addnext(fld)
    p2=doc.add_paragraph("目录页码将在Word或WPS中打开并更新域后显示。标题已经使用标准标题样式，左侧导航栏可直接跳转。")
    for r in p2.runs: r.font.size=Pt(9); r.font.color.rgb=RGBColor(89,89,89)

# Header/footer
for section in doc.sections:
    hp=section.header.paragraphs[0]; hp.text="企业财务软件与ERP实操手册"; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs: r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(100,100,100)
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); fp._p.append(fld)

# Cover
doc.add_paragraph().paragraph_format.space_after=Pt(80)
p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("企业财务软件与ERP实操手册")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("从业务单据到总账报表"); r.font.size=Pt(18); r.bold=True
doc.add_paragraph().paragraph_format.space_after=Pt(85)
for s in ["面向岗位：总账会计与财务主管","案例企业：华创智联设备有限公司","版本日期：2026年9月10日","适用范围：中国大陆制造业及硬件科技企业"]:
    p=doc.add_paragraph(s); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph().paragraph_format.space_after=Pt(55)
p=doc.add_paragraph("本手册围绕业务、单据、系统、凭证、子模块、总账、月结和报表建立完整工作链。软件名称用于说明常见产品逻辑，具体菜单、字段和配置应以企业实际版本为准。")
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
pagebreak()

doc.add_heading("使用说明",0)
add_para("这套资料的目标，是让有审计基础但缺少企业内部系统操作经验的学习者，能够进入一套陌生ERP后迅速识别业务入口、单据链、凭证生成时点、模块对账和月结依赖。学习时应先理解业务事件，再看软件界面；先找来源单据，再考虑总账调整。")
add_para("手册中的会计分录是贯穿案例的典型处理。企业可能因会计政策、系统配置、成本方法和组织模式不同而采用其他合规路径。系统自动生成凭证并不减轻财务判断责任。")
add_para("建议使用顺序：先阅读第1至4章建立全景；再按第5至12章理解八大循环；随后重点学习第13至16章总账、月结和对账；最后使用训练清单和异常排查表反复练习。")
doc.add_heading("阅读标识",1)
add_table(["标识","含义"],[["【核心结论】","必须掌握的系统或会计逻辑"],["【实务操作】","企业内具体执行步骤"],["【常见错误】","高频系统、业务或核算错误"],["【总账检查点】","月结和对账时应验证的事项"],["【财务主管复核点】","主管需要关注的控制和重大判断"]],[4,12],9)
doc.add_heading("目录",0); add_toc(); pagebreak()

doc.add_heading("第一篇 企业系统和ERP底层逻辑",0)
doc.add_heading("第1章 企业信息系统全景",1)
add_para("一家制造企业通常同时运行CRM、采购协同、WMS、MES、HR、OA、资金、发票和ERP。每套系统解决不同业务问题。ERP承担组织主数据、供应链、生产和财务集成，但不必承担所有前端操作。总账人员要掌握数据从哪里产生、经何种接口进入ERP、何时形成财务记录以及失败后由谁处理。")
add_pic("01_landscape.png","图1 企业业务系统与财务系统全景图。系统操作示意图，实际范围因企业架构而异。")
add_table(["系统","主要使用者","主要数据","进入财务的方式","财务关注点"],[
 ["CRM/销售","销售、客服","客户、订单、价格、交付","接口至ERP销售或AR","客户主体、合同、交付和收入事件"],
 ["SRM/采购","采购、供应商","准入、询比价、PO、交付","接口至ERP采购","供应商账户、价格、订单状态"],
 ["WMS","仓库","收货、上架、领料、出库、盘点","接口至ERP库存","数量、批次、日期和重复漏传"],
 ["MES","生产、质量","工单、领料、报工、完工、质量","接口至ERP生产成本","实际耗用、产量、在制和工单状态"],
 ["OA/费控","员工、审批人","申请、合同、发票、报销","接口至AP或GL","真实性、预算、科目、成本中心"],
 ["HR薪酬","HR、薪酬","人员、考勤、工资、社保个税","汇总或明细接口","人数、总额、费用归属和隐私权限"],
 ["资金银行","资金、授权人","支付指令、结果、流水、回单","银企直连和银行接口","成功状态、重复支付、未达和认领"],
 ["数电发票税务","税务、AR、AP","开票、收票、用途确认、申报","接口或底稿桥接","账票税口径和申报状态"]],font_size=7.8)

doc.add_heading("第2章 业务单据到总账的底层逻辑",1)
add_pic("02_posting.png","图2 业务单据进入总账的底层逻辑。")
add_para("【核心结论】采购订单和销售订单主要记录商业承诺；收货、出库、验收、开票、付款等事件才可能改变资产、负债、收入、费用或现金。是否自动生成凭证取决于业务实质和系统配置，不能只凭按钮名称判断。")
add_table(["对象","解决的问题","是否必然产生凭证","例子"],[
 ["业务单据","发生了什么、数量和责任归属","否","采购订单、销售订单"],
 ["业务事件","是否改变企业权利义务和资源","不一定","收货、签收、验收"],
 ["系统过账","将已批准事件写入正式账簿或库存","通常是关键触发点","收货过账、发货过账"],
 ["会计凭证","用科目、金额和维度反映经济影响","是","借原材料、贷收货暂估"],
 ["子模块","保存客户、供应商、资产、物料等明细","是财务明细账来源","AP未清项、资产卡片"],
 ["GL","汇总所有合法来源并形成试算和报表","是","科目余额表、财务报表"]],font_size=8)
add_para("四类凭证需要分开管理：自动凭证由ERP业务事件按配置生成；接口凭证来自OA、HR、银行等外部系统；手工凭证用于独立的会计判断或系统不覆盖事项；月结凭证用于预提、摊销、税费和结转。来源不同，排错路径也不同。")

doc.add_heading("第3章 组织、期间和主数据",1)
add_para("组织结构决定谁的业务进入哪一本账。法人、公司、工厂、库存组织、成本组织和账簿若配置不一致，同一笔实物流可能落到错误公司或错误利润中心。总账接手一套系统时，应先取得组织关系图、科目表、核算维度、期间日历、接口清单和自动记账规则。")
add_table(["主数据","关键字段","错误后果","财务检查"],[
 ["供应商","主体、税号、银行账户、付款条件","付错款、重复应付、税票不匹配","一户一码、账户变更双检"],
 ["客户","主体、信用额度、开票信息、付款条件","收入主体错误、超信用发货","信用和开票信息定期复核"],
 ["物料","基本单位、估价类、税码、批次","数量和成本自动错误","单位、估价和税码变更审批"],
 ["BOM","版本、用量、损耗、生效日","领料需求和产品成本错误","工程变更与生产订单版本一致"],
 ["科目及映射","控制科目、自动记账、报表项目","凭证和报表自动错","新增科目同时维护报表映射"],
 ["成本/利润中心","责任部门、有效期、层级","费用归属和经营分析错误","异常维度和失效中心扫描"],
 ["资产类别","原值、折旧、减值和处置科目","卡片和折旧凭证错误","类别与会计政策一致"],
 ["税码","税率、进销项和用途","税额、发票和申报错误","税务人员批准变更"]],font_size=7.8)

doc.add_heading("第4章 SAP 金蝶 用友的共同逻辑",1)
add_para("软件模块名称和菜单会变化，业务对象相对稳定。学习时把功能翻译成采购订单、收货、发票、未清项、物料移动、生产订单、资产卡片和会计凭证，就能跨系统迁移。下表只比较常见能力，不代表所有版本的固定菜单。")
add_table(["业务能力","SAP S4HANA","金蝶云星空","用友企业级产品","共同底层逻辑"],data["module_compare"],[3,3.4,3.4,3.4,5],7.2)
add_para("【实务提醒】面试或入职时应先问清产品线和版本。例如“用友”可能指U8、U9 Cloud、NC Cloud或YonBIP；“金蝶”可能指KIS、星辰、云星空或EAS。只说用过某厂商，不能说明实际能力。")

pagebreak(); doc.add_heading("第二篇 八大业务循环",0)

def cycle_chapter(num,title,img,chain,roles,judgments,checks,common):
    doc.add_heading(f"第{num}章 {title}",1)
    add_pic(img,f"图{num} {title}通用数据流。系统操作示意图，实际界面和凭证触发点取决于版本及配置。")
    doc.add_heading("业务和系统链条",2); add_para(chain)
    doc.add_heading("岗位分工",2); add_table(["岗位","主要职责"],roles,[4,12],8.5)
    doc.add_heading("关键会计判断",2)
    for x in judgments: bullet(x)
    doc.add_heading("月结检查",2)
    for x in checks: numbered(x)
    doc.add_heading("常见异常",2)
    for x in common: bullet(x)

cycle_chapter(5,"采购到付款","03_p2p.png","采购申请和订单建立商业承诺，收货记录实物到达，验收决定是否进入可用库存。收货暂估确保存货和负债截止完整，供应商发票经过三单匹配后形成财务应付，付款成功后还必须完成供应商未清项核销。",[
 ["采购","供应商选择、价格和PO；解释商务差异"],["仓库与质量","收货、检验、入库、退货及数量真实性"],["AP","发票验真匹配、应付、付款建议和核销"],["总账","暂估、进项税、AP与GL对账、月结截止"],["资金","付款批次、银行授权、结果和回单"]],[
 "收货是否意味着控制权已取得，是否应确认存货和负债。","暂估采用不含税金额，未取得合规抵扣凭证时通常不确认进项税。","发票差异是录入错误、供应商错误、合同变更，还是存货价格差异。","预付款与应付清账不能混为采购费用。"],[
 "取得未开票收货、未收货发票、未匹配发票和暂估账龄清单。","AP供应商余额与GL控制科目核对，并区分暂估应付和财务应付。","核对付款成功状态、银行流水和应付核销，避免付款完成但未清项。","分析价格差异、数量差异、退货和红字发票。"],["重复收货或重复发票","按含税金额暂估","直接在GL调平AP差异","付款账户变更未经独立验证"])

cycle_chapter(6,"销售到收款","04_o2c.png","销售订单、发货、签收、收入、开票、应收和收款是不同事件。系统可以把多个事件配置在相近时间自动过账，但会计仍需判断履约义务和控制权转移，税务人员另行判断增值税纳税义务。",[
 ["销售","客户、合同、订单、价格、交付和催收"],["仓库物流","拣货、出库、运输和签收证据"],["AR","开票、应收、收款认领、核销和账龄"],["总账","收入截止、成本结转、AR与GL对账"],["信用管理","额度、逾期和订单冻结"]],[
 "发货是否已经满足收入确认条件，是否仍存在验收或退货权。","开票不当然等于会计收入，未开票也不当然不能确认收入。","销售出库和收入凭证可能由不同事件触发，需要共同做截止测试。","第三方付款、预收款和客户之间的对应关系。"],[
 "订单、交付、签收、收入、开票、应收和收款七方核对。","分析已发货未过账、已过账无成本、已收入未开票和先开票未交付。","复核客户账龄、贷方余额、未知收款和跨公司代收。","AR与GL按公司、币种、客户和过账状态对账。"],["把开票日作为所有收入时点","发票红冲但AR未冲","客户收款认领错误","发货接口失败导致库存和收入截止不一致"])

doc.add_heading("第7章 费用报销和服务采购",1)
add_para("员工报销通常从OA或费控系统开始。电子化解决收集、审批和流转，但无法自动证明业务真实性，也无法替代对受益人、费用归属、预算、税务用途和付款主体的判断。服务采购则通常应走供应商、合同、验收和AP流程，不能为了方便全部塞入员工报销。")
add_table(["控制维度","解决的问题","系统字段或证据","常见错误"],[
 ["发票","票面和重复性","数电票标识、验真、查重","真票无业务、重复报销"],["审批","授权和业务确认","申请人、部门负责人、预算人","事后补签、拆单绕限额"],["费用科目","会计性质","差旅、招待、服务、资产等","按发票名称机械选科目"],["成本中心项目","责任和受益对象","部门、项目、客户、工单","默认部门未更新"],["预算","是否获资源授权","预算科目、占用和释放","预算通过被误解为会计合规"],["付款主体","债务和资金流","员工垫付、供应商直付","公司、发票、合同主体混乱"]],font_size=8)
add_para("【总账检查点】核对费用系统已审批未生成凭证、已付款未核销、跨期未报销、员工借款和成本中心异常。服务已接受但未到票时，应根据验收和合同评估预提，不应等待发票才确认费用。")

doc.add_heading("第8章 存货 生产和成本核算",1)
add_para("制造业成本链条是采购入库、原材料、生产订单、BOM、领退料、人工和制造费用、报工、在产品、完工入库、产成品和销售成本。WMS负责实物仓储，MES负责车间执行，ERP库存和生产模块保存正式业务对象，成本模块进行价值归集与分配，GL承接最终财务影响。")
add_table(["错误源","系统表现","成本后果","排查入口"],[
 ["BOM版本或用量错误","需求和领料异常","材料数量差异、单位成本失真","BOM版本、ECN和工单展开"],["计量单位错误","数量放大或缩小","负库存、异常耗用和成本","物料单位及换算"],["负库存","先出后入或跨期补单","移动平均等计价错误","物料收发明细和过账时间"],["工单未关闭","成本和在制长期滚动","在产品及差异不能结清","工单状态和结算日志"],["报工或产量错误","投入与产出不配比","单位成本异常","MES、报工和完工入库"],["费用未归集","分配池不完整","产品成本少计、期间费用错计","成本中心和费用科目"],["结账顺序错误","成本运行失败或取数不全","存货和销售成本未最终确定","月结任务依赖及日志"]],font_size=8)
add_para("【核心结论】ERP自动算成本只说明系统按既定主数据和规则执行。若BOM、数量、价格、费用池、分配基础或工单状态错误，系统会稳定而快速地生成错误成本。财务主管必须复核成本桥：期初在制加本期投入，等于完工转出、期末在制和单独确认的异常损失。")

doc.add_heading("第9章 固定资产和在建工程",1)
add_para("固定资产模块以资产卡片为明细账。采购发票可以先进入AP，验收和达到预定可使用状态决定资本化和折旧起点；复杂建设项目先在在建工程归集，再按可使用状态转固。资产类别通常决定原值、累计折旧、减值和处置科目，使用部门决定折旧费用成本中心。")
add_table(["环节","系统对象","会计关注","月结检查"],[
 ["采购验收","PO、收货、发票、验收单","资产还是费用；可使用状态","已验收未建卡"],["在建工程","项目/WBS/CIP明细","可直接归属成本、资本化时点","长期挂账和已使用未转固"],["建卡","资产卡片、类别、部门","原值、年限、残值、折旧方法","卡片与发票、实物一一对应"],["折旧","折旧运行和凭证","起止期间和费用归属","运行成功、金额波动、未过账"],["处置","报废/出售单、清理凭证","批准、收入、税费和损益","已报废仍折旧或实物已无"]],font_size=8)

doc.add_heading("第10章 薪酬系统和财务",1)
add_para("HR保存人员和组织，考勤及绩效形成工资计算基础，薪酬系统计算应发、社保公积金、个税和实发，ERP接收按人员或汇总维度生成的会计数据。工资不能全部记管理费用。生产人员进入生产成本或制造费用，研发人员按研发项目和资本化判断归集，销售及管理人员进入相应期间费用。")
add_para("【总账检查点】以人员花名册为起点，核对工资人数、应发总额、公司和个人承担社保、个税、实发银行批次及应付职工薪酬余额。接口总额相等还不够，还需检查部门、成本中心和薪酬项目映射。")

doc.add_heading("第11章 资金 银行和银企直连",1)
add_para("银企直连是资金系统或ERP通过银行接口发送支付指令、查询余额和获取流水回单。它缩短录入链路，但不能消除未达账项。银行可能已受理而ERP状态未回传，ERP可能已生成付款而银行拒付，手续费和利息也可能没有对应业务单据。")
add_table(["状态","财务判断","系统处理"],[
 ["待发送","尚未进入银行","可按权限修改或撤回"],["发送中/未知","不能假定失败，也不能盲目重发","使用唯一指令号向银行查询"],["成功","银行已执行","生成或确认银行过账并完成清账"],["失败","银行未执行","查失败原因，冲销预过账或恢复未清项"],["退汇","先成功后退回","记录退汇、恢复应付或客户款项并重新审批"]],font_size=8.2)
add_para("银行余额调节表按每个账户、币种和月末日期编制。系统联网后仍应核对银行对账单余额与ERP账面余额，并逐项解释在途、未达、错账、手续费、利息和未知收款。")

doc.add_heading("第12章 发票 ERP和税务系统",1)
add_para("ERP、数电发票平台和电子税务局不是同一系统。ERP以业务和会计为核心；发票平台记录开票、收票、验真、用途和红字；电子税务局承接申报缴税。三者的数据可以接口连接，但会因未开票收入、预收开票、暂估、进项用途、转出和跨期更正产生合理差异。")
add_para("【核心结论】发票真实只能证明票面数据来自合法开票渠道，不能单独证明交易真实、费用与企业经营相关或进项可以抵扣。AP必须把发票与合同、订单、收货验收、付款主体和实际受益人相匹配。")

pagebreak(); doc.add_heading("第三篇 总账 月结 对账和报表",0)
doc.add_heading("第13章 GL总账的真正职责",1)
add_para("总账接收AP、AR、存货、成本、FA、薪酬、资金和税务数据，并负责会计期间、凭证质量、科目及维度、模块对账和报表完整性。总账人员每天可能不录入大量业务凭证，但必须知道每类凭证由什么事件生成、如何冲销、怎样追到源单、失败时找谁处理。")
add_table(["凭证类型","来源","适用事项","主要风险","复核方法"],[
 ["自动凭证","ERP业务模块","收货、出库、开票、折旧、成本结算","配置错会批量自动错","抽样追源单并核自动记账规则"],["接口凭证","OA、HR、银行等","费用、工资、银行流水","漏传、重复、字段映射错","记录数、金额、批次和唯一号"],["手工凭证","GL人工录入","独立会计判断、系统未覆盖事项","缺依据、绕过子模块、维度缺失","底稿、审批、附件及控制科目限制"],["月结凭证","模板或批任务","预提、摊销、税费、结转","模板过期、重复运行、跨期","本期依据、反转设置和总额复算"]],font_size=8)

doc.add_heading("第14章 ERP月结完整SOP",1)
add_pic("05_close.png","图14 ERP月结依赖关系。实际顺序应结合企业系统触发点和结账依赖确定。")
add_table(["序号","阶段","关键动作","计划时间","责任人","完成证据","主要风险"],data["close_steps"],[1,2.5,5,1.8,2.5,4.5,3.5],7.5)
add_para("【财务主管复核点】主管不只查看“已结账”状态，还要检查未完成事项、异常金额、差异桥、接口批次和反结账申请。技术上能够关账，不代表会计和业务已完整关闭。")

doc.add_heading("第15章 子模块与GL核对矩阵",1)
add_table(["模块","GL范围","子模块报表","总账依据","统一口径","常见差异","排查路径","责任"],data["recons"],[1.5,3.2,4,3,3,4,4,2.5],6.8)
add_para("核对前先统一公司、账簿、币种、截止时间、过账状态和控制科目范围。汇总差额确定后，依次下钻到客户或供应商、物料或资产、具体单据和会计凭证。余额相等也应检查是否存在同额错记、错误主体、错误维度或手工绕过子模块。")

doc.add_heading("第16章 报表 期间关闭和错误更正",1)
add_para("GL形成试算平衡表，再由科目和维度映射到财务报表。新增科目、报表版本、现金流量项目和重分类规则都可能导致账平但表错。报表复核包括资产负债表平衡、利润表与权益变动、现金及现金等价物勾稽、当期损益结转和重要科目变动分析。")
add_table(["错误类型","应在哪一层处理","处理原则"],[
 ["原始业务事实错误","业务源系统或ERP业务模块","按可追溯的冲销和重做流程修改源单，再让凭证随单据更新"],
 ["主数据或系统配置错误","主数据/配置层","先停止错误扩散，IT或实施修复，财务验证受影响交易"],
 ["接口遗漏或重复","接口和目标模块","按消息ID及批次补传或冲销，防止重跑形成重复"],
 ["独立会计判断错误","GL或相应子模块","按会计依据和审批做调整，记录期间及报表影响"],
 ["报表列报映射错误","报表层","账务正确时修改取数或列报规则，不为报表位置机械调账"]],font_size=8)

pagebreak(); doc.add_heading("第四篇 主数据 接口 权限和异常",0)
doc.add_heading("第17章 主数据治理",1)
add_para("主数据治理包括申请、查重、审批、创建、变更、冻结和停用。业务部门通常拥有客户、供应商、物料和BOM的业务属性，财务拥有会计科目、税码、付款条件和财务视图，IT维护系统平台。任何关键字段变更都应有生效日、理由、批准人和操作日志。")
add_para("【财务主管复核点】每月检查新增及变更供应商银行账户、异常税码、失效成本中心、重复物料和未经批准的BOM变更；每季度复核主数据权限和长期未使用记录。")

doc.add_heading("第18章 接口 批处理和日志",1)
add_para("接口控制不能只看“成功”标志。最少应核对源系统记录数和金额、目标系统成功数和金额、失败清单、重复清单、批次号及重跑记录。对于部分成功的批次，只补传失败记录；对状态未知的银行指令，必须先向银行查询，不能直接重发。")
add_table(["接口控制","目的","示例"],[
 ["唯一业务键","防止重复创建","源单号加行号、发票唯一标识、支付指令号"],["记录数控制","发现漏行和重复","源系统100行，目标成功99行、失败1行"],["金额控制","发现字段丢失或符号错误","工资应发和实发总额、报销批次金额"],["状态闭环","确认业务是否最终完成","发送、受理、成功、失败、退回"],["重跑审批","控制重复和期间影响","说明原因、范围、预期结果及回退方案"]],font_size=8)

doc.add_heading("第19章 权限与职责分离",1)
add_table(["不相容职责","主要风险","优先控制","人员不足时的补偿控制"],[
 ["供应商账户维护与付款","篡改账户并转出资金","角色分离、账户变更独立回拨","财务负责人逐笔复核变更后首次付款"],["制单与审核过账","无独立检查的错误或舞弊","制单、审核、过账按权限拆分","金额限额、事后独立复核和日志"],["资产建卡与处置批准","虚构或私自处置资产","实物、卡片和处置审批分离","定期盘点和处置收入核对"],["期间关闭与反结账","操纵历史数据","反结账紧急权限和工单","修改前后报表快照及高层批准"],["系统管理员与业务审批","管理员直接改变业务结果","管理员只维护平台，业务变更走工单","管理员日志由独立人员复核"]],font_size=8)

doc.add_heading("第20章 69项ERP常见异常问题库",1)
add_para("每项异常都按现象、原因、检查数据、财务与IT分工、会计和月结影响、解决方案及预防控制展开。排查时先确认业务事实和源单，再检查主数据、接口、凭证和报表。")
cat_names={"MD":"主数据","P2P":"采购与应付","O2C":"销售与应收","COST":"库存生产与成本","EFP":"费用资产与薪酬","BT":"银行与税务","IF":"接口与批处理","GL":"总账月结与报表","AUTH":"权限与内控"}
for cat in cat_names:
    doc.add_heading(cat_names[cat],2)
    for e in [x for x in data["exceptions"] if x["category"]==cat]:
        doc.add_heading(f'{e["id"]} {e["symptom"]}',3)
        add_table(["项目","处理内容"],[
            ["可能原因",e["cause"]],["检查数据",e["check"]],["处理人",e["owner"]],["财务处理",e["finance"]],["IT处理",e["it"]],["会计影响",e["impact"]],["是否影响月结",e["close"]],["预防控制",e["prevent"]]], [3.2,12.8],7.6)

pagebreak(); doc.add_heading("第五篇 贯穿案例与岗位训练",0)
doc.add_heading("第21章 模拟公司和案例数据",1)
add_table(["项目","设定"],[["公司",data["company"]["name"]],["期间",data["company"]["period"]],["业务特征",data["company"]["profile"]],["产品",data["company"]["product"]],["系统",data["company"]["systems"]]], [3.5,12.5],8.5)
add_para("案例从供应商创建开始，贯穿采购订单、分批收货、月末暂估、次月到票和价格差异、付款、生产领料、费用归集、完工成本、销售订单、发货、收入开票、收款、资产折旧和月结。各案例的单据编号、金额和后续状态相互衔接。")

doc.add_heading("第22章 20项ERP实操训练",1)
for c in data["cases"]:
    doc.add_heading(f'{c["id"]} {c["title"]}',2)
    add_table(["项目","内容"],[
      ["所属模块",c["module"]],["业务背景",c["background"]],["原始单据",c["docs"]],["系统操作",c["actions"]],["会计分录",c["journal"]],["系统结果",c["result"]],["总账检查",c["check"]],["常见错误",c["errors"]]
    ],[3,13],8)
    add_para("完成标准：能够说明业务由谁发起、哪个动作改变账务、生成什么子模块记录和凭证、月底与什么报表核对，以及错误应在源单、配置、接口还是GL处理。",boldlead="【训练要求】")

doc.add_heading("第23章 上岗与学习路线",1)
doc.add_heading("入职第一周取得的资料",2)
for x in ["组织架构、法人和账簿清单","月结日历、任务依赖和责任矩阵","科目表、辅助核算和报表映射","客户、供应商、物料、BOM和资产主数据规则","自动记账规则及控制科目","接口清单、批任务、错误日志和IT联系人","AR、AP、存货、成本、FA、银行和税务对账底稿","权限矩阵、反结账及紧急权限流程"]: bullet(x)
doc.add_heading("面对陌生ERP的提问顺序",2)
for x in ["业务从哪个模块和哪张源单开始？","哪个状态代表业务完成，哪个动作触发会计过账？","系统生成什么子模块记录和会计凭证？","凭证科目和维度由哪些主数据或配置决定？","月末使用哪张子模块报表与GL核对？","异常如何追到源单、接口批次和操作日志？","错误必须回业务端改，还是属于独立会计调整？","该动作未完成会阻断哪个后续模块或月结任务？"]: numbered(x)
doc.add_heading("建议观看的实操主题",2)
add_table(["软件","主题","优先来源","学习重点","版本风险"],[
 ["SAP S4HANA","采购收货与供应商发票校验","SAP Help Portal及官方Learning","PO、收货、GR/IR、发票和差异","Fiori应用、云版和本地版可能不同"],
 ["SAP S4HANA","固定资产与月末关闭","SAP官方帮助","卡片、折旧运行和关闭依赖","版本和配置差异"],
 ["金蝶云星空","采购入库、应付和凭证生成","金蝶官方学习中心","单据下推、核销和智能会计平台","需确认企业版及版本"],
 ["金蝶云星空","存货核算、产品成本和系统结账","金蝶官方学习中心","库存关账、成本计算、凭证和总账结账","旧版课程菜单可能变化"],
 ["用友企业级","采购供应链、财务及资产","用友官方产品和培训资料","业务域集成、凭证和月结","先确认YonBIP或NC Cloud等产品线"]],font_size=7.8)

doc.add_heading("附录A SAP 金蝶 用友资料来源",1)
add_table(["软件产品","资料名称","版本或时间","用途","链接"],data["sources"],[2.5,4,3,5,7],6.6)
add_para("资料检索基准日为2026年9月10日。厂商页面可能更新，正式照着界面操作前应再次核验产品线、版本和企业配置。社区个人经验只用于发现线索，不作为版本无关的统一规则。")

doc.add_heading("附录B 总账会计的每日 每周 每月检查",1)
add_table(["频率","检查事项","输出"],[
 ["每日","接口失败、银行未知状态、大额付款、重复发票、未认领收款","异常日报和责任人"],
 ["每周","暂估和未清项账龄、负库存、零成本、未关闭工单、CIP账龄","模块异常清单"],
 ["月末","截止、所有子模块对账、成本结算、税务桥接、报表校验、期间关闭","月结签字包"],
 ["季度","主数据变更、权限冲突、长期差异、接口趋势和报表映射","主管复核报告"],
 ["年度","资产盘点、主数据清理、权限年度复核、系统变更及归档","年度控制证据"]],font_size=8.2)

out=OUT/"企业财务软件与ERP实操手册.docx"
doc.save(out)
print(out)
