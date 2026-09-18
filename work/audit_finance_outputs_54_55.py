import json
import zipfile
from pathlib import Path

from docx import Document
from finance_cases_data_54_55 import CASES

ROOT = Path(r"C:\Users\HZJ\Desktop\Git\work")


def walk(node):
    yield node
    for child in node.get("children", {}).get("attached", []):
        yield from walk(child)


for case in CASES:
    ep = case["episode"]
    stem = f"齐昊老师_学员实战案例_第{ep}集"
    docx_path = ROOT / f"{stem}_文字思维导图.docx"
    xmind_path = ROOT / f"{stem}_思维导图.xmind"
    with zipfile.ZipFile(docx_path) as z:
        assert z.testzip() is None
        assert "word/document.xml" in z.namelist()
        xml = z.read("word/document.xml").decode("utf-8")
        assert "TODO" not in xml and "PLACEHOLDER" not in xml
    doc = Document(docx_path)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert case["title"] in all_text and case["video_id"] in all_text
    assert "不是逐字稿" in all_text
    assert len(doc.tables) == 1
    headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert len(headings) >= len(case["sections"])
    with zipfile.ZipFile(xmind_path) as z:
        assert z.testzip() is None
        assert {"content.json", "manifest.json"}.issubset(z.namelist())
        sheets = json.loads(z.read("content.json"))
        assert [s["title"] for s in sheets] == ["01 总览图", "02 详细知识树"]
        groups = [list(walk(s["rootTopic"])) for s in sheets]
        ids = [topic["id"] for group in groups for topic in group]
        assert len(ids) == len(set(ids))
        assert all(case["title"] in s["rootTopic"]["title"] for s in sheets)
    print(ep, "OK", "headings", len(headings), "topics", [len(group) for group in groups])
