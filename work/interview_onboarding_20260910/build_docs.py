from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent;O=R/'deliverables';D=json.loads((R/'content.json').read_text(encoding='utf8'));NAVY='17365D';PALE='F3F6F9'
def setup(title,sub):
 d=Document();s=d.sections[0];s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(2);s.right_margin=Cm(2)
 for n,z in [('Normal',10),('Title',25),('Heading 1',18),('Heading 2',14),('Heading 3',11.2)]:
  st=d.styles[n];st.font.name='Microsoft YaHei';st._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.size=Pt(z)
  if n!='Normal':st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.keep_with_next=True
 d.styles['Normal'].paragraph_format.line_spacing=1.2;d.styles['Normal'].paragraph_format.space_after=Pt(4)
 h=s.header.paragraphs[0];h.text=title;h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 d.add_paragraph().paragraph_format.space_after=Pt(75);p=d.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(title);p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(sub);r.bold=True;r.font.size=Pt(17);d.add_paragraph().paragraph_format.space_after=Pt(65)
 for x in ['适用岗位：总账会计、财务主管','版本日期：2026年9月10日','原则：真实表达审计经历，不虚构企业执行经验']:p=d.add_paragraph(x);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 d.add_page_break();return d
def shade(c,x):
 pr=c._tc.get_or_add_tcPr();e=pr.find(qn('w:shd'))
 if e is None:e=OxmlElement('w:shd');pr.append(e)
 e.set(qn('w:fill'),x)
def table(d,h,rows,fs=7.2):
 t=d.add_table(rows=1,cols=len(h));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,x in enumerate(h):
  c=t.rows[0].cells[i];c.text=str(x);shade(c,NAVY);c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
  for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(fs)
 for ri,row in enumerate(rows):
  cs=t.add_row().cells
  for i,x in enumerate(row):
   cs[i].text='' if x is None else str(x);cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if ri%2:shade(cs[i],PALE)
   for p in cs[i].paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.02
    for r in p.runs:r.font.size=Pt(fs)
 pr=t.rows[0]._tr.get_or_add_trPr();e=OxmlElement('w:tblHeader');e.set(qn('w:val'),'true');pr.append(e);d.add_paragraph().paragraph_format.space_after=Pt(0)
def toc(d):p=d.add_paragraph();e=OxmlElement('w:fldSimple');e.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u');p._p.append(e)
def bullets(d,xs):
 for x in xs:d.add_paragraph(x,style='List Bullet')
def question(d,x,full=False):
 d.add_heading(f"{x['id']} {x['question']}",2)
 rows=[['面试官真正想考什么',x['testing']],['回答思路',x['logic']],['30秒回答',x['short']]]
 if full:rows += [['1分钟回答',x['minute']],['2分钟深入回答',x['deep']]]
 rows += [['结合审计经验可以说什么',x['audit']],['容易踩的坑',x['pitfall']],['可能追问',x['followup']]]
 table(d,['项目','内容'],rows,7.5)

# interview book
d=setup('总账会计与财务主管面试实战手册','从真实经历到结构化表达')
d.add_heading('使用说明',0);d.add_paragraph('专业题优先使用“结论—判断逻辑—操作步骤—风险点”；行为和项目管理题才使用STAR。审计中接触、检查和参与的事项不得表述为企业财务中负责或执行。')
d.add_heading('目录',0);toc(d);d.add_page_break()
d.add_heading('第一篇 岗位能力模型',0)
table(d,['模块','总账要求','主管要求','准备优先级'],[['会计、月结、科目核对、报表','必须熟练','必须熟练并能复核','最高'],['税务、成本、ERP、资金','需要掌握并能排查','需要复核和组织','最高'],['内控、协同、Issue','能配合','必须建立闭环','主管最高'],['财务分析、团队、流程优化','需要掌握','需要组织和改进','中高']],7.6)
d.add_heading('第二篇 审计转企业表达',0)
table(d,['长度','结构','示范重点'],[['30秒','背景→优势→目标','四年审计、报表和风险基础；正在补企业执行'],['1分钟','再加迁移能力和补缺行动','不夸大ERP、报税、成本和月结经验'],['3分钟','增加真实审计案例','只说检查、识别、沟通和影响评估，不说企业执行']],7.5)
bullets(d,['离开审计：希望从外部检查转向内部持续负责和改进，不贬低事务所。','短板：企业日常操作和系统执行；用系统学习、案例训练和入职分阶段复核补足。','胜任理由：报表逻辑、异常识别、证据、截止和项目推进基础较强，且对缺口认知清楚。','不要说：审计什么都懂、企业财务很简单、我以前替客户做过月结。'])
d.add_heading('第三篇 情景题五步法',0);bullets(d,['事实：真实业务、合同、单据、系统和期间。','影响：科目、税务、现金、月结和报表。','处理：找Owner、取得资料、调账或调表。','复核：子模块、GL、税表、报表和凭证。','预防：责任人、截止、系统规则和异常清单。'])
d.add_heading('第四篇 核心30题',0)
for x in D['core']:question(d,x,True)
d.add_heading('第五篇 总账会计90题',0)
mods=[]
for x in D['gl']:
 if x['module'] not in mods:mods.append(x['module'])
for m in mods:
 d.add_heading(m,1)
 for x in [z for z in D['gl'] if z['module']==m]:question(d,x,False)
d.add_heading('第六篇 财务主管65题',0)
mods=[]
for x in D['sup']:
 if x['module'] not in mods:mods.append(x['module'])
for m in mods:
 d.add_heading(m,1)
 for x in [z for z in D['sup'] if z['module']==m]:question(d,x,False)
d.add_heading('第七篇 ERP经验不足怎么回答',0)
table(d,['场景','可信表达'],[['招聘要求SAP','我没有作为企业财务用户独立负责SAP月结，这一点会如实说明；但我在审计中查看过系统资料，正在系统训练P2P、O2C、AP、AR、成本和GL逻辑。'],['追问会不会操作','我目前能解释业务单据如何进入子模块和GL，以及常见对账和异常路径；具体事务码或菜单需要在贵公司环境中按权限和配置熟悉。'],['证明学习能力','入职后先取得流程、权限、历史凭证和月结清单，在复核下完成一轮，再独立承担。']],7.5)
d.add_heading('第八篇 面试官可能担心的弱点',0)
table(d,['担心','如何解释','如何证明','不要说'],[['不适应重复工作','理解企业财务需要持续执行和闭环','愿意从对账、凭证和月结做起','企业工作比审计轻松'],['ERP不熟','承认没有独立操作','展示系统底层流程和训练记录','SAP我看过所以算熟练'],['报税不足','区分账票税逻辑与申报执行','能解释数据来源并继续训练','我帮客户报过税'],['成本不足','已系统学习工单/BOM/在制/分配','能完整排查案例','成本就是审计存货'],['没有企业月结','熟悉结果和风险但未直接负责','展示Close Calendar和第一次月结计划','我审过所以会做']],6.9)
d.add_heading('第九篇 反向面试和JD拆解',0)
bullets(d,['月结通常T+几完成？哪些模块最容易延误？','团队如何分工，是否有专职成本和税务？','ERP、法人主体和报表体系是什么？','岗位空缺原因及前3个月最希望解决的问题？','历史账、Excel手工量和审计调整情况？','全盘账务是否意味着兼税务、成本、资金和报表？','主管岗位是否有团队，还是单人全盘会计？'])
d.add_heading('第十篇 六套完整模拟面试',0)
for i in range(1,7):
 role='总账会计' if i<=3 else '财务主管';d.add_heading(f'模拟面试{i}：{role}',1)
 table(d,['环节','建议时间','问题'],[['自我介绍','2分钟','请结合岗位介绍自己'],['转型动机','4分钟','为什么从审计转企业'],['专业题','10分钟',D['gl'][(i*7)%90]['question'] if i<=3 else D['sup'][(i*7)%65]['question']],['实操题','8分钟','给出一个月结异常并说明第一步查什么'],['追问','6分钟','你没有直接经验，如何降低入职风险'],['反向提问','5分钟','询问月结、团队、系统和岗位空缺']],7.2)
d.add_heading('附录 官方资料与版本提示',0);table(d,['文件','状态','用途','官方链接'],D['sources'],6.7)
d.save(O/'总账会计与财务主管面试实战手册.docx')

# cases
c=setup('总账、财务主管情景案例库','162个企业财务实战Case')
c.add_heading('使用说明',0);c.add_paragraph('每个案例要求先讲事实和影响，再讲处理、复核和预防。账错做凭证；账正确但列报位置错误时做报表重分类。')
c.add_heading('目录',0);toc(c);c.add_page_break()
groups=[]
for x in D['cases']:
 if x['category'] not in groups:groups.append(x['category'])
for g in groups:
 c.add_heading(g,0)
 for x in [z for z in D['cases'] if z['category']==g]:
  c.add_heading(f"{x['id']} {x['topic']}",2);table(c,['问题','实务处理'],[['先做什么',x['first']],['查什么',x['check']],['找谁',x['who']],['怎么判断',x['judgment']],['怎么调整',x['entry']],['如何复核和关闭',x['close']],['如何防止重复',x['prevent']]],7.3)
c.add_heading('附录 情景题自我训练',0);bullets(c,['先口头回答30秒，再看参考逻辑。','第二次回答必须说出资料、系统、责任人和复核证据。','不会的问题不猜：确认事实、查制度和准则、看历史、咨询专业人员、形成判断、复核处理。'])
c.save(O/'总账、财务主管情景案例库.docx')

# onboarding
o=setup('入职30、60、90天实战手册','从第一天到独立月结和小范围改善')
o.add_heading('使用说明',0);o.add_paragraph('前30天重点是看懂公司，不急于证明自己或全面改革。历史问题按资金、税务、报表、月结和法律风险排序。')
o.add_heading('目录',0);toc(o);o.add_page_break()
o.add_heading('第1章 Day 1 Checklist',0)
table(o,['事项','具体动作','当天成果'],[['汇报线和团队','确认直接上级、模块Owner和备岗','联系人清单'],['账号权限','申请ERP、OA、邮箱和共享盘','可登录但不索要他人密码'],['主体业务','了解法人、产品、收入采购生产模式','一页业务图'],['月结','取得Close Calendar和本月状态','最近截止和紧急事项'],['风险','询问当前最需要新员工解决的问题','Top 3问题']],7.4)
o.add_heading('第2章 Week 1资料和验证',0)
table(o,['优先级','资料','验证动作'],[['必须马上看','最近TB、报表、月结清单、资金和紧急Issue','从TB追一项到明细和原始单据'],['第一周看','科目表、账龄、暂估、预提、税务、存货、成本、固定资产','完成一项模块与GL核对'],['第一个月逐步看','借款、关联方、审计报告、审计调整、制度和历史问题','形成风险排序']],7.4)
plans=[['1-30天','看懂公司','业务、产品、系统、团队、月结、税务、报表和主要风险','业务财务全景、月结地图、Top10风险','业务负责人、各财务模块、IT','最近12个月TB和报表、关键台账','急于改革、误解历史处理','能解释账如何形成'],['31-60天','开始独立','独立完成一次月结、主要科目核对、理解申报和报表','月结包、对账底稿、Issue Log','税务、成本、仓库、销售采购','模块报表、申报和期后数据','关键路径延误','在复核下稳定完成月结'],['61-90天','小范围改善','清理高风险历史问题，选择2-3项流程优化','Action Plan和效果指标','流程Owner和IT','逾期、返工和异常数据','改革范围过大','改进可实施且不破坏现有控制']]
o.add_heading('第3章 完整30/60/90天计划',0);table(o,['阶段','目标','任务','交付','认识谁','取得数据','风险','成功标准'],plans,6.5)
o.add_heading('第4章 第一次月结',0)
month=[['T-5','取得日历、确认Owner和关键路径','不重新设计流程；先跟随历史机制'],['T-3','暂估迟交、验收缺失','列缺口、找采购销售、确认金额和时限'],['T-1','负库存、工单未关、工资修改','判断阻断项；升级仓库生产HR负责人'],['T','业务截止','例外必须有原因、批准、期限和补救'],['T+1','回款未认领、资产未转固','AR/资金认领；资产与使用部门确认状态'],['T+2','税务差异、应付借方、长期未达','查业务实质、账票税和银行证据'],['T+3','报表和锁账','复核重大调整、重分类、勾稽和未决事项']]
table(o,['时间','问题与动作','新人原则'],month,7.3)
o.add_heading('第5章 接手历史烂账',0)
table(o,['时间','先做','不要做'],[['第一天','保护数据和权限；了解资金、申报、月结和重大风险','批量冲销、改科目、删除资料'],['第一周','银行、税务、TB、审计调整、权限和重大往来','承诺短期清完所有问题'],['第一个月','按资金/税务/报表/法律/持续扩大排序','平均用力处理所有小差异'],['后续','逐项Owner、期限、方案、复核和Closure','问题解决后不做根因分析']],7.3)
o.add_heading('第6章 成熟企业接岗',0);o.add_paragraph('ERP成熟、月结T+3和分工清晰的企业，更强调按流程快速执行、异常分析、系统控制、风险复核和流程优化。新人应先掌握SOP、接口、报表口径和服务水平，不要把成熟流程改成个人Excel。')
o.add_heading('第7章 不会怎么办',0);bullets(o,['确认事实和问题边界。','查公司制度、会计政策和系统流程。','查适用准则或税法及执行日期。','查看历史处理但不盲从。','咨询业务、内部专业人员、IT、税务顾问、审计师或律师。','形成书面判断并经适当人员复核后处理。'])
o.add_heading('第8章 个人补缺清单',0)
table(o,['能力','当前基础','补强方式','真正熟练的条件'],[['报表、科目、审计风险','较好','转换为企业处理和Closure表达','参与真实月结'],['ERP操作','不足','视频/沙盘/模拟单据链','在真实权限环境反复操作'],['税务申报','不足','练账票税桥和官方申报资料','至少参与一轮申报'],['制造成本','不足','工单/BOM/在制/分配案例','参与成本月结'],['资金银行','一般','银行调节和支付控制案例','接触真实银企流程'],['跨部门推进','较好','把审计沟通转为Owner/期限/关闭','承担企业Issue']],7.2)
o.add_heading('附录 90天复盘问题',0);bullets(o,['我能否解释主要业务如何进入GL？','是否独立完成过一次月结并保留证据？','哪些能力仍需复核或帮助？','Top10风险是否有Owner和期限？','提出的改进是否少量、明确、可实施？'])
o.save(O/'入职30、60、90天实战手册.docx')
print('docs done')
