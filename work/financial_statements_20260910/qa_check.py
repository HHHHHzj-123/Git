from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from zipfile import ZipFile
import json
r=Path(__file__).parent;o=r/'deliverables'
doc=Document(o/'企业财务报表编制、调整、重分类与校验实操手册.docx');txt='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
for x in ['负数不等于错误','账错了做凭证','调整前TB借贷各12,787万元','期末1,405万元','50个报表错误案例']:assert x in txt,x
wb=load_workbook(o/'财务报表编制与校验实操工具.xlsx',data_only=False);assert len(wb.sheetnames)==22;assert wb['01-调整前原始TB'].max_row>=66;assert wb['03-AJ账务调整'].max_row==16;assert wb['05-RJ报表重分类'].max_row==13;assert wb['17-报表异常整改'].max_row==55
with ZipFile(o/'财务报表编制与校验全景.xmind')as z:c=json.loads(z.read('content.json'));assert len(c)==3;raw=json.dumps(c,ensure_ascii=False);assert '负数≠错误' in raw and '现金流量表' in raw
pages=len(list((r/'qa'/'rendered').glob('*.png')))
print('pages',pages,'tables',len(doc.tables),'sheets',len(wb.sheetnames),'xmind',len(c))
