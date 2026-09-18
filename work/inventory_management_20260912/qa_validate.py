from pathlib import Path
from zipfile import ZipFile
import json,re
from openpyxl import load_workbook
from docx import Document
b=Path(__file__).resolve().parent; out=b/'deliverables'
x=out/'企业存货管理工具包.xlsx'; wbv=load_workbook(x,data_only=True,read_only=True); wbf=load_workbook(x,data_only=False,read_only=True)
dash=[wbv['00-管理看板'][f'B{i}'].value for i in range(6,19)]
errs=[]; formulas=0
for ws in wbf.worksheets:
 for row in ws.iter_rows():
  for c in row:
   if isinstance(c.value,str) and c.value.startswith('='): formulas+=1
for ws in wbv.worksheets:
 for row in ws.iter_rows():
  for c in row:
   if isinstance(c.value,str) and re.match(r'^#(REF!|DIV/0!|VALUE!|NAME\?|N/A|NUM!|NULL!|SPILL!|CALC!)$',c.value): errs.append((ws.title,c.coordinate,c.value))
doc=Document(out/'企业存货管理实战手册.docx'); headings=[p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
with ZipFile(out/'企业存货管理全景.xmind') as z: xm=json.loads(z.read('content.json'))
report={'xlsx_sheets':len(wbf.sheetnames),'formulas':formulas,'dashboard':dash,'formula_errors':errs,'doc_headings':len(headings),'doc_tables':len(doc.tables),'xmind_sheets':len(xm),'files':{p.name:p.stat().st_size for p in out.iterdir()}}
(b/'qa'/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False,indent=2))
