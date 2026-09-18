from pathlib import Path
from zipfile import ZipFile
import json,re
from datetime import date
from openpyxl import load_workbook
from PIL import Image,ImageDraw

BASE=Path(__file__).resolve().parent;OUT=BASE/'deliverables';QA=BASE/'qa';C=QA/'contacts';C.mkdir(parents=True,exist_ok=True)

def contact(files,out,cols=4,size=(340,440)):
    files=list(files);rows=(len(files)+cols-1)//cols;can=Image.new('RGB',(cols*size[0],rows*size[1]),'white');dr=ImageDraw.Draw(can)
    for i,p in enumerate(files):
        im=Image.open(p).convert('RGB');im.thumbnail((size[0]-12,size[1]-28));x=(i%cols)*size[0]+6;y=(i//cols)*size[1]+20;can.paste(im,(x+(size[0]-12-im.width)//2,y));dr.text((x,(i//cols)*size[1]+3),p.stem,fill='black')
    can.save(out)

d=json.loads((BASE/'content.json').read_text(encoding='utf-8'));asof=date(2026,8,31)
total=sum(x[9] for x in d['ar']);overdue=sum(x[9] for x in d['ar'] if date.fromisoformat(x[4])<asof);over90=sum(x[9] for x in d['ar'] if (asof-date.fromisoformat(x[4])).days>90)
wb=load_workbook(OUT/'应收账款管理工具包.xlsx',data_only=False);forms=[];errs=[]
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('='):forms.append(f'{ws.title}!{c.coordinate}:{c.value}')
            if isinstance(c.value,str) and re.search(r'#REF!|#DIV/0!|#VALUE!|#NAME\?|#N/A|#NUM!|#SPILL!|#CALC!',c.value):errs.append(f'{ws.title}!{c.coordinate}')
assert wb['02-应收明细'].max_column==18
assert not any('DATEVALUE' in x for x in forms)
with ZipFile(OUT/'应收账款管理实战手册.docx') as z:doc_bad=z.testzip()
with ZipFile(OUT/'应收账款管理全景.xmind') as z:xm_bad=z.testzip();xm_sheets=len(json.loads(z.read('content.json')))
pages=sorted((QA/'docx_render').glob('page-*.png'),key=lambda p:int(re.search(r'(\d+)',p.stem).group()))
for i in range(0,len(pages),12):contact(pages[i:i+12],C/f'doc_{i+1}_{min(i+12,len(pages))}.png')
imgs=sorted([p for p in QA.glob('*.png') if not p.name.startswith('doc')])
for i in range(0,len(imgs),8):contact(imgs[i:i+8],C/f'xlsx_{i+1}_{min(i+8,len(imgs))}.png',cols=2,size=(650,390))
report={'sheets':len(wb.sheetnames),'formulas':len(forms),'formula_errors':errs,'total_ar':total,'overdue':overdue,'overdue_rate':overdue/total,'over90':over90,'pages':len(pages),'doc_zip_bad':doc_bad,'xmind_sheets':xm_sheets,'xmind_zip_bad':xm_bad}
(QA/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False,indent=2))
