from __future__ import annotations
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE=Path(__file__).resolve().parent; OUT=BASE/'deliverables'; D=json.loads((BASE/'content.json').read_text(encoding='utf-8'))
NAV='17365D'; PALE='F4F7FA'; LINE='D9D9D9'

def shade(c,color):
    p=c._tc.get_or_add_tcPr(); e=p.find(qn('w:shd'))
    if e is None: e=OxmlElement('w:shd'); p.append(e)
    e.set(qn('w:fill'),color)
def borders(t):
    p=t._tbl.tblPr; b=p.find(qn('w:tblBorders'))
    if b is None: b=OxmlElement('w:tblBorders'); p.append(b)
    for x in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement(f'w:{x}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),LINE); b.append(e)
def repeat(row):
    p=row._tr.get_or_add_trPr(); e=OxmlElement('w:tblHeader'); e.set(qn('w:val'),'true'); p.append(e)
def table(doc,h,rows,widths=None,font=8.4):
    t=doc.add_table(rows=1,cols=len(h)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,v in enumerate(h):
        c=t.rows[0].cells[i]; c.text=str(v); shade(c,NAV); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(font)
    repeat(t.rows[0])
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text='' if v is None else str(v);cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2:shade(cs[i],PALE)
            for p in cs[i].paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.05
                for r in p.runs:r.font.name='Microsoft YaHei';r.font.size=Pt(font)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths):row.cells[i].width=Cm(w)
    borders(t);doc.add_paragraph().paragraph_format.space_after=Pt(2);return t
def para(doc,text='',lead=None):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.25
    if lead:p.add_run(lead).bold=True
    p.add_run(text);return p
def bullets(doc,items):
    for x in items:
        p=doc.add_paragraph(style='List Bullet');p.paragraph_format.space_after=Pt(2);p.add_run(x)
def head(doc,text,level=1):
    p=doc.add_heading(text,level=level);p.paragraph_format.keep_with_next=True;return p
def flow(doc,steps):
    table(doc,['节点','动作','管理结果'],[[i+1,s,'继续' if i<len(steps)-1 else '关闭'] for i,s in enumerate(steps)],[1.2,12.2,2.2],8.7)
def toc(doc):
    p=doc.add_paragraph();r=p.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u');r._r.addnext(f)
def setup():
    doc=Document();s=doc.sections[0];s.top_margin=Cm(2);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2.1);s.right_margin=Cm(2)
    for n,z in [('Normal',10),('Title',26),('Subtitle',12),('Heading 1',18),('Heading 2',14),('Heading 3',11.5)]:
        x=doc.styles[n];x.font.name='Microsoft YaHei';x._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');x.font.size=Pt(z);x.font.color.rgb=RGBColor(0,0,0)
        if n.startswith('Heading'):x.font.bold=True;x.paragraph_format.space_before=Pt(10);x.paragraph_format.space_after=Pt(5);x.paragraph_format.keep_with_next=True
    return doc

def case(doc,c):
    head(doc,f"{c[0]} {c[1]}",3)
    table(doc,['处理环节','财务主管动作'],[['第一步',c[2]],['沟通对象',c[3]],['处理方案',c[4]],['升级条件',c[5]],['工作留痕','保存客户明细、证据、书面沟通、审批、下一步、责任人和关闭结果。'],['复核关闭','确认账务、客户余额、信用状态及后续业务限制均已同步。']],[3.0,12.7],8.7)

def build():
    doc=setup();p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('应收账款管理实战手册')
    p=doc.add_paragraph(style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('总账会计 财务主管 财务经理工作版')
    doc.add_paragraph();p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('版本日期 2026年9月11日').bold=True
    doc.add_paragraph();para(doc,'本手册从客户准入、信用审批、销售放行、收入与应收确认、对账催收、回款核销、争议处理、预期信用损失到坏账核销，完整说明财务主管如何把应收账款从账面余额转化为可管理的经营事项。')
    para(doc,'应收管理的目标不是单纯压低应收余额，而是在可接受的信用风险下支持有质量的销售，并让收入按约转化为现金。',lead='核心结论：')
    doc.add_page_break();head(doc,'目录');toc(doc);doc.add_page_break()

    head(doc,'第一章 应收账款管理全景')
    head(doc,'1.1 应收管理管什么',2)
    table(doc,['对象','核算关注','管理关注'],[['客户','名称、主体和辅助核算','偿付能力、信用变化和集中度'],['合同与订单','应收确认和金额','账期、付款条件、验收与违约条款'],['应收余额','发生、收款和核销','是否逾期、争议、风险和责任人'],['现金回收','银行到账和核销','回款预测、未认领和现金安排'],['坏账准备','预期信用损失入账','风险分层、模型假设和个别评估'],['历史坏账','核销和收回','账销案存、追索和责任复盘']],[2.3,6.5,6.9],9)
    head(doc,'1.2 业务闭环',2);flow(doc,['客户准入和信用调查','信用额度与账期审批','销售合同和订单','信用检查与发货放行','履约、验收和开票','收入及应收确认','客户对账和到期提醒','逾期催收与争议解决','银行到账和回款认领','应收核销','ECL计提或坏账核销','复盘信用政策并关闭Issue'])
    head(doc,'1.3 四类要求必须分开',2)
    table(doc,['类型','内容','例子'],[['会计准则','决定收入、应收、合同资产和预期信用损失的确认计量','不能用企业内部账龄政策替代准则判断'],['法律要求','合同权利、诉讼时效和证据效力','一般诉讼时效期间为三年，但要结合具体事实'],['常见实务','信用分级、催收会议、对账和冻结发货','行业习惯不等于强制规定'],['企业政策','额度、账期、审批层级和停发货阈值','应结合规模、风险和系统能力制定']],[2.6,6.0,7.1],8.8)

    head(doc,'第二章 岗位责任与跨部门协同')
    table(doc,['事项','执行人','财务主管','协同方','批准人'],D['raci'],[3.0,3.0,4.2,2.8,2.7],8)
    head(doc,'2.1 财务不能替业务承担的责任',2)
    bullets(doc,['销售对客户关系、商业承诺和催收行动负责。','项目或交付部门对履约、验收和质量事实负责。','法务对诉讼、仲裁和法律策略提供专业意见。','财务对数据、信用风险、账务、减值、监督和升级负责。'])
    head(doc,'2.2 业务部门推责任时怎么回应',2)
    para(doc,'“财务可以提供账龄、额度和风险影响，也可以组织跟踪，但客户为什么拒付、何时解决验收以及由谁催收，需要销售和项目负责人给出可验证的行动计划。请明确Owner和日期，财务按计划跟踪并向管理层报告。”')

    head(doc,'第三章 财务主管的工作节奏')
    table(doc,['频率','主管关注','亲自处理','会计执行后复核','升级事项'],[
        ['每日','大额到账、未认领、冻结客户和例外发货','重大放货例外、异常退款','回款核销和银行匹配','超额发货、客户重大负面信息'],['每周','逾期、承诺回款、争议和现金预测','Top风险客户会议','催收清单和承诺更新','反复失约、争议长期无Owner'],['每月','AR-GL、账龄、ECL、对账和管理报告','重大判断和主管复核','全量对账及底稿','重大减值、收入或列报问题'],['每季度','信用复评、集中度、账龄迁徙','重大客户复评','资料更新和指标分析','信用政策或客户结构恶化'],['每年','政策、核销、审计和制度','模型与重大核销审批','年度对账及档案','重大损失和责任追究']],[1.7,4.0,4.0,4.0,3.0],8)
    head(doc,'3.1 什么必须由主管亲自看',2)
    bullets(doc,['逾期90天以上、金额重大或客户经营恶化的余额。','冻结客户仍继续发货的全部例外。','争议款、诉讼款和个别评估客户。','应收贷方、长期未认领回款和跨客户核销。','ECL参数变化、坏账核销和以前年度核销款收回。'])

    head(doc,'第四章 单据 系统 数据和台账')
    head(doc,'4.1 单据链',2)
    table(doc,['资料','提供部门','取得时间','关键字段','勾稽对象','高风险缺陷'],[
        ['客户准入资料','销售/客户','合作前','主体、股东、地址、账户','工商与主数据','主体冒用或资料过期'],['信用审批单','销售/财务','额度启用前','额度、账期、有效期、审批人','ERP信用主数据','审批后系统未更新'],['销售合同','销售/法务','签订时','主体、货物、验收、付款、违约','订单、发货和发票','口头承诺改变账期'],['销售订单','销售','发货前','客户、物料、价格、账期','合同和出库','超额度绕过'],['出库与签收','仓库/物流','发货时','数量、批次、签收人','订单和收入','代签或缺少签收'],['验收资料','项目/客户','验收时','条件、日期、保留事项','合同和收入','内部验收替代客户验收'],['发票','税务/AR','开票时','主体、内容、金额、税率','收入和客户','主体或内容不符'],['银行流水','资金','到账日','户名、账号、金额、摘要','收款与核销','第三方付款不明'],['对账函','AR/销售','月或季度','期末余额和差异','明细账','未回函无替代证据'],['催收记录','销售/AR','持续','承诺金额、日期、联系人','账龄和预测','只写“客户走流程”']],[2.4,2.2,2.0,3.4,2.8,3.2],7.5)
    head(doc,'4.2 ERP数据流',2);flow(doc,['CRM客户与信用信息','销售合同和订单','ERP信用额度校验','出库和验收','SD或销售模块生成结算数据','AR生成客户明细','自动或接口凭证进入GL','银行流水进入资金模块','AR执行回款认领和核销','账龄 ECL和管理报表'])
    para(doc,'SAP、金蝶和用友的按钮和表名不同，但底层逻辑一致：主数据决定客户和账期，业务单据形成应收来源，清账或核销动作改变未清项状态，总账接收汇总或明细凭证。系统差异为零仍不能证明收入真实性和催收质量。',lead='系统逻辑：')
    head(doc,'4.3 台账字段',2)
    table(doc,['字段','为什么需要','维护人','频率'],[['客户编码','连接合同、订单、发票、收款和总账','主数据管理员','发生变化时'],['信用额度及有效期','控制风险敞口和过期额度','信用管理','审批后立即'],['合同账期和到期日','计算逾期和催收节点','AR/销售','逐笔'],['责任销售和区域','落实催收Owner并分析绩效','销售运营','人员变化时'],['争议类型和金额','区分信用问题与履约问题','销售/项目','每周'],['承诺回款金额和日期','形成现金预测和失约记录','销售','每周'],['最后催收日及证据','判断催收是否真实执行','AR','每次行动后'],['信用冻结状态','防止风险继续扩大','信用管理','实时'],['ECL分组和个别评估标识','支持减值计量','总账','月末'],['Issue状态和Closure','证明问题已解决','财务主管','持续']],[3.0,6.5,3.0,3.2],8.7)

    head(doc,'第五章 信用管理和销售放行')
    head(doc,'5.1 信用额度不是财务拍脑袋',2)
    bullets(doc,['客户外部资信、经营年限、股东和诉讼信息。','历史采购规模、付款及时性、争议和退货。','拟交易毛利、替代客户、担保和预付款比例。','最大风险敞口，包括已到期、未到期和已审批未发货订单。','额度有效期、复评频率和自动冻结条件。'])
    head(doc,'5.2 信用例外审批',2)
    table(doc,['条件','最低处理','不应接受的理由'],[['额度不足但无逾期','确认回款计划和新增订单金额，限次限期审批','“客户一直不错”'],['存在逾期','解释逾期原因、证据和新增风险敞口','“销售目标紧”'],['超过90天逾期','原则上冻结，重大商业理由由高层审批','“客户答应下周付”'],['诉讼或经营异常','法务和财务经理评估担保、预付或停止交易','“以后还有大单”']],[3.0,7.8,4.9],8.7)
    head(doc,'5.3 案例 客户逾期90天仍要求发货',2)
    para(doc,'北方自动化现有应收904万元中的90.4万元已逾期115天，另有22.6万元为冻结后的例外发货。销售提出再发货80万元。财务主管不能只回答“可以”或“不可以”。')
    flow(doc,['核对客户全部风险敞口','确认逾期是信用还是验收争议','取得客户付款计划和书面证据','计算新增订单后的最大敞口','提出预付款、部分回款或担保方案','按权限审批一次性例外','设置金额和失效日期','发货后跟踪并复盘'])

    head(doc,'第六章 到期管理 催收和争议')
    head(doc,'6.1 分层催收',2)
    table(doc,['阶段','动作','Owner','主管关注'],[['到期前7天','发送提醒并确认发票、验收和付款资料','销售/AR','是否存在流程障碍'],['逾期1至30天','电话和邮件确认付款日','销售','承诺是否具体'],['逾期31至90天','书面催款、限制新增额度','销售负责人','是否重复失约'],['逾期91至180天','冻结发货、专项会议、法务评估','销售总监/法务','减值和诉讼时效'],['超过180天','管理层决定诉讼、重组、担保或核销路径','财务经理/法务','个别ECL和责任追究']],[2.7,6.0,3.0,4.0],8.5)
    head(doc,'6.2 空泛催收记录为什么没有管理价值',2)
    para(doc,'“客户正在走流程”不能支持回款预测，也不能证明催收有效。合格记录至少包括客户联系人、付款审批节点、尚缺资料、承诺金额、承诺日期、下次行动、销售Owner和证据路径。')
    head(doc,'6.3 争议款分类',2)
    table(doc,['争议类型','先找谁','关键证据','可能处理'],[['未验收','项目/客户','合同条件、验收单、异议','继续履约、补验收或重评收入'],['质量问题','质量/售后','检测、退换货、索赔','维修、折让、退货或抗辩'],['价格差异','销售','报价、合同、变更','补充协议或纠错'],['发票问题','税务/AR','开票申请和客户要求','红冲重开或解释'],['客户资金困难','销售/法务','财务状况、付款计划','担保、分期、诉讼和ECL'],['无依据拖延','销售/法务','履约证据和催款','停止供货并法律追索']],[2.3,2.4,5.4,5.6],8.4)

    head(doc,'第七章 回款认领 核销和对账')
    head(doc,'7.1 收到钱不等于应收已经关闭',2)
    flow(doc,['银行确认到账','提取付款户名 金额 日期 摘要','匹配客户主数据','匹配发票和应收未清项','识别一对多 多对一和第三方付款','低置信度转销售或客户确认','AR执行核销','主管复核重大和跨客户核销','未认领款列入月结Issue'])
    head(doc,'7.2 不允许直接自动核销的场景',2)
    bullets(doc,['同金额存在多个客户或多个未清项目。','付款主体与合同客户不一致且没有可靠说明。','客户存在折扣、返利、退货、索赔或手续费扣减。','回款金额包含预付款和历史欠款。','销售或客户指定用途与系统建议不一致。'])
    head(doc,'7.3 客户对账',2)
    para(doc,'对账不是把财务账发给客户。应先内部核对订单、出库、验收、发票、回款和退款，再向客户确认余额及差异。重大客户不回函时，需要销售参与，并利用后续回款、门户记录、签收验收和往来邮件形成替代证据。')

    head(doc,'第八章 会计核算与预期信用损失')
    head(doc,'8.1 核算链条',2)
    table(doc,['节点','典型分录','管理含义'],[['满足收入确认条件','借：应收账款；贷：主营业务收入、应交税费','确认收入不等于客户信用风险消失'],['收到客户款','借：银行存款；贷：应收账款','需正确认领和核销'],['计提ECL','借：信用减值损失；贷：坏账准备','反映预计现金短缺'],['核销坏账','借：坏账准备；贷：应收账款','需审批并账销案存'],['核销后收回','按适用会计政策恢复或直接确认收回并记录银行款','不得形成账外资金']],[2.8,6.8,6.1],8.5)
    head(doc,'8.2 应收账款 合同资产和合同负债',2)
    table(doc,['项目','核心判断','常见错误'],[['应收账款','收取对价的权利仅取决于时间流逝','未达到无条件收款权就转应收'],['合同资产','已履约但收款权还取决于其他条件','与普通应收混用导致账龄错误'],['合同负债','在转让商品前已收或已到期应收客户对价','看到应收贷方就机械重分类']],[2.8,7.0,5.9],9)
    head(doc,'8.3 ECL不能只套账龄比例',2)
    para(doc,'账龄可以作为组合计量的重要基础，但财务还应考虑历史损失率、客户类型、账龄迁徙、当前经营状况、前瞻信息、担保以及个别客户重大不利事件。客户破产时，即使账龄只有60天，也可能需要个别评估。')
    table(doc,['风险层级','起始天数','截止天数','案例率','使用依据'],D['ecl_rates'],[2.4,2.0,2.0,2.0,7.3],8.5)
    para(doc,'上述比例仅为模拟企业练习参数，不代表准则、税法或行业统一比例。企业应根据自身数据和合理证据制定，并按权限审批。',lead='重要说明：')

    head(doc,'第九章 核心指标和深层原因')
    table(doc,['指标','定义','为什么看','恶化原因','向下追查'],D['metrics'],[2.5,3.0,3.5,4.0,4.2],7.8)
    head(doc,'9.1 应收增加的原因树',2)
    bullets(doc,['规模：赊销收入增长是否足以解释余额增长。','政策：账期、额度或例外是否放宽。','结构：大客户、项目客户或高风险客户占比是否增加。','执行：催收、验收、发票和回款认领是否变慢。','质量：退货、争议、提前确认收入或客户经营困难。','数据：串户、重复确认、未核销和截止错误。'])
    head(doc,'9.2 指标不能机械比较',2)
    para(doc,'项目型国企客户可能合同账期长但违约概率较低；经销商账期短却可能风险集中。DSO和逾期率必须结合合同条款、业务季节性、客户结构和收入真实性分析，不能用单一阈值评价所有客户。')

    head(doc,'第十章 高频异常案例库')
    for c in D['cases']:case(doc,c)

    head(doc,'第十一章 内部控制 舞弊和红旗信号')
    table(doc,['控制','主要风险','控制动作','证据'],D['controls'],[3.2,3.8,5.7,3.0],8.2)
    head(doc,'11.1 红旗信号',2)
    bullets(doc,['月末集中新增应收、次月退货或冲红。','客户额度频繁被同一人员例外放开。','第三方付款多、退款账户与原付款账户不同。','客户长期不对账，但销售坚持余额没有问题。','催收记录长期只有“客户走流程”。','核销坏账后仍有回款，但未进入财务账。','应收贷方和跨客户核销频繁出现。','销售人员既能创建客户又能修改账期或发货状态。'])

    head(doc,'第十二章 新财务主管接手检查')
    head(doc,'12.1 第一周取得什么',2)
    table(doc,['资料','检查目标','高风险信号'],[['客户主数据和信用权限','谁能创建、修改和审批','销售可自行改额度'],['最近12个月AR明细和GL','模块总账是否一致','长期差异无解释'],['账龄及逾期清单','识别Top风险','账龄按发票日而非到期日'],['信用政策和例外日志','制度是否真正执行','例外没有期限'],['客户对账和催收记录','证据是否可用','重大客户长期未对账'],['ECL底稿和核销台账','模型及审批','只套比例或老板口头决定'],['争议和诉讼清单','财务是否获得风险信息','法务与财务数据不一致'],['银行未认领款','核销是否及时','大额长期挂账']],[3.5,6.0,6.2],8.5)
    head(doc,'12.2 第一月怎么判断管理好不好',2)
    bullets(doc,['抽取前十大客户，从合同追到发货、验收、应收、回款和核销。','重算账龄并检查到期日来源。','检查逾期90天以上是否逐笔有Owner、证据和日期。','抽查冻结客户是否仍能发货。','复核ECL参数、个别评估和管理层审批。','比较销售回款承诺与实际到账，评价预测可信度。'])
    head(doc,'12.3 保护责任边界',2)
    para(doc,'接手时应保存基准日余额、历史问题清单、已取得资料、无法验证事项及向上汇报记录。不要在证据不足时签署“余额全部正确”或替历史责任人补做虚假对账。对重大问题提出临时控制、Owner和期限，并持续记录Closure。')

    head(doc,'第十三章 月末Checklist')
    table(doc,['明确动作','责任人','期限','证据','关闭标准'],D['checklist'],[5.2,2.2,1.4,4.2,4.0],8)

    head(doc,'第十四章 财务经理和管理层视角')
    head(doc,'14.1 一页管理报告只保留重要结论',2)
    table(doc,['区域','管理层要看什么','表达方式'],[['规模','应收余额、赊销收入和环比变化','余额增加多少，其中多少由收入增长解释'],['质量','逾期率、90天以上、争议和集中度','风险集中在哪些客户和原因'],['现金','本月回款、下月预测和预测准确率','现金缺口及承诺客户'],['风险','冻结客户、额度例外、诉讼和ECL','最坏情形及报表影响'],['行动','Top Issue、Owner和截止','需要管理层决定什么']],[2.0,6.4,6.8],8.8)
    head(doc,'14.2 管理层决策问题',2)
    bullets(doc,['继续增长是否需要放宽信用，能承受多大现金占用和损失。','大客户高毛利但账期长，回报是否覆盖资金成本和风险。','现金紧张时，催收、保理、信用保险和停止发货如何组合。','逾期问题来自客户、销售激励还是交付质量，应该改变哪项机制。'])

    head(doc,'第十五章 20%核心知识')
    table(doc,['核心能力','达到的工作标准'],[['客户信用','知道风险敞口，不只看账面应收'],['账龄','能解释到期日来源，区分逾期和未到期'],['催收','每笔重大逾期有Owner、行动和证据'],['争议','能区分信用、验收、质量、价格和发票问题'],['核销','到账、认领、核销和银行勾稽一致'],['ECL','组合参数有证据，重大客户做个别评估'],['系统','理解CRM、销售、AR、银行和GL数据流'],['控制','额度、发货、收款、核销和核销坏账职责分离'],['主管管理','用Issue和Closure推动跨部门解决'],['管理报告','把余额压缩为现金、风险、原因和行动']],[4.0,11.7],9)

    head(doc,'附录A 贯穿案例客户与应收明细')
    table(doc,['客户','类型','信用级别','额度','账期','销售','状态'],[[x[1],x[2],x[3],x[4],x[5],x[6],x[8]] for x in D['customers']],[3.5,2.0,1.6,2.0,1.4,1.8,2.2],7.8)
    table(doc,['应收ID','客户','到期日','金额','状态'],[[x[0],x[2],x[4],x[7],x[10]] for x in D['ar']],[2.2,4.0,2.4,2.5,4.6],8)
    head(doc,'附录B 主要依据和使用边界')
    table(doc,['法规或准则','文件号','与本专题关系','官方来源'],D['sources'],[4.3,3.0,4.4,4.0],7.2)
    para(doc,'财政部内部控制应用指引主要适用于执行企业内部控制规范体系的企业；其他企业可结合规模和风险参考其控制思想。民法典诉讼时效的具体起算、中断、中止和争议处理应由法务结合个案判断。')
    out=OUT/'应收账款管理实战手册.docx';OUT.mkdir(parents=True,exist_ok=True);doc.save(out);return out

def topic(title,children=None):
    x={'id':title,'title':title}
    if children:x['children']={'attached':[topic(a,b if isinstance(b,list) else None) if isinstance(a,str) else topic(str(a)) for a,b in children]}
    return x
def build_xmind():
    sheets=[]
    full=[('业务流程',[(x,[]) for x in ['客户准入','信用审批','合同订单','发货验收','应收确认','催收争议','回款核销','ECL核销']]),('主管工作',[(x,[]) for x in ['每日','每周','月结','季度复评','年度政策']]),('工具包',[(x,[]) for x in ['客户主数据','应收明细','催收计划','争议台账','ECL','Checklist','Issue Log']]),('风险',[(x,[]) for x in ['超额发货','空泛催收','错配回款','收入截止','舞弊','坏账']])]
    sheets.append({'id':'s1','class':'sheet','title':'应收全景','rootTopic':topic('应收账款管理',full)})
    sheets.append({'id':'s2','class':'sheet','title':'异常处理','rootTopic':topic('应收异常闭环',[('发现',[]),('核实',[]),('证据',[]),('会计处理',[]),('管理整改',[]),('升级',[]),('Closure',[])])})
    sheets.append({'id':'s3','class':'sheet','title':'新主管接手','rootTopic':topic('新财务主管接手',[('第一周资料',[]),('第一月抽查',[]),('Top风险客户',[]),('系统权限',[]),('历史问题',[]),('责任边界',[])])})
    p=OUT/'应收账款管理全景.xmind'
    with ZipFile(p,'w',ZIP_DEFLATED) as z:
        z.writestr('content.json',json.dumps(sheets,ensure_ascii=False));z.writestr('metadata.json',json.dumps({'creator':{'name':'Codex'},'activeSheetId':'s1'},ensure_ascii=False));z.writestr('manifest.json',json.dumps({'file-entries':{'content.json':{},'metadata.json':{}}}))
    return p
if __name__=='__main__':print(build());print(build_xmind())
