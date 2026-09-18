from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\Users\HZJ\Desktop\Git\work\齐昊老师_学员实战案例_第1集_文字思维导图.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "203040"
MUTED = "64748B"
LIGHT_BLUE = "E8F1F8"
LIGHT_GRAY = "F3F5F7"
GOLD = "B7791F"
GREEN = "2F6B4F"
RED = "9B1C1C"
WHITE = "FFFFFF"
LATIN_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"


def set_run(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = LATIN_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_inches):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    dxa = [int(round(w * 1440)) for w in widths_inches]
    for width in dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[i])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(dxa[i]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run(run, 9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run(run2, 9, color=MUTED)


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    p.paragraph_format.space_after = Pt(10)
    set_run(p.add_run(text), 27, True, DARK_BLUE)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(20)
        set_run(p2.add_run(subtitle), 15, False, BLUE)


def add_label_line(doc, label, value, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(label + "："), 10.5, True, MUTED)
    set_run(p.add_run(value), 10.5, False, INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), {1: 16, 2: 13, 3: 12}[level], True, {1: BLUE, 2: BLUE, 3: DARK_BLUE}[level])
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        set_run(p.add_run(bold_prefix), 11, True, INK)
        set_run(p.add_run(text[len(bold_prefix):]), 11, False, INK)
    else:
        set_run(p.add_run(text), 11, False, INK)
    return p


def add_bullet(doc, text, level=0, color=INK):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text), 11, False, color)
    return p


def add_number(doc, text, level=0):
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text), 11, False, INK)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, label_color=BLUE):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), label_color)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.06)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.2
    set_run(p.add_run(label + "  "), 10.5, True, label_color)
    set_run(p.add_run(text), 10.5, False, INK)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = LATIN_FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    style = doc.styles[style_name]
    style.font.name = LATIN_FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run(hp.add_run("齐昊老师 · 学员一手真实案例分享"), 9, True, MUTED)
footer = section.footer
add_page_number(footer.paragraphs[0])

add_title(doc, "学员实战案例 第 1 集", "亏损企业如何实现业绩 5 倍增长，并成长为省级龙头")
add_label_line(doc, "主讲账号", "齐昊老师")
add_label_line(doc, "案例学员", "囍桐（空降财务负责人）")
add_label_line(doc, "案例行业", "啤酒经销商")
add_label_line(doc, "视频时长", "16 分 41 秒")
add_label_line(doc, "发布时间", "2024 年 7 月 10 日")

doc.add_paragraph().paragraph_format.space_after = Pt(22)
add_callout(doc, "一句话结论", "真正的财务管理不是把账做得更辛苦，而是通过统一系统、重构流程、明确责任、精细核算和绩效驱动，让经营数据能支持决策并推动业务增长。", "EAF3EE", GREEN)
add_callout(doc, "阅读说明", "“视频内容”来自抖音页面的智能文稿；“整理提炼”是为财务经理/财务总监学习而增加的结构化总结，不等同于老师逐字原话。", LIGHT_GRAY, MUTED)

doc.add_page_break()
add_heading(doc, "一、文字思维导图（总览）", 1)
add_body(doc, "中心主题：亏损啤酒经销商的财务管理改革")

mindmap = [
    (0, "1. 初始困境：持续亏损 + 数据混乱 + 决策失灵"),
    (1, "财务长期加班仍频繁出错；老板无法据此制定绩效奖金。"),
    (1, "6 套系统割裂，名称与口径不统一；业务靠微信报单、财务重复录入。"),
    (0, "2. 第一抓手：统一业务与财务口径"),
    (1, "停用旧系统，上线统一手机 APP，业务下单、配送、监控、数据推送一体化。"),
    (1, "责任归位：谁发生业务，谁负责录入；财务由录单转向审核和管理。"),
    (0, "3. 落地难点：系统透明化触碰既得利益"),
    (1, "前半年将系统录入率与绩效强绑定；低于 70% 绩效归零。"),
    (1, "归零绩效即时奖励给录入率超过 90% 的前五名；半年后再调整权重。"),
    (1, "改革必须取得老板明确、持续的支持。"),
    (0, "4. 上线基础：数据治理与制度重建"),
    (1, "连续两个月亲自盘点，SKU 从 2,000 个压缩至 635 个。"),
    (1, "同步更新仓储、物流等制度并宣讲，确保制度与外部法规一致。"),
    (0, "5. 核算创新：把包装物作为管理对象"),
    (1, "将酒水、瓶子、瓶盖、箱子拆分核算；系统跟踪客户应退包装物。"),
    (1, "用回瓶率识别异常客户、调整回收价，并制止二批商私下收购。"),
    (1, "赠酒照样收瓶箱押金，兼顾促销吸引力与资产回收。"),
    (0, "6. 风险控制：应收账款三维机制"),
    (1, "设置最长账期、最大欠款金额；超限自动预警。"),
    (1, "上一笔未结清最多再开 3 单，超额锁单；借客户投诉倒逼回款入账。"),
    (0, "7. 增长机制：客户绑定 + 绩效改革"),
    (1, "包量合同：3 个月完成 8 万元销量，每箱奖励 2 元；提前完成奖励 3 元。"),
    (1, "奖励折算为下一批订单赠酒，并继续收包装物押金。"),
    (1, "绩效加入回款、新客户、拜访频率、高毛利占比、专销、利润分红等指标。"),
    (1, "扣除坏账、掉店和风险金，把灰色收入转化为透明激励。"),
    (0, "8. 最终结果：亏损企业走向省级龙头，业绩约增长 5 倍"),
]
for level, text in mindmap:
    add_bullet(doc, text, level=level, color=DARK_BLUE if level == 0 else INK)

add_callout(doc, "主线", "数据统一 → 流程重构 → 责任归位 → 核算精细 → 风险受控 → 绩效驱动 → 经营改善", "FFF7E8", GOLD)

doc.add_page_break()
add_heading(doc, "二、案例背景与问题诊断", 1)
add_heading(doc, "1. 学员与企业背景", 2)
add_body(doc, "【视频内容】2021 年，囍桐从医疗器械行业跨到一家啤酒经销商担任财务负责人。公司当时持续亏损，财务数据混乱，财务人员经常加班却仍不断出错。老板无法依靠数据制定绩效奖金，也难以作出经营决策。")
add_body(doc, "【整理提炼】她面对的并不是一个单纯的“账务问题”，而是系统、流程、责任、资产、信用和激励机制共同失灵的经营管理问题。")

add_heading(doc, "2. 跨行业是否构成障碍", 2)
add_body(doc, "【视频内容】学员认为“财务管理无行业区分”，因此把跨行业挑战当成机会。")
add_body(doc, "【整理提炼】会计科目和业务细节会随行业变化，但财务管理的底层逻辑相对稳定：确认业务事实、统一数据口径、控制风险、配置资源、建立激励并用数据支持经营。")

add_heading(doc, "3. 根因树", 2)
root_table = doc.add_table(rows=1, cols=3)
root_table.style = "Table Grid"
headers = ["表面症状", "直接原因", "管理根因"]
for i, h in enumerate(headers):
    shade_cell(root_table.rows[0].cells[i], LIGHT_BLUE)
    p = root_table.rows[0].cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(h), 10.5, True, DARK_BLUE)
rows = [
    ("财务加班仍出错", "业务微信报单，财务二次录入", "责任不在业务源头；流程设计错误"),
    ("数据对不上", "6 套系统、名称和口径不一致", "缺乏统一数据标准与主系统"),
    ("公司持续亏损", "SKU 冗余、包装物流失、坏账与激励失真", "财务未深入业务链条实施控制"),
    ("绩效无法制定", "经营数据不可信、指标单一", "数据治理与绩效机制脱节"),
]
for row in rows:
    cells = root_table.add_row().cells
    for i, value in enumerate(row):
        p = cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(value), 10, False, INK)
set_repeat_table_header(root_table.rows[0])
set_table_widths(root_table, [1.65, 2.25, 2.6])

add_heading(doc, "三、改革路径详解", 1)
add_heading(doc, "1. 系统整合：先统一业务和财务口径", 2)
add_body(doc, "【视频内容】企业原有业务考勤、仓储配送、财务、线下商城等 6 套独立系统，维护主体分散，同一个客户在不同系统里的名称也不一致。学员推动停用旧系统，使用统一手机 APP。")
add_bullet(doc, "业务人员在 APP 直接下单，并同步选择司机、生成配送单。")
add_bullet(doc, "系统记录车辆轨迹，停留超过 15 分钟可自动报警。")
add_bullet(doc, "销售人员可看到附近商户采购记录、应收账款、复购率以及商品成本和定价变化。")
add_bullet(doc, "数据录入责任由财务转回业务：谁做业务谁录入，减少二次录入和责任推诿。")
add_callout(doc, "财务总监视角", "信息化的第一目标不是“上一个系统”，而是确定唯一数据源、统一主数据和口径，并把责任固化到业务发生的源头。")

add_heading(doc, "2. 变革管理：识别“不会用”背后的利益冲突", 2)
add_body(doc, "【视频内容】销售抵触新系统，表面理由是工作忙，实质是透明化会压缩订单与收入中的灰色空间。")
add_bullet(doc, "前半年提高“系统录入率”在绩效中的权重。")
add_bullet(doc, "录入率低于 70% 的人员，相关绩效归零。")
add_bullet(doc, "归零部分即时奖励给录入率超过 90% 的前五名，形成强烈示范效应。")
add_bullet(doc, "半年后习惯形成，再降低录入率权重，增加销售额、复购率等经营指标。")
add_bullet(doc, "关键改革由老板拍板，为跨部门冲突提供权威支持。")
add_callout(doc, "整理提炼", "制度能否落地，取决于利益机制是否同步改变。培训解决“不会”，绩效和奖惩解决“不愿”。")

add_heading(doc, "3. 数据治理：系统上线前先把基础数据做干净", 2)
add_body(doc, "【视频内容】学员连续两个月亲自盘点，将 SKU 从约 2,000 个精简至 635 个，淘汰滞销品；同时更新仓储、物流等制度并上墙宣讲。")
add_bullet(doc, "盘点不仅核对数量，还要确认品名、规格、包装物、库龄和可售状态。")
add_bullet(doc, "SKU 精简降低了库存资金占用、盘点复杂度和系统维护成本。")
add_bullet(doc, "制度更新还要关注《会计法》《公司法》等外部法规变化，避免内部规则失效。")

add_heading(doc, "4. 核算创新：将包装物拆开管理", 2)
add_body(doc, "【视频内容】啤酒业务不能只核算“整箱酒”，而要将酒水、瓶子、瓶盖、箱子拆分核算。系统自动记录客户应退包装物数量，并计算回瓶率，目标约为 70%—80%。")
add_bullet(doc, "低回瓶率客户：分析原因，必要时提高回收价。")
add_bullet(doc, "二批商私下收购公司包装物：通过停货约束制止，减少资产流失。")
add_bullet(doc, "开业赠酒或买十送一：酒可以赠，但瓶子和箱子必须收押金。")
add_callout(doc, "整理提炼", "财务核算对象决定管理视野。把包装物从商品中拆出来，才可能看见回收率、周转效率、流失责任和真实促销成本。")

add_heading(doc, "5. 应收账款：账期、额度、订单三维联控", 2)
add_body(doc, "【视频内容】每个客户都设置最长账期和最大欠款金额，超限预警；上一笔应收未结清时，销售员最多还能开 3 单，超过后系统锁单。")
add_bullet(doc, "账期控制时间风险。")
add_bullet(doc, "信用额度控制金额风险。")
add_bullet(doc, "锁单控制新增敞口，并借客户无法下单后的主动投诉倒逼销售及时上交回款。")
add_callout(doc, "财务总监视角", "应收管理不能只靠月末催款，应嵌入订单环节，形成“交易前授信、交易中预警、超限锁单、回款解锁”的闭环。")

add_heading(doc, "6. 客户绑定：用包量合同设计复购", 2)
add_body(doc, "【视频内容】客户在 3 个月内完成 8 万元销量，每箱奖励 2 元；提前完成则奖励 3 元。奖励不直接返现，而折算成下一批订单的赠酒，且赠酒继续收取瓶箱押金。")
add_bullet(doc, "奖励与未来订单绑定，提升复购和客户黏性。")
add_bullet(doc, "奖励使用赠酒形式，兼顾客户感知与企业现金流。")
add_bullet(doc, "押金机制确保促销不造成包装物资产失控。")

add_heading(doc, "7. 绩效改革：把灰色利益转成透明收入", 2)
add_body(doc, "【视频内容】新绩效由基本工资、绩效工资、提成、阶段性奖励和扣除项目组成。")
perf_table = doc.add_table(rows=1, cols=2)
perf_table.style = "Table Grid"
for i, h in enumerate(("奖励维度", "扣除与约束")):
    shade_cell(perf_table.rows[0].cells[i], LIGHT_BLUE)
    p = perf_table.rows[0].cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(h), 10.5, True, DARK_BLUE)
perf_rows = [
    ("回款比率", "坏账"),
    ("新客户开发", "掉店（客户流失）"),
    ("老客户拜访频率", "风险金扣除"),
    ("高毛利产品占比", "系统录入与业务合规"),
    ("归口专销、年终利润分红", "异常订单与资产流失责任"),
]
for left, right in perf_rows:
    cells = perf_table.add_row().cells
    set_run(cells[0].paragraphs[0].add_run(left), 10, False, INK)
    set_run(cells[1].paragraphs[0].add_run(right), 10, False, INK)
set_repeat_table_header(perf_table.rows[0])
set_table_widths(perf_table, [3.25, 3.25])
add_body(doc, "【整理提炼】绩效指标从“只看销量”升级为“销量、回款、毛利、客户质量、过程行为和风险结果”并重，使业务人员能够通过合规经营获得更高、可持续的收入。")

add_heading(doc, "四、改革成效与因果链", 1)
effect_table = doc.add_table(rows=1, cols=3)
effect_table.style = "Table Grid"
for i, h in enumerate(("改革动作", "直接变化", "经营价值")):
    shade_cell(effect_table.rows[0].cells[i], LIGHT_BLUE)
    p = effect_table.rows[0].cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(h), 10.5, True, DARK_BLUE)
effect_rows = [
    ("统一系统与源头录入", "数据口径一致、差错减少", "老板获得可信经营数据"),
    ("盘点和 SKU 精简", "库存更真实、品类更聚焦", "降低占用与管理复杂度"),
    ("包装物分项核算", "回瓶率可量化、流失可追责", "减少资产损失、提升周转"),
    ("应收三维控制", "欠款超限受控、回款更及时", "降低坏账与销售截留风险"),
    ("客户与绩效机制重构", "复购、毛利和行为目标一致", "形成增长与风险兼顾的激励"),
]
for row in effect_rows:
    cells = effect_table.add_row().cells
    for i, value in enumerate(row):
        set_run(cells[i].paragraphs[0].add_run(value), 9.6, False, INK)
set_repeat_table_header(effect_table.rows[0])
set_table_widths(effect_table, [2.0, 2.15, 2.35])
add_callout(doc, "视频结论", "在这些改革推动下，企业由亏损走向省级龙头经销商，视频标题概括为业绩增长约 5 倍。", "EAF3EE", GREEN)

add_heading(doc, "五、对财务经理 / 财务总监的能力启示", 1)
capabilities = [
    ("业务理解力", "能走进仓库、配送、销售和客户现场，理解数据是怎样产生的。"),
    ("系统与数据治理", "能定义主数据、统一口径、设计源头录入和责任机制。"),
    ("流程与内控设计", "把控制点嵌入下单、配送、回款、促销和绩效流程。"),
    ("管理会计能力", "通过拆分核算对象，让资产周转、客户质量和真实利润可见。"),
    ("变革推动力", "识别利益冲突，设计奖惩节奏，并争取老板持续授权。"),
    ("经营结果意识", "财务工作的终点不是报表完成，而是经营改善和企业价值提升。"),
]
for name, detail in capabilities:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(name + "："), 11, True, DARK_BLUE)
    set_run(p.add_run(detail), 11, False, INK)

add_heading(doc, "六、可复用行动清单", 1)
checklist = [
    "列出企业现有系统、台账和报表，找出重复录入及口径冲突。",
    "选取一个业务闭环，画出从接单到回款的流程并标注数据责任人。",
    "检查系统上线前的客户、商品、仓库、供应商等主数据是否准确。",
    "识别行业中被笼统核算、但实际需要拆分管理的资产或成本对象。",
    "为客户建立账期、额度、订单状态和回款动作的联控规则。",
    "审查绩效是否同时覆盖收入、回款、毛利、客户质量、过程行为与风险。",
    "为重大改革明确老板授权、试运行周期、奖惩规则和复盘节点。",
]
for item in checklist:
    add_number(doc, item)

add_callout(doc, "建议的学习方式", "看完本集后，不只记住制度和指标，更要复盘学员的行动顺序：先观察与盘点，再统一数据和责任，随后重构核算与控制，最后用绩效把新流程稳定下来。", "FFF7E8", GOLD)

add_heading(doc, "七、来源与准确性说明", 1)
add_body(doc, "视频来源：抖音账号“齐昊老师”，“学员一手真实案例分享”合集第 1 集，视频编号 7389889299036097819。")
add_body(doc, "内容依据：抖音页面公开展示的标题、时长、发布时间及“智能文稿”。本文进行了结构化改写与学习提炼，并非逐字稿。涉及具体阈值、金额和效果的表述均按智能文稿保留；如用于企业实际制度设计，应结合公司业务、法律法规和系统条件重新论证。")

doc.core_properties.title = "齐昊老师学员实战案例第1集：文字思维导图"
doc.core_properties.subject = "财务管理与财务总监实战案例整理"
doc.core_properties.author = "Codex"
doc.core_properties.keywords = "财务管理, 财务经理, 财务总监, 业财融合, 内控, 应收账款, 绩效管理"
doc.save(OUT)
print(OUT)
