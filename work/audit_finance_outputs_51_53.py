import json
import zipfile
from pathlib import Path

from docx import Document
from finance_cases_data_51_53 import CASES

ROOT = Path(r"C:\Users\HZJ\Desktop\Git\work")


def walk(node):
    yield node
    for child in node.get("children", {}).get("attached", []):
        yield from walk(child)


for case in CASES:
    episode = case["episode"]
    stem = f"齐昊老师_学员实战案例_第{episode}集"
    docx_path = ROOT / f"{stem}_文字思维导图.docx"
    xmind_path = ROOT / f"{stem}_思维导图.xmind"
    with zipfile.ZipFile(docx_path) as z:
        assert z.testzip() is None
        assert "word/document.xml" in z.namelist()
    doc = Document(docx_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert case["title"] in text
    assert "不是逐字稿" in text
    assert len(doc.tables) == 1
    assert len([p for p in doc.paragraphs if p.style.name.startswith("Heading")]) >= len(case["sections"])
    for video_id in (str(case["video_id"]).split(";") if ";" in str(case["video_id"]) else [str(case["video_id"])]):
        assert video_id.strip() in text
    with zipfile.ZipFile(xmind_path) as z:
        assert z.testzip() is None
        sheets = json.loads(z.read("content.json"))
        assert [s["title"] for s in sheets] == ["01 总览图", "02 详细知识树"]
        topics = [list(walk(s["rootTopic"])) for s in sheets]
        ids = [t["id"] for group in topics for t in group]
        assert len(ids) == len(set(ids))
        assert all(case["title"] in s["rootTopic"]["title"] for s in sheets)
    print(episode, "OK", "paragraphs", len(doc.paragraphs), "topics", [len(t) for t in topics])
