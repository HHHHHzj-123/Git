from pathlib import Path
import json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent; OUT=ROOT/'deliverables'; IMG=ROOT/'images'; IMG.mkdir(exist_ok=True)
D=json.loads((ROOT/'content.json').read_text(encoding='utf-8'))
NAVY='17365D'; BLUE='2F75B5'; PALE='F3F6F9'; AMBER='FFF2CC'; GREEN='E2F0D9'; RED='FCE4D6'

def ifont(sz,bold=False):
    for p in [Path(r'C:\Windows\Fonts\msyh.ttc'),Path(r'C:\Windows\Fonts\simhei.ttf')]:
        if p.exists(): return ImageFont.truetype(str(p),sz,index=0)
    return ImageFont.load_default()
def flow(name,title,levels):
    w=1500; m=70; bh=100; gap=45; h=150+len(levels)*(bh+gap)
    im=Image.new('RGB',(w,h),'white'); dr=ImageDraw.Draw(im); dr.text((m,30),title,font=ifont(38),fill='black')
    y=110
    for ri,row in enumerate(levels):
        bw=(w-2*m-(len(row)-1)*28)/len(row)
        for i,s in enumerate(row):
            x=m+i*(bw+28); fill=(232,242,250) if ri%2==0 else (244,247,249)
            dr.rounded_rectangle((x,y,x+bw,y+bh),14,fill=fill,outline=(47,117,181),width=3)
            lines=[]; cur=''
            for ch in s:
                if dr.textlength(cur+ch,font=ifont(22))>bw-24: lines.append(cur); cur=ch
                else: cur+=ch
            if cur: lines.append(cur)
            for j,t in enumerate(lines[:3]):
                tw=dr.textlength(t,font=ifont(22)); dr.text((x+(bw-tw)/2,y+25+j*28),t,font=ifont(22),fill='black')
            if i<len(row)-1:
                ax=x+bw+4; ay=y+bh/2; dr.line((ax,ay,ax+17,ay),fill=(90,90,90),width=3); dr.polygon([(ax+17,ay),(ax+9,ay-6),(ax+9,ay+6)],fill=(90,90,90))
        if ri<len(levels)-1:
            cx=w/2; dr.line((cx,y+bh+3,cx,y+bh+28),fill=(90,90,90),width=3); dr.polygon([(cx,y+bh+28),(cx-7,y+bh+18),(cx+7,y+bh+18)],fill=(90,90,90))
        y+=bh+gap
    im.save(IMG/name,dpi=(180,180))

flow('panorama.png','企业财务资料全景',[['业务单据','系统数据','会计凭证','财务账簿'],['管理台账','月结底稿','税务资料','报表资料'],['合同与法规','审计PBC','档案索引','岗位交接'],['形成','复核','归档','调阅','到期鉴定']])
flow('decision.png','是否另建Excel台账的判断',[['ERP是否完整记录对象和状态'],['能否按责任人、到期日和异常筛选'],['是否需要跨系统勾稽或留复核证据'],['满足：使用ERP报表并保存参数','不满足：建轻量Excel补充台账'],['明确唯一负责人、更新频率和停用条件']])
flow('close.png','月结资料包形成路径',[['T-5通知截止与资料需求'],['T-3催收业务单据、合同和人员变动'],['T-1接口、库存、银行和发票预检查'],['T+1子模块结账与对账'],['T+2调整、税务测算和报表校验'],['主管复核、锁定版本、归档索引']])
flow('handover.png','财务岗位交接SOP',[['确定交接范围和截止日'],['系统权限、银行U盾、税务身份盘点'],['台账、底稿、合同和未决事项逐项演示'],['抽样重做关键月结任务'],['问题清单、责任人和完成期限'],['三方签字、权限回收、后续观察期']])
flow('archive.png','会计资料生命周期',[['形成与接收'],['真实性完整性检查'],['分类命名与版本锁定'],['月结包/年度包归档'],['权限控制与调阅登记'],['保管期满鉴定'],['销毁审批或继续保管']])

doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(2.1); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2.1); sec.right_margin=Cm(2.0)
for st,sz in [('Normal',10),('Title',26),('Heading 1',18),('Heading 2',14),('Heading 3',11.5)]:
    s=doc.styles[st]; s.font.name='Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑'); s.font.size=Pt(sz)
    if st!='Normal': s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].paragraph_format.line_spacing=1.25; doc.styles['Normal'].paragraph_format.space_after=Pt(4)
for s in ['Heading 1','Heading 2','Heading 3']:
    doc.styles[s].paragraph_format.keep_with_next=True

def shade(c,color):
    pr=c._tc.get_or_add_tcPr(); e=pr.find(qn('w:shd'))
    if e is None: e=OxmlElement('w:shd'); pr.append(e)
    e.set(qn('w:fill'),color)
def table(headers,rows,fs=7.7,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); shade(c,NAVY); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(fs)
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text='' if v is None else str(v); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cs[i],PALE)
            for p in cs[i].paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
                for r in p.runs: r.font.size=Pt(fs)
    if widths:
        for rr in t.rows:
            for i,w in enumerate(widths): rr.cells[i].width=Cm(w)
    pr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); pr.append(rep)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def para(s,lead=None):
    p=doc.add_paragraph();
    if lead: p.add_run(lead).bold=True
    p.add_run(s); return p
def bullets(items):
    for x in items: doc.add_paragraph(x,style='List Bullet')
def pic(name,caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(IMG/name),width=Cm(16.4))
    p=doc.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.font.size=Pt(8); r.font.italic=True; r.font.color.rgb=RGBColor(89,89,89)
def page(): doc.add_page_break()
def toc():
    p=doc.add_paragraph(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u'); p._p.append(fld)

for section in doc.sections:
    h=section.header.paragraphs[0]; h.text='企业财务资料、台账、底稿、档案与岗位交接实操手册'; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in h.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100)
    f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)

doc.add_paragraph().paragraph_format.space_after=Pt(75)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('企业财务资料、台账、底稿、档案与岗位交接实操手册')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('从资料形成、月结复核到可验证交接'); r.bold=True; r.font.size=Pt(17)
doc.add_paragraph().paragraph_format.space_after=Pt(70)
for s in ['面向岗位：总账会计、财务主管','案例企业：启辰智能硬件（苏州）有限公司','版本日期：2026年9月10日','适用范围：中国大陆制造业及硬件科技企业']:
    p=doc.add_paragraph(s); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
page(); doc.add_heading('使用说明',0)
para('本手册解决三个实际问题：资料从哪里来、月结时如何成为可复核证据、换人时怎样把隐性经验变成可验证的交接。表格中的保存年限仅在《会计档案管理办法》明确规定时标注法定期限；合同、经营资料和非会计档案需结合其他法规及企业制度确定。')
para('Excel台账不是越多越好。成熟ERP已能完整记录对象、状态、责任人和变更日志时，应优先使用系统报表；只有跨系统勾稽、异常跟踪、监管取证或管理闭环无法由ERP满足时，才建立轻量补充台账。')
doc.add_heading('目录',0); toc(); page()

doc.add_heading('第一篇 财务资料体系全景',0)
doc.add_heading('第1章 财务资料为什么不是“存文件”',1)
pic('panorama.png','图1 企业财务资料全景')
table(['资料层','回答的问题','典型使用者','失控后果'],[
['业务证据','交易是否真实、由谁批准、何时发生','业务、AP、AR、总账','错账、跨期、无法追责'],['系统记录','单据如何流转、何时过账、是否被修改','业务财务、IT','接口遗漏、重复、版本不一致'],['账簿与底稿','账是否完整、对账是否完成、判断是否复核','总账、主管、审计','月结失控、报表错报'],['税务资料','票、账、申报和缴税能否闭环','税务、总账、主管','少缴多缴、凭证不足'],['档案与交接','能否长期调阅并由新任接手','档案员、管理层','人员离职后业务瘫痪']],8)
doc.add_heading('第2章 资料分类、责任与来源',1)
rows=[[m['category'],m['name'],m['cycle'],m['source'],m['provider'],m['owner'],m['frequency']] for m in D['materials']]
table(['类别','资料','循环','来源系统','提供方','维护人','更新频率'],rows,6.9)
doc.add_heading('第3章 是否需要单独建Excel台账',1); pic('decision.png','图2 台账必要性判断')
table(['判断问题','建表信号','不建表信号','结论'],[
['系统能否提供完整字段','缺少责任人、期限、异常原因','字段完整且可导出','优先修系统或补轻量台账'],['是否跨多个系统','合同、ERP、银行和税局需桥接','全部在同一模块闭环','跨系统通常需要桥接底稿'],['是否需要复核留痕','需签字、锁版、解释差异','系统有审批日志和快照','保留复核证据'],['是否只是重复抄数','人工重复录入且无人使用','台账驱动催收或决策','无用途即停用']],7.5)
doc.add_heading('第4章 贯穿案例公司与月结假设',1)
table(['项目','设定'],[[k,v] for k,v in D['company'].items()],8)
para('贯穿案例以2026年6月月结为主。资料量按380人、制造工厂、多个银行账户、采购生产销售并存的企业估计；实际耗时受系统成熟度、业务量、异常率和团队分工影响。')
page()

doc.add_heading('第二篇 15类财务资料如何形成和使用',0)
cats=[]
for m in D['materials']:
    if m['category'] not in cats: cats.append(m['category'])
for idx,cat in enumerate(cats,1):
    doc.add_heading(f'第{idx+4}章 {cat}',1)
    ms=[m for m in D['materials'] if m['category']==cat]
    for m in ms:
        doc.add_heading(m['name'],2)
        table(['工作问题','实务答案'],[
            ['为什么需要',m['use']],['从哪里取得',f"{m['source']}；通常由{m['provider']}提供"],['谁负责',f"维护：{m['owner']}；复核：{m['reviewer']}"],['用什么工具',m['tool']],['何时整理',f"首次：{m['first']}；日常：{m['daily']}；月度：{m['month']}；季度：{m['quarter']}；年度：{m['year']}"],['预计耗时',f"{m['time']}，工作量：{m['workload']}"],['ERP与Excel',f"ERP覆盖：{m['erp']}；Excel建议：{m['excel']}"],['如何保存',m['save']],['保管口径',m['retention']],['敏感级别',m['sensitive']],['高频问题',m['problems']],['缺失后果',m['consequence']]],7.7)
page()

doc.add_heading('第三篇 核心台账和月结底稿',0)
doc.add_heading('第20章 20张核心台账设计',1)
for i,l in enumerate(D['ledger_specs'],1):
    doc.add_heading(f'{i:02d}. {l[0]}',2)
    table(['用途','来源','负责人/频率','核心字段','月结使用','预计耗时','常见异常','没有时的后果'],[[l[1],l[2],l[3],l[5],l[6],l[7],l[8],l[9]]],7.3)
doc.add_heading('第21章 月结资料包',1); pic('close.png','图3 月结资料包形成路径')
table(['编号','资料','频率','责任岗位','最小证据'],D['close_package'],7.5)
doc.add_heading('第22章 T-5至T+2催收和关闭日历',1)
table(['时间','动作','对象','输出','未完成处理'],D['chase'],7.4)
doc.add_heading('第23章 底稿复核的四层逻辑',1)
bullets(['完整性：系统范围、公司、期间、币种、过账状态是否齐全。','准确性：汇总数是否与子模块和总账一致。','判断性：暂估、预提、减值、重分类和截止依据是否充分。','可追溯性：从结论能否追到明细、原始单据、系统参数、编制人与复核人。'])
doc.add_heading('第24章 资料催收与未决事项闭环',1)
table(['场景','第一动作','升级条件','关闭证据'],[
['业务未交合同/验收','列明单据编号、金额和截止日','超过T-1或影响重大','正式合同/验收与入账结果'],['银行流水未取得','核对网银权限和接口状态','余额无法证明或疑似重复支付','对账单、调节表和复核签字'],['成本模块未结','锁定工单、报工、费用和接口异常','影响毛利或库存重大','成本运行日志、差异解释和GL对账'],['发票状态未清','区分未收票、待勾选、异常票','影响抵扣或申报','发票平台清单和申报勾稽']],7.6)
page()

doc.add_heading('第四篇 文件夹、命名、版本和档案',0)
doc.add_heading('第25章 标准文件夹与命名',1)
para('推荐命名：日期_公司_模块_资料名称_期间_版本_状态_编制人。例如：20260703_QCZN_GL_银行余额调节表_202606_v1.0_已复核_张三.xlsx。文件名承担检索作用，文件索引承担跨目录定位、版本和权限管理。')
table(['控制点','实务要求','禁止或高风险做法'],[
['版本','草稿v0.x、送审v1.0、修改v1.1、锁定FINAL；锁版后只通过变更记录修改','final_final2、覆盖已复核文件'],['路径','法人/年度/月度/循环/资料类型形成稳定层级','个人桌面作为唯一存储'],['索引','记录唯一ID、路径、版本、责任人、保密级别和保管期限','只靠文件名猜位置'],['权限','按岗位和最小权限授权；敏感工资、网银、税务资料隔离','共享盘全员可修改'],['电子原件','保留来源格式、元数据和验真信息','只打印或截图后删除电子原件']],7.7)
doc.add_heading('第26章 会计档案生命周期',1); pic('archive.png','图4 会计资料生命周期')
doc.add_heading('第27章 法定会计档案保管期限',1)
table(['档案类别','保管期限','层级/依据','起算/说明'],D['retention'],7.5)
para('【法规边界】上表来自财政部、国家档案局令第79号所附会计档案保管期限表。未了结债权债务会计凭证等，即使保管期满也不得销毁，应单独抽出保管至事项完结。电子会计凭证仅在满足来源真实、可验证、防篡改、可调阅和备份等条件时按电子形式归档。')
doc.add_heading('第28章 权限、保密与调阅',1)
table(['等级','资料例子','访问控制','外发要求'],[
['一般内部','流程制度、空白模板','员工只读、责任人可改','注明版本'],['敏感','客户供应商余额、合同价格、税表','岗位授权、下载留痕','脱敏或主管批准'],['高度敏感','工资明细、身份证、银行U盾、税务身份','极少数实名授权、禁止共用账号','加密、审批、接收确认'],['法定档案','凭证、账簿、报表、移交清册','档案制度和调阅登记','复制件标注用途']],7.5)
page()

doc.add_heading('第五篇 审计、税务和管理层资料包',0)
doc.add_heading('第29章 审计PBC资料包',1)
table(['阶段','资料包','财务动作','复核点'],[
['计划','TB、报表、科目表、流程、系统清单','锁定版本和公司范围','与总账最终版一致'],['实质性程序','明细账、台账、合同、发票、回函信息','提供索引和抽样证据','防止前后版本冲突'],['盘点与函证','盘点计划、银行/往来清单、控制表','确认地址、联系人、截止日','独立寄发和回函轨迹'],['调整与完成','审计调整、未更正差异、报表勾稽','跟踪每项是否入账或仅调表','管理层批准和版本锁定']],7.5)
doc.add_heading('第30章 税务资料包与账票税桥接',1)
table(['链条','数据来源','核对对象','典型差异'],[
['销售台账→开票→销项','销售/AR/开票系统','收入、应收、销项、申报表','未开票收入、红冲跨期、税会时点'],['采购→收票→进项','AP/发票平台','暂估、应付、进项、勾选','票到货未到、不可抵扣、待认证'],['利润→所得税','GL/纳税调整台账','利润总额、应纳税所得额、预缴','限额扣除、暂时性差异、优惠'],['工资→个税','HR/薪酬/银行','工资表、职工薪酬、实发、申报','人员变动、补发、专项扣除']],7.4)
doc.add_heading('第31章 管理层报告与经营分析资料',1)
para('管理层资料应从已锁定的财务版本取数，经营口径与法定报表口径若不同，必须保留桥接表。常见包包括损益桥、预算差异、现金预测、应收回款、库存周转、产品毛利和重大未决事项。')
page()

doc.add_heading('第六篇 岗位交接和接手验证',0)
doc.add_heading('第32章 岗位交接SOP',1); pic('handover.png','图5 财务岗位交接SOP')
para('【核心结论】交接完成不是“文件已发”。新任应在监督下重做一次银行调节、子模块对账、关键税务勾稽或月结任务，以结果验证能否独立执行。')
doc.add_heading('第33章 入职第一周',1)
table(['日期','目标','应取得资料','应完成验证'],[
['第1天','认识组织和关账节奏','组织图、岗位表、月结日历','确认公司、期间和责任边界'],['第2天','理解系统和权限','系统清单、接口图、角色权限','登录并导出关键报表'],['第3天','理解账和台账','科目表、TB、台账、上月底稿','从TB追到一个明细和源单'],['第4天','理解税务和资金','税种、账户、申报、银行调节','核对一项申报数和一张调节表'],['第5天','理解未决事项','长期挂账、审计调整、税务风险','形成接手风险清单和行动计划']],7.5)
doc.add_heading('第34章 100个核心交接问题',1)
groups=[]
for q in D['questions']:
    if q['category'] not in groups: groups.append(q['category'])
for g in groups:
    doc.add_heading(g,2)
    qs=[q for q in D['questions'] if q['category']==g]
    table(['编号','必须问的问题','为什么问','应取得的证据','初始状态'],[[q['id'],q['question'],q['why'],q['evidence'],q['status']] for q in qs],6.8)
doc.add_heading('第35章 交接签字、责任边界与后续观察期',1)
table(['事项','建议文本或动作'],[
['截止日','双方确认资料、余额和未决事项截至YYYY-MM-DD。'],['责任边界','签字证明已按清单移交和接收，不免除交接前后各自依法及依职权承担的责任。'],['未决事项','逐项列明金额、原因、下一动作、责任人和期限，不以口头说明替代。'],['权限','移交仅办理新任授权；离任账号、网银、税务身份按制度回收或停用。'],['观察期','建议覆盖至少一次月结；重大税务、审计或融资事项可延长并指定答疑人。']],7.6)
doc.add_heading('第36章 离职交接清单',1)
bullets(['锁定最后工作日和交接基准日；导出系统权限、银行权限、税务身份及设备清单。','完成本期可完成工作，未完成事项形成结构化问题清单。','移交正式路径中的文件，不以个人微信、私人网盘或个人桌面为唯一载体。','由接任者现场演示关键任务，主管抽查余额、版本和证据。','签字后回收权限和实物介质，并保留交接清册。'])
page()

doc.add_heading('第七篇 团队分工、问题案例和自动化',0)
doc.add_heading('第37章 不同规模企业的最小配置',1)
table(['企业阶段','岗位现实','最少应保留的10项台账/底稿','补偿控制'],[
['小型企业','总账兼税务、资金或报表','月结清单、银行调节、AR/AP、暂估、预提、税务、固定资产、工资核对、未决事项','老板/外部顾问独立复核付款和申报'],['中型企业','AP、AR、成本、总账、税务逐步分工','在最少10项上增加成本、库存、合同、关联方、预算和PBC','财务主管统一关账日历和版本'],['大型集团','共享中心、业务财务、集团报告分层','ERP工作流、数据仓库、标准底稿和档案平台','系统权限、接口监控和集团一致口径']],7.2)
doc.add_heading('第38章 50个真实问题与补救',1)
groups=[]
for x in D['problems']:
    if x['category'] not in groups: groups.append(x['category'])
for g in groups:
    doc.add_heading(g,2)
    ps=[x for x in D['problems'] if x['category']==g]
    table(['编号','场景','根因','风险','补救','预防'],[[x['id'],x['scenario'],x['cause'],x['risk'],x['remedy'],x['prevention']] for x in ps],6.7)
doc.add_heading('第39章 AI、RPA和Power Query适用边界',1)
table(['任务','适合工具','可自动化内容','必须人工判断'],[
['文件归档','RPA/脚本','按元数据命名、移动、生成索引','资料是否真实、保密级别、最终销毁'],['资料催收','流程机器人','按清单和截止日提醒、升级','重大程度和跨部门沟通'],['台账更新','Power Query/接口','合并ERP、银行、发票数据','字段映射、异常原因、会计结论'],['月结检查','Python/规则引擎','负数余额、长期挂账、异常变动扫描','暂估、减值、截止、重分类'],['合同阅读','AI辅助','提取主体、金额、付款、履约条款','收入确认、税务和法律结论'],['交接问答','知识库/AI检索','定位制度、流程和历史问题','权限审批、责任确认、敏感信息处理']],7.2)
doc.add_heading('第40章 案例公司一个月的资料时间分配',1)
table(['事项','预计耗时','责任岗位','说明'],D['time_allocation'],7.5)
table(['工作阶段','预计耗时','工作量','完成证据','自动化建议'],D['time_map'],7.5)
para('耗时是案例公司的容量估计，用于排班和识别低效点，不是行业定额。成熟接口可显著减少搬运数据时间，但异常调查和专业判断通常仍是月结瓶颈。')

doc.add_heading('附录A 官方依据与版本控制',0)
table(['文件','发布机关/文号','效力/状态','本手册用途','官方链接'],D['sources'],6.9)
doc.add_heading('附录B 财务主管月末复核十问',0)
bullets(['本月所有关键系统是否已按统一截止日关闭？','银行、AR、AP、存货、成本、固定资产、薪酬和税务是否与GL核对？','暂估、预提、待摊和减值是否有业务依据并经复核？','长期挂账和负数余额是否逐项说明经济实质？','手工凭证和重大调整是否有独立审批？','税务申报数据能否从申报表追到账、票和业务台账？','报表与管理层报告是否引用同一锁定版本？','未决事项是否有责任人、期限和升级记录？','档案是否包含电子原件、来源信息、编制复核和版本？','若关键人员明天离职，接任者能否依清单完成一次月结？'])

OUT.mkdir(exist_ok=True)
path=OUT/'企业财务资料、台账、底稿与交接实操手册.docx'; doc.save(path)
print(path)
