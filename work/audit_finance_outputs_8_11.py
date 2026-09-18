import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\HZJ\Desktop\Git\work")
for episode in (8, 9, 10, 11):
    docx_path = ROOT / f"齐昊老师_学员实战案例_第{episode}集_文字思维导图.docx"
    doc = Document(docx_path)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "TODO" not in all_text and "PLACEHOLDER" not in all_text
    headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    numbered = [
        p for p in doc.paragraphs
        if p._p.pPr is not None and p._p.pPr.numPr is not None
    ]
    assert headings and numbered
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:tblW")).get(qn("w:w")) == "9360"
        assert tbl_pr.find(qn("w:tblInd")).get(qn("w:w")) == "120"
        grid = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
        assert sum(grid) == 9360
        for row in table.rows:
            assert [int(cell._tc.tcPr.tcW.get(qn("w:w"))) for cell in row.cells] == grid

    xmind_path = ROOT / f"齐昊老师_学员实战案例_第{episode}集_思维导图.xmind"
    with zipfile.ZipFile(xmind_path) as archive:
        assert archive.testzip() is None
        content = json.loads(archive.read("content.json"))
        assert len(content) == 2
        assert {sheet["title"] for sheet in content} == {"01 总览图", "02 详细知识树"}
    print(
        f"EP{episode}\theadings={len(headings)}\tlists={len(numbered)}\t"
        f"tables={len(doc.tables)}\tsheets={len(content)}"
    )
