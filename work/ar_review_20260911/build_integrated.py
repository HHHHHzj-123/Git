from pathlib import Path
from copy import deepcopy, copy
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table

BASE=Path(__file__).resolve().parent
ORIG=Path(r'C:\Users\HZJ\Desktop\Git\work\ar_management_20260911\deliverables')
SUPP=BASE/'deliverables'
OUT=BASE/'integrated'; OUT.mkdir(parents=True,exist_ok=True)

def ptext(el):
    return ''.join(t.text or '' for t in el.findall('.//'+qn('w:t'))).strip()

def style_name(el):
    ppr=el.find(qn('w:pPr'))
    if ppr is None: return ''
    ps=ppr.find(qn('w:pStyle'))
    return '' if ps is None else (ps.get(qn('w:val')) or '')

def is_heading(el, max_level=None):
    if el.tag != qn('w:p'): return False
    s=style_name(el)
    if not s.startswith('Heading'): return False
    try: level=int(s.replace('Heading',''))
    except: return False
    return level <= max_level if max_level else True

def section_elements(doc, start_text, end_at_same_or_higher=True):
    els=list(doc.element.body)
    start=None; level=None
    for i,el in enumerate(els):
        if el.tag==qn('w:p') and ptext(el).startswith(start_text):
            start=i; s=style_name(el); level=int(s.replace('Heading','')) if s.startswith('Heading') else 2; break
    if start is None: raise ValueError(start_text)
    end=len(els)-1
    for j in range(start+1,len(els)):
        if els[j].tag==qn('w:p') and is_heading(els[j],level): end=j; break
    return [deepcopy(x) for x in els[start:end]]

def rename_first(blocks,new_text):
    for el in blocks:
        if el.tag==qn('w:p'):
            ts=el.findall('.//'+qn('w:t'))
            if ts:
                ts[0].text=new_text
                for t in ts[1:]: t.text=''
                return blocks
    return blocks

def find_heading(doc,prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

def insert_before(doc,target_prefix,blocks):
    target=find_heading(doc,target_prefix)._p
    for el in blocks: target.addprevious(el)

def insert_before_prior_page_break(doc,target_prefix,blocks):
    target=find_heading(doc,target_prefix)._p
    prev=target.getprevious(); page_break=None
    while prev is not None:
        if prev.findall('.//'+qn('w:br')):
            page_break=prev; break
        prev=prev.getprevious()
    anchor=page_break if page_break is not None else target
    for el in blocks: anchor.addprevious(el)

def append_before_sectpr(doc,blocks):
    body=doc.element.body; sect=body.find(qn('w:sectPr'))
    for el in blocks:
        if sect is not None: sect.addprevious(el)
        else: body.append(el)

def integrate_docx():
    main=Document(ORIG/'应收账款管理实战手册.docx')
    supp=Document(SUPP/'应收账款管理第一版复核与补强报告.docx')
    # Put each supplement into the chapter where it is used; no separate “supplement” part remains.
    mapping=[
        ('3.4','2.3 责任边界与销售绩效','第三章'),
        ('3.1','5.4 客户信用调查、评分、额度和账期','第六章'),
        ('3.2','6.1.1 会计账龄和逾期账龄必须分开','6.2'),
        ('3.3','6.1.2 催收升级SLA','6.2'),
        ('3.5','7.4 对账差异与复杂回款核销','第八章'),
        ('3.6','8.4 坏账的管理、会计与税务三条线','第九章'),
        ('3.7','8.5 ECL实操底稿','第九章'),
        ('3.8','9.3 风险预警、集中度与指标','第十章'),
        ('3.9','9.4 应收增长与现金流量分析树','第十章'),
        ('3.10','9.5 应收融资选择','第十章'),
        ('3.12','10.2 十五个复杂案例','第十一章'),
        ('3.11','14.3 应收周会SOP','第十五章'),
    ]
    for start,new_title,target in mapping:
        blocks=rename_first(section_elements(supp,start),new_title)
        insert_before(main,target,blocks)
    source_blocks=rename_first(section_elements(supp,'第五部分'),'附录C 补充依据与适用边界')
    append_before_sectpr(main,source_blocks)
    # Add a static chapter index so LibreOffice/WPS preview is useful even before TOC field refresh.
    mini=Document(); p=mini.add_paragraph('若自动目录未显示，可在WPS中按 Ctrl+A 后按 F9 更新；也可通过左侧导航栏按标题跳转。')
    t=mini.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    t.rows[0].cells[0].text='章节'; t.rows[0].cells[1].text='主题'
    chapters=['应收账款管理全景','岗位责任与跨部门协同','财务主管的工作节奏','单据、系统、数据和台账','信用管理和销售放行','到期管理、催收和争议','回款认领、核销和对账','会计核算与预期信用损失','核心指标和深层原因','高频异常案例库','内部控制、舞弊和红旗信号','新财务主管接手','月末Checklist','财务经理和管理层视角','20%核心知识']
    for i,x in enumerate(chapters,1):
        cells=t.add_row().cells; cells[0].text=f'第{i}章'; cells[1].text=x
    blocks=[deepcopy(mini.paragraphs[0]._p),deepcopy(mini.tables[0]._tbl)]
    insert_before_prior_page_break(main,'第一章',blocks)
    # Update title/subtitle/date without changing heading hierarchy.
    for p in main.paragraphs:
        if p.text.strip()=='应收账款管理实战手册':
            p.text='应收账款管理实战手册（融合修订版）'; p.style='Title'
        elif '总账会计 财务主管 财务经理工作版' in p.text:
            p.text='总账会计 · 财务主管 · 财务经理工作版｜已融合第一版复核补强内容'; p.style='Subtitle'
        elif '版本日期 2026年9月11日' in p.text:
            p.text='版本日期 2026年9月12日'; p.runs[0].bold=True
    out=OUT/'应收账款管理实战手册（融合修订版）.docx'; main.save(out); return out

def copy_sheet(src_ws,dst_wb,new_name):
    dst=dst_wb.create_sheet(new_name)
    for row in src_ws.iter_rows():
        for c in row:
            nc=dst[c.coordinate]; nc.value=c.value
            if c.has_style:
                nc._style=copy(c._style); nc.font=copy(c.font); nc.fill=copy(c.fill); nc.border=copy(c.border); nc.alignment=copy(c.alignment); nc.number_format=c.number_format; nc.protection=copy(c.protection)
            if c.hyperlink: nc._hyperlink=copy(c.hyperlink)
            if c.comment: nc.comment=copy(c.comment)
    for k,v in src_ws.column_dimensions.items():
        dst.column_dimensions[k].width=v.width; dst.column_dimensions[k].hidden=v.hidden; dst.column_dimensions[k].outlineLevel=v.outlineLevel
    for k,v in src_ws.row_dimensions.items():
        dst.row_dimensions[k].height=v.height; dst.row_dimensions[k].hidden=v.hidden; dst.row_dimensions[k].outlineLevel=v.outlineLevel
    for rng in src_ws.merged_cells.ranges: dst.merge_cells(str(rng))
    dst.freeze_panes=src_ws.freeze_panes; dst.sheet_view.showGridLines=src_ws.sheet_view.showGridLines
    dst.auto_filter.ref=src_ws.auto_filter.ref
    for dv in src_ws.data_validations.dataValidation: dst.add_data_validation(deepcopy(dv))
    # Preserve conditional formatting rules.
    for cf,rules in src_ws.conditional_formatting._cf_rules.items():
        for rule in rules: dst.conditional_formatting.add(str(cf.sqref),deepcopy(rule))
    # Preserve Excel tables under unique names.
    for t in src_ws.tables.values():
        nt=deepcopy(t); old_name=getattr(t,'name','Table'); nt.name=f'INT_{new_name.split("-")[0]}_{old_name}'.replace(' ','_'); nt.displayName=nt.name; dst.add_table(nt)
    dst.sheet_properties=copy(src_ws.sheet_properties)
    dst.page_margins=copy(src_ws.page_margins); dst.page_setup=copy(src_ws.page_setup); dst.print_options=copy(src_ws.print_options)
    return dst

def integrate_xlsx():
    base=load_workbook(ORIG/'应收账款管理工具包.xlsx')
    sup=load_workbook(SUPP/'应收账款管理补强工具包.xlsx')
    selected=[
        ('02-客户信用评分','16-客户信用评分'),('03-账龄与逾期账龄','17-双账龄分析'),('04-催收升级SLA','18-催收升级SLA'),
        ('05-催收记录','19-催收记录'),('06-风险预警','20-风险预警'),('07-集中度分析','21-集中度分析'),
        ('08-坏账与税务','22-坏账与税务'),('09-ECL参数与迁徙','23-ECL参数迁徙'),('10-应收月报','24-应收月报'),
        ('11-客户对账函','25-客户对账函'),('12-催款函','26-催款函'),('13-回款核销异常','27-回款核销异常'),
        ('14-销售绩效挂钩','28-销售绩效挂钩'),('15-应收会议','29-应收会议'),('16-融资评估','30-融资评估'),
        ('17-应收增长分析树','31-应收增长分析'),('18-补强案例','32-复杂案例'),
    ]
    for old,new in selected: copy_sheet(sup[old],base,new)
    # Remap formulas that referred to the supplement workbook's old sheet names.
    m=base['24-应收月报']
    m['B6']="=SUM('21-集中度分析'!B6:B11)"; m['D7']="='21-集中度分析'!C6"; m['F7']="=SUM('21-集中度分析'!C6:C10)"; m['H7']="='23-ECL参数迁徙'!B12"
    base.calculation.fullCalcOnLoad=True; base.calculation.forceFullCalc=True; base.calculation.calcMode='auto'
    out=OUT/'应收账款管理完整工具包（融合修订版）.xlsx'; base.save(out); return out

if __name__=='__main__':
    print(integrate_docx()); print(integrate_xlsx())
