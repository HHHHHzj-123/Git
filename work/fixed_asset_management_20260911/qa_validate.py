import json,zipfile
from pathlib import Path
from openpyxl import load_workbook
root=Path(r'C:\Users\HZJ\Desktop\Git\work\fixed_asset_management_20260911');d=json.loads((root/'content.json').read_text(encoding='utf-8'));p=root/'deliverables'/'企业固定资产管理工具包.xlsx'
cost=sum(x[8] for x in d['assets']);dep=sum(x[16] for x in d['assets']);net=sum(x[17] for x in d['assets']);idle=sum(x[17] for x in d['assets'] if x[22]=='闲置');missing=sum(x[17] for x in d['assets'] if x[22]=='盘亏待查');expected=[cost,dep,net,idle,idle/net,missing,6000000,2,0,2,0]
wb=load_workbook(p,data_only=True,read_only=True);ws=wb['00-管理看板'];actual=[ws[f'B{i}'].value for i in range(6,17)]
for i,(a,e) in enumerate(zip(actual,expected),6):assert a is not None and abs(float(a)-float(e))<.01,(i,a,e)
wbf=load_workbook(p,data_only=False,read_only=True);formulas=0;errs=[]
for sh in wbf.worksheets:
 for row in sh.iter_rows():
  for c in row:
   v=c.value
   if isinstance(v,str) and v.startswith('='):formulas+=1
   if isinstance(v,str) and v in {'#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A'}:errs.append((sh.title,c.coordinate,v))
assert not errs,errs
for q in [p,root/'deliverables'/'企业固定资产管理实战手册.docx',root/'deliverables'/'企业固定资产管理全景.xmind']:
 assert zipfile.is_zipfile(q);z=zipfile.ZipFile(q);assert not z.testzip();z.close()
with zipfile.ZipFile(root/'deliverables'/'企业固定资产管理全景.xmind') as z:assert len(json.loads(z.read('content.json')))==3
print({'dashboard':actual,'sheets':len(wb.sheetnames),'formulas':formulas,'errors':errs,'doc_pages':len(list((root/'qa'/'doc_render').glob('page-*.png')))})
